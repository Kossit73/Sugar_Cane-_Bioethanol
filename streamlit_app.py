"""Streamlit front-end for the Sugar Cane Bioethanol multi-product finance model.

This app wraps the `model.py` engine and exposes interactive controls for
adjusting key drivers and reviewing the resulting financial outputs, dashboards,
sensitivities, and scenarios.

Run with:
    streamlit run streamlit_app.py
"""

from __future__ import annotations

import copy
import json
import importlib
import math
import re
import sys
import concurrent.futures
import ast
import urllib.parse
import urllib.request
import datetime
import time
from dataclasses import dataclass
from collections import OrderedDict
from contextlib import contextmanager
from io import BytesIO
from typing import Callable, Dict, Iterable, List, Mapping, Optional, Sequence, Tuple

import numpy as np
import pandas as pd
import streamlit as st
from streamlit.errors import StreamlitAPIException
from pandas.testing import assert_frame_equal

from dependencies import ensure_package, get_package_error

MATPLOTLIB_INSTALL_ERROR: Optional[str]
if ensure_package("matplotlib"):
    try:  # pragma: no cover - optional dependency
        import matplotlib.dates as mdates
        import matplotlib.pyplot as plt
    except Exception as exc:  # pragma: no cover - unexpected import failure
        MATPLOTLIB_INSTALL_ERROR = str(exc)
        plt = None
        mdates = None
else:
    MATPLOTLIB_INSTALL_ERROR = get_package_error("matplotlib")
    plt = None
    mdates = None

_MATPLOTLIB_WARNING_SHOWN = False


CHAT_PROVIDER_DEFAULTS: Dict[str, Dict[str, object]] = {
    "OpenAI": {"base_url": "https://api.openai.com/v1/chat/completions", "supports_reasoning": True, "supports_tools": True, "supports_web_search": True},
    "Anthropic": {"base_url": "https://api.anthropic.com/v1/messages", "supports_reasoning": True, "supports_tools": True, "supports_web_search": False},
    "Google Gemini": {"base_url": "", "supports_reasoning": True, "supports_tools": True, "supports_web_search": True},
    "Mistral": {"base_url": "https://api.mistral.ai/v1/chat/completions", "supports_reasoning": True, "supports_tools": True, "supports_web_search": False},
    "Cohere": {"base_url": "https://api.cohere.com/v2/chat", "supports_reasoning": True, "supports_tools": True, "supports_web_search": False},
    "DeepSeek": {"base_url": "https://api.deepseek.com/v1/chat/completions", "supports_reasoning": True, "supports_tools": True, "supports_web_search": False},
    "xAI": {"base_url": "https://api.x.ai/v1/chat/completions", "supports_reasoning": True, "supports_tools": True, "supports_web_search": True},
    "Llama-compatible": {"base_url": "http://localhost:8000/v1/chat/completions", "supports_reasoning": False, "supports_tools": False, "supports_web_search": False},
}

CHAT_EVAL_PROMPTS: List[Dict[str, str]] = [
    {"prompt": "Estimate DSCR sensitivity if revenue drops by 10%.", "expected_tool": "tool_calc"},
    {"prompt": "Compare lender cases by Project_NPV and explain the difference.", "expected_tool": "tool_scenario_compare"},
    {"prompt": "Aggregate revenue by product and summarize concentration risk.", "expected_tool": "tool_dataframe_aggregate"},
    {"prompt": "Should we prioritize IRR or DSCR for this lender discussion?", "expected_tool": "none"},
    {"prompt": "Benchmark ethanol price assumptions against market references.", "expected_tool": "web"},
    {"prompt": "What are the key assumptions behind the current NPV?", "expected_tool": "none"},
    {"prompt": "Run a quick check on capex shock impact using a simple formula.", "expected_tool": "tool_calc"},
    {"prompt": "Explain why DSCR falls in downside cases.", "expected_tool": "none"},
    {"prompt": "Compare reserve account implications across lender cases.", "expected_tool": "tool_scenario_compare"},
    {"prompt": "Summarize top risks and practical mitigations.", "expected_tool": "none"},
    {"prompt": "Aggregate monthly production by product and identify weak spots.", "expected_tool": "tool_dataframe_aggregate"},
    {"prompt": "Check if current assumptions are conservative versus industry norms.", "expected_tool": "web"},
    {"prompt": "Which KPI should govern covenant negotiations?", "expected_tool": "none"},
    {"prompt": "Calculate breakeven change for a 5% opex increase.", "expected_tool": "tool_calc"},
    {"prompt": "Compare Project_IRR across available lender cases.", "expected_tool": "tool_scenario_compare"},
    {"prompt": "Provide a recommendation with caveats for debt structuring.", "expected_tool": "none"},
    {"prompt": "Benchmark inflation assumptions with external references.", "expected_tool": "web"},
    {"prompt": "Aggregate cash waterfall components and flag dominant drains.", "expected_tool": "tool_dataframe_aggregate"},
    {"prompt": "Evaluate causality between capex delays and covenant stress.", "expected_tool": "none"},
    {"prompt": "Give a concise action plan for improving both NPV and DSCR.", "expected_tool": "none"},
]

INSTITUTION_WHITELIST: Tuple[Tuple[str, str], ...] = (
    ("iea.org", "International Energy Agency"),
    ("worldbank.org", "World Bank"),
    ("imf.org", "International Monetary Fund"),
    ("oecd.org", "OECD"),
    ("ifc.org", "International Finance Corporation"),
    ("reuters.com", "Reuters"),
    ("bloomberg.com", "Bloomberg"),
    ("fao.org", "FAO"),
    ("un.org", "United Nations"),
    ("gov", "Government source"),
    ("edu", "Academic source"),
)

BENCHMARK_METRIC_DEFINITIONS: Dict[str, Tuple[str, ...]] = {
    "Project_IRR": ("internal rate of return", "irr", "equity irr", "project irr"),
    "DSCR_min": ("debt service coverage ratio", "dscr", "coverage ratio"),
    "Project_NPV": ("net present value", "npv", "discounted cash flow"),
}


@dataclass
class ChatProviderSettings:
    provider_name: str
    model_name: str
    api_key: str
    base_url: str
    temperature: float
    max_tokens: int
    reasoning_mode: str
    use_tools: bool
    use_web_search: bool


class BaseChatAdapter:
    """Shared provider adapter interface."""

    provider_name: str = "base"

    def __init__(self, settings: ChatProviderSettings):
        self.settings = settings

    def supports_tools(self) -> bool:
        return _chat_supports(self.settings.provider_name, "supports_tools")

    def supports_reasoning(self) -> bool:
        return _chat_supports(self.settings.provider_name, "supports_reasoning")

    def _request_with_retries(self, url: str, payload: Mapping[str, object], headers: Mapping[str, str]) -> Dict[str, object]:
        """Issue HTTP request with retry/backoff and basic rate-limit handling."""

        max_attempts = 3
        for attempt in range(max_attempts):
            req = urllib.request.Request(
                url,
                data=json.dumps(payload).encode("utf-8"),
                headers=dict(headers),
                method="POST",
            )
            try:
                with urllib.request.urlopen(req, timeout=30) as response:
                    return json.loads(response.read().decode("utf-8"))
            except Exception as exc:
                message = str(exc).lower()
                is_retryable = "429" in message or "rate" in message or "timed out" in message or "temporarily" in message
                if attempt == max_attempts - 1 or not is_retryable:
                    raise
                time.sleep(0.6 * (2 ** attempt))
        return {}

    def _capability_check(self) -> Optional[str]:
        model_name = self.settings.model_name.lower()
        if "reasoning" in model_name and not self.supports_reasoning():
            return f"Selected model '{self.settings.model_name}' may not support reasoning mode on {self.settings.provider_name}."
        return None

    def generate(self, messages: List[Dict[str, str]]) -> str:
        raise NotImplementedError

    def stream(self, messages: List[Dict[str, str]]) -> Iterable[str]:
        text = self.generate(messages)
        yield text


class OpenAIAdapter(BaseChatAdapter):
    provider_name = "OpenAI"

    def generate(self, messages: List[Dict[str, str]]) -> str:
        warning = self._capability_check()
        payload: Dict[str, object] = {
            "model": self.settings.model_name,
            "messages": messages,
            "temperature": float(self.settings.temperature),
            "max_tokens": int(self.settings.max_tokens),
        }
        if self.settings.reasoning_mode and self.supports_reasoning():
            payload["reasoning"] = {"effort": self.settings.reasoning_mode}
        data = self._request_with_retries(
            self.settings.base_url,
            payload,
            {"Content-Type": "application/json", "Authorization": f"Bearer {self.settings.api_key}"},
        )
        content = _extract_text_from_provider_payload(data)
        if warning and content:
            return f"{warning}\n\n{content}"
        return content


class AnthropicAdapter(BaseChatAdapter):
    provider_name = "Anthropic"

    def generate(self, messages: List[Dict[str, str]]) -> str:
        warning = self._capability_check()
        system_msgs = [m.get("content", "") for m in messages if m.get("role") == "system"]
        non_system = [{"role": m.get("role", "user"), "content": m.get("content", "")} for m in messages if m.get("role") != "system"]
        payload: Dict[str, object] = {
            "model": self.settings.model_name,
            "messages": non_system,
            "max_tokens": int(self.settings.max_tokens),
            "temperature": float(self.settings.temperature),
        }
        if system_msgs:
            payload["system"] = "\n\n".join(system_msgs)
        data = self._request_with_retries(
            self.settings.base_url,
            payload,
            {
                "Content-Type": "application/json",
                "x-api-key": self.settings.api_key,
                "anthropic-version": "2023-06-01",
            },
        )
        content = _extract_text_from_provider_payload(data)
        if warning and content:
            return f"{warning}\n\n{content}"
        return content


class GeminiAdapter(BaseChatAdapter):
    provider_name = "Google Gemini"

    def generate(self, messages: List[Dict[str, str]]) -> str:
        warning = self._capability_check()
        # Gemini native endpoint shape varies by SDK/version; this adapter assumes
        # an OpenAI-compatible gateway/base_url for portability in this app.
        payload: Dict[str, object] = {
            "model": self.settings.model_name,
            "messages": messages,
            "temperature": float(self.settings.temperature),
            "max_tokens": int(self.settings.max_tokens),
        }
        data = self._request_with_retries(
            self.settings.base_url,
            payload,
            {"Content-Type": "application/json", "Authorization": f"Bearer {self.settings.api_key}"},
        )
        content = _extract_text_from_provider_payload(data)
        if warning and content:
            return f"{warning}\n\n{content}"
        return content


class GenericOpenAICompatibleAdapter(BaseChatAdapter):
    provider_name = "generic"

    def generate(self, messages: List[Dict[str, str]]) -> str:
        warning = self._capability_check()
        payload: Dict[str, object] = {
            "model": self.settings.model_name,
            "messages": messages,
            "temperature": float(self.settings.temperature),
            "max_tokens": int(self.settings.max_tokens),
        }
        if self.settings.reasoning_mode and self.supports_reasoning():
            payload["reasoning"] = {"effort": self.settings.reasoning_mode}
        data = self._request_with_retries(
            self.settings.base_url,
            payload,
            {"Content-Type": "application/json", "Authorization": f"Bearer {self.settings.api_key}"},
        )
        content = _extract_text_from_provider_payload(data)
        if warning and content:
            return f"{warning}\n\n{content}"
        return content


def _streamlit_runtime_exists() -> bool:
    """Return True when executed inside an active Streamlit runtime."""
    try:  # Streamlit >= 1.35 exposes runtime.exists()
        from streamlit.runtime import exists as runtime_exists  # type: ignore import

        return bool(runtime_exists())
    except Exception:  # pragma: no cover - fallback for older versions
        try:
            from streamlit.runtime.scriptrunner import get_script_run_ctx  # type: ignore

            return get_script_run_ctx() is not None
        except Exception:
            return False


def _safe_rerun() -> None:
    """Attempt to trigger a Streamlit rerun only when a runtime exists."""

    if not _streamlit_runtime_exists():
        return
    try:
        st.experimental_rerun()
    except StreamlitAPIException:  # pragma: no cover - defensive catch
        pass
    except Exception:
        # Older Streamlit builds or embedded executions may surface bespoke
        # rerun exceptions; swallow them so non-Streamlit contexts keep running.
        pass


def _editor_state_key(table_name: str) -> str:
    """Return the canonical widget state key for a given table editor."""

    return f"table_editor__{table_name}"


def _update_editor_state(table_name: str, tables: "InputTables") -> None:
    """Synchronise the Streamlit data editor state with the backing table."""

    _ = tables  # retained for signature compatibility
    state_key = _editor_state_key(table_name)
    # Streamlit forbids direct writes to widget-managed keys once the widget is
    # instantiated. Clearing the state entry ensures the next render picks up
    # the refreshed DataFrame without violating session-state policies. Also
    # purge the legacy key from earlier builds to avoid conflicts.
    st.session_state.pop(state_key, None)
    st.session_state.pop(f"editor_{table_name}", None)


def _sync_table_mutation_state(table_name: str, tables: "InputTables") -> None:
    """Keep editor/default/cache state coherent after table mutations."""

    _update_editor_state(table_name, tables)
    # Clear derived caches so downstream pages (scenarios/downloads) recompute
    # from the latest edited inputs.
    st.session_state.pop("scenario_payload_cache", None)
    st.session_state.pop("excel_bytes_map", None)
    # Close default row editor if active table changed/reset.
    edit_state = st.session_state.get(DEFAULT_EDIT_STATE_KEY)
    if isinstance(edit_state, Mapping) and edit_state.get("table") == table_name:
        st.session_state.pop(DEFAULT_EDIT_STATE_KEY, None)


def _total_capex_from_inputs(tables: "InputTables") -> float:
    """Return the aggregate CAPEX amount from the current input schedule."""

    try:
        capex_df = tables.ensure_table("capex_lines").copy()
    except KeyError:  # pragma: no cover - defensive guard if schema missing
        return 0.0

    if capex_df.empty:
        return 0.0

    amounts = pd.to_numeric(capex_df.get("amount"), errors="coerce").fillna(0.0)
    return float(amounts.sum())


def _option_index(options: Sequence[str], value: str, default: int = 0) -> int:
    """Return the index of ``value`` inside ``options`` with a safe fallback."""

    try:
        return options.index(value)
    except ValueError:
        return default if 0 <= default < len(options) else 0


def _float_option_values(start: float, stop: float, step: float, digits: int = 6) -> List[float]:
    """Return an inclusive list of float options suitable for dropdown widgets."""

    if step <= 0:
        return [round(start, digits)]
    count = int(round((stop - start) / step))
    if count < 0:
        return [round(start, digits)]
    values = [round(start + idx * step, digits) for idx in range(count + 1)]
    if values and values[-1] < round(stop, digits):
        values.append(round(stop, digits))
    return values


def _chat_supports(provider_name: str, capability: str) -> bool:
    defaults = CHAT_PROVIDER_DEFAULTS.get(provider_name, {})
    return bool(defaults.get(capability, False))


def _run_sandbox(code: str) -> Dict[str, object]:
    """Execute controlled Python snippets in a restricted sandbox namespace."""

    risky_tokens = [
        "__",
        "import ",
        "open(",
        "exec(",
        "eval(",
        "os.",
        "sys.",
        "subprocess",
        "socket",
        "shutil",
        "pathlib",
        "requests",
        "urllib",
        "globals(",
        "locals(",
    ]
    lowered = code.lower()
    denied = [token for token in risky_tokens if token in lowered]
    if denied:
        return {"used": True, "ok": False, "error": f"Blocked by sandbox denial list: {', '.join(denied)}"}
    if len(code) > 4000:
        return {"used": True, "ok": False, "error": "Code exceeds sandbox size limit (4000 chars)."}

    safe_builtins = {
        "abs": abs,
        "min": min,
        "max": max,
        "sum": sum,
        "len": len,
        "range": range,
        "round": round,
        "sorted": sorted,
    }
    restricted_globals = {"__builtins__": safe_builtins, "math": math, "np": np, "pd": pd}

    def _execute() -> Dict[str, object]:
        restricted_locals: Dict[str, object] = {}
        exec(code, restricted_globals, restricted_locals)
        output = restricted_locals.get("result")
        output_text = str(output)
        if len(output_text) > 1200:
            output_text = output_text[:1200] + "... [truncated]"
        return {"used": True, "ok": True, "result": output_text, "locals": list(restricted_locals.keys())}

    try:
        with concurrent.futures.ThreadPoolExecutor(max_workers=1) as executor:
            return executor.submit(_execute).result(timeout=2.0)
    except concurrent.futures.TimeoutError:
        return {"used": True, "ok": False, "error": "Sandbox execution timed out (2s)."}
    except Exception as exc:
        return {"used": True, "ok": False, "error": str(exc)}


def _tool_calc(expression: str) -> Dict[str, object]:
    """Safely evaluate arithmetic-style expressions."""

    expr = expression.strip()
    if not expr:
        return {"used": True, "ok": False, "error": "No expression provided."}
    if len(expr) > 500:
        return {"used": True, "ok": False, "error": "Expression too long."}
    blocked_tokens = ["__", "import", "open(", "exec(", "eval(", "lambda", "os.", "sys.", "subprocess"]
    if any(token in expr.lower() for token in blocked_tokens):
        return {"used": True, "ok": False, "error": "Expression blocked by security policy."}
    try:
        node = ast.parse(expr, mode="eval")
        allowed_nodes = (
            ast.Expression,
            ast.BinOp,
            ast.UnaryOp,
            ast.Constant,
            ast.Num,
            ast.Add,
            ast.Sub,
            ast.Mult,
            ast.Div,
            ast.Pow,
            ast.Mod,
            ast.FloorDiv,
            ast.USub,
            ast.UAdd,
            ast.Load,
            ast.Call,
            ast.Name,
        )
        allowed_names = {"abs": abs, "round": round, "min": min, "max": max}
        for child in ast.walk(node):
            if not isinstance(child, allowed_nodes):
                return {"used": True, "ok": False, "error": f"Unsupported expression element: {type(child).__name__}"}
            if isinstance(child, ast.Name) and child.id not in allowed_names:
                return {"used": True, "ok": False, "error": f"Unknown identifier: {child.id}"}
        result = eval(compile(node, "<tool_calc>", "eval"), {"__builtins__": {}}, allowed_names)
        return {"used": True, "ok": True, "result": str(result)}
    except Exception as exc:
        return {"used": True, "ok": False, "error": str(exc)}


def _tool_dataframe_aggregate(model_results: Mapping[str, object], table_name: str, group_by: str, metric: str, agg: str) -> Dict[str, object]:
    """Aggregate a model output DataFrame with basic guardrails."""

    table = model_results.get(table_name)
    if not isinstance(table, pd.DataFrame) or table.empty:
        return {"used": True, "ok": False, "error": f"Table '{table_name}' is unavailable or empty."}
    if group_by and group_by not in table.columns:
        return {"used": True, "ok": False, "error": f"Group-by column '{group_by}' not found."}
    if metric not in table.columns:
        return {"used": True, "ok": False, "error": f"Metric column '{metric}' not found."}
    agg_funcs = {"sum": "sum", "mean": "mean", "min": "min", "max": "max", "median": "median"}
    agg_func = agg_funcs.get(agg, "sum")
    try:
        if group_by:
            out = table.groupby(group_by, dropna=False)[metric].agg(agg_func).reset_index()
        else:
            out = pd.DataFrame([{metric: getattr(pd.to_numeric(table[metric], errors="coerce"), agg_func)()}])
        preview = out.head(20).to_dict(orient="records")
        return {"used": True, "ok": True, "result": preview}
    except Exception as exc:
        return {"used": True, "ok": False, "error": str(exc)}


def _tool_scenario_compare(model_results: Mapping[str, object], metric_name: str, top_n: int = 10) -> Dict[str, object]:
    """Compare base metrics with lender-case metrics when available."""

    base_metrics = model_results.get("metrics", {})
    lender_df = model_results.get("lender_case_results")
    if not isinstance(lender_df, pd.DataFrame) or lender_df.empty:
        return {"used": True, "ok": False, "error": "No lender_case_results found for scenario comparison."}
    if metric_name not in lender_df.columns:
        return {"used": True, "ok": False, "error": f"Metric '{metric_name}' not found in lender_case_results."}
    try:
        comp = lender_df[["case_name", metric_name]].copy()
        base_value = base_metrics.get(metric_name) if isinstance(base_metrics, Mapping) else None
        comp["base_value"] = base_value
        comp["delta_vs_base"] = pd.to_numeric(comp[metric_name], errors="coerce") - (float(base_value) if isinstance(base_value, (int, float, np.floating)) else 0.0)
        comp = comp.head(max(1, min(int(top_n), 50)))
        return {"used": True, "ok": True, "result": comp.to_dict(orient="records")}
    except Exception as exc:
        return {"used": True, "ok": False, "error": str(exc)}


def _web_comparison_search(query: str, limit: int = 5) -> List[Dict[str, str]]:
    """Perform lightweight web lookup and evidence enrichment for comparative references."""

    if not query.strip():
        return []
    encoded = urllib.parse.urlencode({"q": query, "format": "json", "no_html": "1", "skip_disambig": "1"})
    url = f"https://api.duckduckgo.com/?{encoded}"
    try:
        with urllib.request.urlopen(url, timeout=8) as response:
            payload = json.loads(response.read().decode("utf-8"))
    except Exception:
        return []

    results: List[Dict[str, str]] = []
    for item in payload.get("RelatedTopics", []):
        if isinstance(item, dict) and item.get("Text") and item.get("FirstURL"):
            title = str(item.get("Text"))
            url = str(item.get("FirstURL"))
            results.append({"title": title, "url": url})
        if len(results) >= limit:
            break
    return results


def _extract_comparable_metrics(text: str) -> Dict[str, str]:
    """Extract rough comparable fields (date, region, value) from source text."""

    lowered = text.lower()
    year_match = re.search(r"\b(20\d{2}|19\d{2})\b", text)
    date_value = year_match.group(0) if year_match else ""
    value_match = re.search(r"(\$?\d[\d,]*(?:\.\d+)?%?)", text)
    metric_value = value_match.group(1) if value_match else ""
    region = ""
    for candidate in ("global", "us", "usa", "europe", "asia", "brazil", "india", "china", "africa", "latin america"):
        if candidate in lowered:
            region = candidate.upper() if len(candidate) <= 3 else candidate.title()
            break
    return {"date": date_value, "region": region, "value": metric_value}


def _match_whitelisted_institution(url: str) -> Tuple[str, bool]:
    lowered = url.lower()
    for domain, label in INSTITUTION_WHITELIST:
        if domain in lowered:
            return label, True
    return "Unclassified source", False


def _fetch_source_text(url: str, max_chars: int = 6000) -> str:
    """Fetch lightweight source text snippet for parsing benchmark signals."""

    try:
        request = urllib.request.Request(
            url,
            headers={"User-Agent": "Mozilla/5.0 (compatible; StreamlitBot/1.0)"},
            method="GET",
        )
        with urllib.request.urlopen(request, timeout=8) as response:
            raw = response.read(max_chars * 2).decode("utf-8", errors="ignore")
    except Exception:
        return ""
    text = re.sub(r"<script.*?>.*?</script>", " ", raw, flags=re.IGNORECASE | re.DOTALL)
    text = re.sub(r"<style.*?>.*?</style>", " ", text, flags=re.IGNORECASE | re.DOTALL)
    text = re.sub(r"<[^>]+>", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text[:max_chars]


def _normalize_benchmark_value(raw_value: str) -> Optional[float]:
    token = raw_value.replace(",", "").strip()
    if not token:
        return None
    try:
        if token.endswith("%"):
            return float(token[:-1]) / 100.0
        if token.startswith("$"):
            return float(token[1:])
        return float(token)
    except Exception:
        return None


def _parse_benchmark_definition_alignment(text: str, query: str) -> Dict[str, object]:
    """Parse metric mentions and score alignment to known benchmark definitions."""

    combined = f"{query} {text}".lower()
    best_metric = ""
    best_score = 0.0
    for metric_name, phrases in BENCHMARK_METRIC_DEFINITIONS.items():
        hits = sum(1 for phrase in phrases if phrase in combined)
        score = min(1.0, hits / max(1, len(phrases)))
        if score > best_score:
            best_score = score
            best_metric = metric_name
    return {"metric": best_metric, "definition_alignment_score": round(best_score, 2)}


def _score_credibility(url: str, title: str) -> Tuple[float, str]:
    """Return heuristic credibility score with rationale."""

    score = 0.4
    reasons: List[str] = []
    trusted_domains = ("gov", "edu", "org", "iea.org", "worldbank.org", "oecd.org", "reuters.com", "bloomberg.com")
    lowered_url = url.lower()
    lowered_title = title.lower()
    institution_label, is_whitelisted = _match_whitelisted_institution(url)
    if is_whitelisted:
        score += 0.35
        reasons.append(f"whitelisted institution: {institution_label}")
    elif any(domain in lowered_url for domain in trusted_domains):
        score += 0.2
        reasons.append("trusted-like domain pattern")
    if any(keyword in lowered_title for keyword in ("report", "index", "statistics", "official", "benchmark")):
        score += 0.15
        reasons.append("benchmark-like content")
    if re.search(r"\b20\d{2}\b", title):
        score += 0.1
        reasons.append("contains explicit year")
    score = max(0.0, min(score, 1.0))
    reason_text = ", ".join(reasons) if reasons else "generic web reference"
    return score, reason_text


def _build_evidence_pipeline(
    query: str,
    model_metrics: Mapping[str, object],
    top_n: int = 5,
) -> List[Dict[str, object]]:
    """Retrieve sources and enrich them into comparable evidence rows."""

    raw_sources = _web_comparison_search(query, limit=max(top_n * 2, top_n))
    current_year = datetime.datetime.utcnow().year
    evidence_rows: List[Dict[str, object]] = []
    metric_snapshot = {
        "Project_NPV": model_metrics.get("Project_NPV"),
        "Project_IRR": model_metrics.get("Project_IRR"),
        "DSCR_min": model_metrics.get("DSCR_min"),
    }
    whitelisted_rows: List[Dict[str, object]] = []
    non_whitelisted_rows: List[Dict[str, object]] = []
    for src in raw_sources:
        title = str(src.get("title", ""))
        url = str(src.get("url", ""))
        source_text = _fetch_source_text(url)
        parse_source = f"{title} {source_text}".strip()
        extracted = _extract_comparable_metrics(parse_source)
        normalized_value = _normalize_benchmark_value(extracted.get("value", ""))
        alignment = _parse_benchmark_definition_alignment(parse_source, query)
        score, why = _score_credibility(url, title)
        institution_label, is_whitelisted = _match_whitelisted_institution(url)
        year_value = extracted.get("date", "")
        recency = "unknown"
        if year_value.isdigit():
            delta = current_year - int(year_value)
            recency = "current (<=1y)" if delta <= 1 else f"{delta} years old"
        confidence = round(min(1.0, score * 0.6 + float(alignment.get("definition_alignment_score", 0.0)) * 0.4), 2)
        row = {
            "source_title": title,
            "source_url": url,
            "institution": institution_label,
            "whitelisted_institution": is_whitelisted,
            "extracted_date": extracted.get("date", ""),
            "extracted_region": extracted.get("region", ""),
            "benchmark_metric": alignment.get("metric", ""),
            "extracted_value": extracted.get("value", ""),
            "normalized_benchmark_value": normalized_value,
            "definition_alignment_score": alignment.get("definition_alignment_score", 0.0),
            "credibility_score": round(score, 2),
            "confidence_score": confidence,
            "why_used": why if why else "Relevant to query and benchmark extraction",
            "date_recency": recency,
            "model_Project_NPV": metric_snapshot.get("Project_NPV"),
            "model_Project_IRR": metric_snapshot.get("Project_IRR"),
            "model_DSCR_min": metric_snapshot.get("DSCR_min"),
        }
        if is_whitelisted:
            whitelisted_rows.append(row)
        else:
            non_whitelisted_rows.append(row)

    evidence_rows = (whitelisted_rows + non_whitelisted_rows)[:top_n]
    return evidence_rows


def _call_chat_provider(settings: ChatProviderSettings, messages: List[Dict[str, str]]) -> str:
    """Send chat request through provider adapters with shared interface."""

    if not settings.api_key.strip():
        return ""
    if not settings.base_url.strip():
        return ""
    try:
        adapter = _create_provider_adapter(settings)
        return adapter.generate(messages).strip()
    except Exception:
        return ""


def _extract_text_from_provider_payload(data: Mapping[str, object]) -> str:
    """Extract assistant text from common provider response payload shapes."""

    if not isinstance(data, Mapping):
        return ""
    choices = data.get("choices")
    if isinstance(choices, list) and choices:
        message = choices[0].get("message", {}) if isinstance(choices[0], Mapping) else {}
        content = message.get("content") if isinstance(message, Mapping) else None
        if isinstance(content, str) and content.strip():
            return content
        if isinstance(content, list):
            blocks: List[str] = []
            for block in content:
                if isinstance(block, Mapping):
                    text = block.get("text")
                    if isinstance(text, str) and text.strip():
                        blocks.append(text.strip())
            if blocks:
                return "\n\n".join(blocks)
    content_blocks = data.get("content")
    if isinstance(content_blocks, list):
        blocks: List[str] = []
        for block in content_blocks:
            if isinstance(block, Mapping):
                text = block.get("text")
                if isinstance(text, str) and text.strip():
                    blocks.append(text.strip())
        if blocks:
            return "\n\n".join(blocks)
    output = data.get("output_text")
    if isinstance(output, str) and output.strip():
        return output
    return ""


def _create_provider_adapter(settings: ChatProviderSettings) -> BaseChatAdapter:
    """Factory for provider adapters."""

    if settings.provider_name == "OpenAI":
        return OpenAIAdapter(settings)
    if settings.provider_name == "Anthropic":
        return AnthropicAdapter(settings)
    if settings.provider_name == "Google Gemini":
        return GeminiAdapter(settings)
    return GenericOpenAICompatibleAdapter(settings)


def _build_chatbot_fallback_reply(
    user_prompt: str,
    metrics: Mapping[str, object],
    sandbox_output: Mapping[str, object],
    web_sources: Sequence[Mapping[str, str]],
) -> str:
    """Return a prose-first fallback reply when provider calls are unavailable."""

    npv = metrics.get("Project_NPV")
    irr = metrics.get("Project_IRR")
    dscr = metrics.get("DSCR_min")
    metric_fragments: List[str] = []
    if isinstance(npv, (int, float, np.floating)):
        metric_fragments.append(f"project NPV is about {float(npv):,.2f}")
    if isinstance(irr, (int, float, np.floating)):
        metric_fragments.append(f"project IRR is roughly {float(irr) * 100:.2f}%")
    if isinstance(dscr, (int, float, np.floating)):
        metric_fragments.append(f"minimum DSCR is around {float(dscr):.2f}")
    metric_sentence = "I checked the loaded model context and " + ", ".join(metric_fragments) + "." if metric_fragments else ""

    sandbox_sentence = ""
    if sandbox_output.get("used"):
        if sandbox_output.get("ok"):
            sandbox_sentence = f"I also executed your sandbox snippet successfully and obtained: {sandbox_output.get('result')}."
        else:
            sandbox_sentence = f"I attempted the sandbox snippet, but it returned an error: {sandbox_output.get('error')}."

    web_sentence = ""
    if web_sources:
        web_sentence = (
            f"For external comparison, I found {len(web_sources)} web references that can be used to benchmark assumptions."
        )

    return (
        "Here is a practical interpretation of your request: "
        f"{user_prompt.strip()}.\n\n"
        f"{metric_sentence} {sandbox_sentence} {web_sentence}\n\n"
        "Recommendation: if you want a tighter answer, share the exact KPI you want to optimize "
        "(for example NPV, DSCR, or payback) and I will provide a targeted scenario-based response."
    ).strip()


def _build_lightweight_plan(
    user_prompt: str,
    prior_plans: Sequence[Mapping[str, object]],
    sandbox_enabled: bool,
    web_enabled: bool,
) -> Dict[str, object]:
    """Generate a lightweight per-turn plan that can be reused by follow-up turns."""

    prompt = user_prompt.strip()
    prompt_lower = prompt.lower()
    prior_objective = ""
    for item in reversed(prior_plans):
        prior_objective = str(item.get("objective", "")).strip()
        if prior_objective:
            break

    objective = prompt
    if not objective and prior_objective:
        objective = prior_objective
    if not objective:
        objective = "Provide a reasoned analytical response to the user's request."

    sandbox_needed = sandbox_enabled and any(
        token in prompt_lower
        for token in (
            "calculate",
            "calc",
            "simulate",
            "simulation",
            "python",
            "code",
            "table",
            "dataframe",
            "model this",
        )
    )
    web_needed = web_enabled and any(
        token in prompt_lower
        for token in (
            "benchmark",
            "compare",
            "market",
            "industry",
            "latest",
            "current",
            "reference",
            "best practice",
        )
    )

    return {
        "objective": objective,
        "clarification": "Objective set from current prompt; falls back to prior turn objective when needed.",
        "sandbox_needed": sandbox_needed,
        "web_needed": web_needed,
        "execution": "Run selected tools, then synthesize a concise prose answer with interpretation and recommendation.",
    }


def _update_conversation_summary(history: Sequence[Mapping[str, object]], existing_summary: str) -> str:
    """Maintain a rolling summary every five user turns."""

    user_turns = [item for item in history if str(item.get("role", "")).lower() == "user"]
    if not user_turns:
        return existing_summary
    if len(user_turns) % 5 != 0 and existing_summary.strip():
        return existing_summary

    recent = history[-10:]
    summary_lines: List[str] = []
    for item in recent:
        role = str(item.get("role", "")).strip().lower()
        content = str(item.get("content", "")).strip().replace("\n", " ")
        if not content:
            continue
        clipped = content[:180] + ("..." if len(content) > 180 else "")
        summary_lines.append(f"{role}: {clipped}")
    return " | ".join(summary_lines)


def _update_fact_memory(
    user_prompt: str,
    assistant_text: str,
    metrics: Mapping[str, object],
    existing_facts: Mapping[str, object],
) -> Dict[str, object]:
    """Update persistent fact memory with assumptions, KPIs, constraints, and decisions."""

    facts = dict(existing_facts)
    text_blob = f"{user_prompt}\n{assistant_text}".lower()
    constraints = set(str(x) for x in facts.get("constraints", []))
    assumptions = set(str(x) for x in facts.get("assumptions", []))
    decisions = list(facts.get("decisions", [])) if isinstance(facts.get("decisions"), list) else []

    if "assume" in text_blob or "assumption" in text_blob:
        assumptions.add(user_prompt.strip()[:200])
    for keyword in ("must", "limit", "cannot", "constraint", "target", "threshold"):
        if keyword in text_blob:
            constraints.add(keyword)
    for keyword in ("recommend", "should", "decision", "choose", "select"):
        if keyword in text_blob:
            decision_note = assistant_text.strip()[:220]
            if decision_note:
                decisions.append(decision_note)
            break

    kpi_keys = ["Project_NPV", "Project_IRR", "Equity_IRR", "DSCR_min", "Payback_Year"]
    kpis = {k: metrics.get(k) for k in kpi_keys if k in metrics}

    facts["assumptions"] = sorted(x for x in assumptions if x)
    facts["constraints"] = sorted(x for x in constraints if x)
    facts["decisions"] = decisions[-10:]
    facts["kpis"] = kpis
    return facts


def _reasoning_checklist(
    response_text: str,
    user_prompt: str,
    plan: Mapping[str, object],
    sandbox_output: Mapping[str, object],
    web_sources: Sequence[Mapping[str, str]],
) -> Dict[str, bool]:
    """Assess answer quality against internal reasoning checklist criteria."""

    text = response_text.lower()
    prompt = user_prompt.lower()
    assumptions = any(token in text for token in ("assumption", "assume", "assuming"))
    requested_tools = bool(plan.get("sandbox_needed")) or bool(plan.get("web_needed"))
    tool_used = bool(sandbox_output.get("used")) or bool(web_sources)
    used_requested_data_tools = (not requested_tools) or tool_used
    if any(token in prompt for token in ("calculate", "aggregate", "compare")) and not tool_used:
        used_requested_data_tools = False
    causality = any(token in text for token in ("because", "therefore", "due to", "driven by", "leads to"))
    recommendation = any(token in text for token in ("recommend", "should", "action", "next step"))
    caveat = any(token in text for token in ("caveat", "risk", "however", "uncertain", "limitation"))
    return {
        "stated_assumptions": assumptions,
        "used_requested_data_tools": used_requested_data_tools,
        "explained_causality": causality,
        "recommendation_and_caveat": recommendation and caveat,
    }


def _enforce_reasoning_checklist(
    response_text: str,
    checklist: Mapping[str, bool],
    plan: Mapping[str, object],
) -> str:
    """Append missing checklist sections so final answer is complete."""

    additions: List[str] = []
    if not checklist.get("stated_assumptions", False):
        additions.append("Assumption: This guidance assumes current model inputs and base-case KPI definitions remain unchanged.")
    if not checklist.get("used_requested_data_tools", False):
        additions.append(
            f"Tool/Data note: Requested tool usage may be incomplete for this turn. Planned tool path was sandbox={plan.get('sandbox_needed')} web={plan.get('web_needed')}."
        )
    if not checklist.get("explained_causality", False):
        additions.append("Causality: The recommendation is driven by how operating cash flow affects debt service coverage and valuation metrics.")
    if not checklist.get("recommendation_and_caveat", False):
        additions.append("Recommendation + caveat: Prioritize DSCR resilience first, but note outcomes remain sensitive to price and capex uncertainty.")
    if not additions:
        return response_text
    return response_text.strip() + "\n\n" + "\n".join(f"- {line}" for line in additions)


def _run_offline_chat_evaluation(model_results: Mapping[str, object]) -> pd.DataFrame:
    """Run offline heuristic evaluation set for chatbot quality checks."""

    rows: List[Dict[str, object]] = []
    metrics = model_results.get("metrics", {}) if isinstance(model_results.get("metrics", {}), Mapping) else {}
    for case in CHAT_EVAL_PROMPTS:
        prompt = case["prompt"]
        expected_tool = case["expected_tool"]
        start = time.perf_counter()
        synthetic_plan = {
            "sandbox_needed": expected_tool in {"tool_calc", "tool_dataframe_aggregate", "tool_scenario_compare"},
            "web_needed": expected_tool == "web",
        }
        synthetic_sandbox = {"used": synthetic_plan["sandbox_needed"], "ok": True, "result": "synthetic"}
        synthetic_web = [{"title": "Synthetic reference", "url": "https://example.com"}] if synthetic_plan["web_needed"] else []
        response = _build_chatbot_fallback_reply(prompt, metrics, synthetic_sandbox, synthetic_web)
        checklist = _reasoning_checklist(response, prompt, synthetic_plan, synthetic_sandbox, synthetic_web)
        elapsed_ms = (time.perf_counter() - start) * 1000.0
        token_estimate = max(1, int(len(response) / 4))
        rows.append(
            {
                "prompt": prompt,
                "expected_tool": expected_tool,
                "factual_grounding": float(checklist["stated_assumptions"]),
                "coherence": 1.0 if len(response.split()) >= 25 else 0.5,
                "tool_use_correctness": 1.0 if ((expected_tool == "none") or checklist["used_requested_data_tools"]) else 0.0,
                "latency_ms": round(elapsed_ms, 2),
                "cost_tokens_estimate": token_estimate,
            }
        )
    return pd.DataFrame(rows)


def _render_chatbot_tab(model_results: Mapping[str, object]) -> None:
    st.subheader("Intelligent Analytical Chatbot")
    st.caption(
        "Reasoning-first assistant with conversation memory, optional sandbox execution, "
        "selective web comparison, and multi-provider LLM settings."
    )

    history = st.session_state.setdefault("chat_history", [])
    plans = st.session_state.setdefault("chat_plans", [])
    conversation_summary = st.session_state.setdefault("chat_summary", "")
    fact_memory = st.session_state.setdefault(
        "chat_fact_memory",
        {"assumptions": [], "constraints": [], "decisions": [], "kpis": {}},
    )
    provider_options = list(CHAT_PROVIDER_DEFAULTS.keys())
    selected_provider = st.selectbox("Provider", provider_options, key="chat_provider")
    defaults = CHAT_PROVIDER_DEFAULTS[selected_provider]

    col1, col2, col3 = st.columns(3)
    model_name = col1.text_input("Model name", value=st.session_state.get("chat_model_name", "gpt-4.1-mini"), key="chat_model_name")
    api_key = col2.text_input("API key", value=st.session_state.get("chat_api_key", ""), type="password", key="chat_api_key")
    base_url = col3.text_input("Base URL", value=st.session_state.get("chat_base_url", str(defaults.get("base_url", ""))), key="chat_base_url")
    col4, col5, col6 = st.columns(3)
    temperature = col4.number_input("Temperature", min_value=0.0, max_value=2.0, value=float(st.session_state.get("chat_temperature", 0.2)), step=0.1, key="chat_temperature")
    max_tokens = int(col5.number_input("Max tokens", min_value=64, max_value=8192, value=int(st.session_state.get("chat_max_tokens", 1200)), step=64, key="chat_max_tokens"))
    reasoning_mode = col6.selectbox("Reasoning mode", ["", "low", "medium", "high"], index=2, key="chat_reasoning_mode")

    st.write(
        f"Capabilities: reasoning={_chat_supports(selected_provider, 'supports_reasoning')}, "
        f"tools={_chat_supports(selected_provider, 'supports_tools')}, "
        f"web={_chat_supports(selected_provider, 'supports_web_search')}"
    )

    with st.expander("Conversation history", expanded=False):
        if history:
            for item in history:
                st.markdown(f"**{item.get('role', 'assistant').title()}:** {item.get('content', '')}")
        else:
            st.info("No messages yet.")
        if st.button("Clear history", key="clear_chat_history"):
            st.session_state["chat_history"] = []
            st.session_state["chat_plans"] = []
            st.session_state["chat_summary"] = ""
            st.session_state["chat_fact_memory"] = {"assumptions": [], "constraints": [], "decisions": [], "kpis": {}}
            _safe_rerun()

    with st.expander("Planner memory", expanded=False):
        if plans:
            for idx, plan in enumerate(plans[-5:], start=max(1, len(plans) - 4)):
                st.markdown(
                    f"**Turn {idx} objective:** {plan.get('objective', '')}\n\n"
                    f"- Sandbox needed: {plan.get('sandbox_needed')}\n"
                    f"- Web needed: {plan.get('web_needed')}\n"
                    f"- Execution: {plan.get('execution', '')}"
                )
        else:
            st.caption("No saved plans yet.")
    with st.expander("Conversation summary & fact memory", expanded=False):
        st.write(f"Summary: {conversation_summary or 'No summary yet.'}")
        st.markdown("**Fact memory**")
        st.write(f"Assumptions: {', '.join(map(str, fact_memory.get('assumptions', []))) or 'None'}")
        st.write(f"Constraints: {', '.join(map(str, fact_memory.get('constraints', []))) or 'None'}")
        st.write(f"Recent decisions: {' | '.join(map(str, fact_memory.get('decisions', []))) or 'None'}")
        st.write(f"Tracked KPIs: {fact_memory.get('kpis', {})}")

    use_sandbox = st.checkbox("Use sandbox execution for intermediate calculations", value=True, key="chat_use_sandbox")
    use_web = st.checkbox("Use web search for comparative analysis", value=True, key="chat_use_web")
    st.markdown("#### Structured tools")
    tool_choice = st.selectbox(
        "Select tool",
        ["none", "tool_calc(expression)", "tool_dataframe_aggregate(...)", "tool_scenario_compare(...)"],
        key="chat_tool_choice",
    )
    calc_expression = st.text_input(
        "tool_calc expression",
        value=st.session_state.get("chat_calc_expression", ""),
        key="chat_calc_expression",
    )
    aggregate_table = st.selectbox(
        "tool_dataframe_aggregate table",
        ["revenue", "production_monthly", "price_curves", "cash_waterfall", "credit_metrics_yearly", "lender_case_results"],
        key="chat_aggregate_table",
    )
    aggregate_group = st.text_input("Group by column (optional)", value=st.session_state.get("chat_aggregate_group", ""), key="chat_aggregate_group")
    aggregate_metric = st.text_input("Metric column", value=st.session_state.get("chat_aggregate_metric", "revenue"), key="chat_aggregate_metric")
    aggregate_fn = st.selectbox("Aggregation", ["sum", "mean", "min", "max", "median"], key="chat_aggregate_fn")
    scenario_metric = st.text_input("tool_scenario_compare metric", value=st.session_state.get("chat_scenario_metric", "Project_NPV"), key="chat_scenario_metric")
    scenario_top_n = int(st.number_input("Top N scenario rows", min_value=1, max_value=50, value=10, step=1, key="chat_scenario_top_n"))

    advanced_mode = st.checkbox("Advanced mode: run raw Python sandbox code", value=False, key="chat_advanced_mode")
    sandbox_code = ""
    if advanced_mode:
        sandbox_code = st.text_area(
            "Sandbox code (advanced). Set `result = ...` to expose output.",
            value=st.session_state.get("chat_sandbox_code", ""),
            height=120,
            key="chat_sandbox_code",
        )
    user_prompt = st.text_area("Ask the assistant", height=140, key="chat_prompt")

    if st.button("Send", key="chat_send"):
        if not user_prompt.strip():
            st.warning("Enter a prompt to continue.")
            return

        settings = ChatProviderSettings(
            provider_name=selected_provider,
            model_name=model_name.strip() or "gpt-4.1-mini",
            api_key=api_key,
            base_url=base_url.strip(),
            temperature=float(temperature),
            max_tokens=max_tokens,
            reasoning_mode=reasoning_mode,
            use_tools=bool(use_sandbox),
            use_web_search=bool(use_web),
        )
        plan = _build_lightweight_plan(
            user_prompt=user_prompt,
            prior_plans=plans,
            sandbox_enabled=settings.use_tools,
            web_enabled=settings.use_web_search and _chat_supports(selected_provider, "supports_web_search"),
        )
        if tool_choice != "none":
            plan["sandbox_needed"] = True
        plans.append(plan)
        st.session_state["chat_plans"] = plans

        sandbox_output: Dict[str, object] = {"used": False}
        if bool(plan.get("sandbox_needed")):
            if tool_choice == "tool_calc(expression)":
                sandbox_output = _tool_calc(calc_expression)
            elif tool_choice == "tool_dataframe_aggregate(...)":
                sandbox_output = _tool_dataframe_aggregate(
                    model_results=model_results,
                    table_name=aggregate_table,
                    group_by=aggregate_group.strip(),
                    metric=aggregate_metric.strip(),
                    agg=aggregate_fn,
                )
            elif tool_choice == "tool_scenario_compare(...)":
                sandbox_output = _tool_scenario_compare(
                    model_results=model_results,
                    metric_name=scenario_metric.strip(),
                    top_n=scenario_top_n,
                )
            elif advanced_mode and sandbox_code.strip():
                sandbox_output = _run_sandbox(sandbox_code)

        web_sources: List[Dict[str, str]] = []
        evidence_rows: List[Dict[str, object]] = []
        if bool(plan.get("web_needed")):
            evidence_rows = _build_evidence_pipeline(
                query=user_prompt,
                model_metrics=model_results.get("metrics", {}) if isinstance(model_results.get("metrics", {}), Mapping) else {},
                top_n=5,
            )
            web_sources = [{"title": str(row.get("source_title", "")), "url": str(row.get("source_url", ""))} for row in evidence_rows]

        compact_metrics = model_results.get("metrics", {}) if isinstance(model_results.get("metrics", {}), Mapping) else {}
        system_prompt = (
            "You are an intelligent analytical assistant. Respond with sections: "
            "Direct answer, Internal reasoning summary, Sandbox output usage, External comparison, "
            "Interpretation, Recommendation, Sources. Keep answers concise and actionable."
        )
        metric_lines = ", ".join(
            f"{key}={value}"
            for key, value in compact_metrics.items()
            if key in {"Project_NPV", "Project_IRR", "Equity_IRR", "DSCR_min", "Payback_Year"}
        )
        web_lines = "; ".join(f"{src.get('title', '')} ({src.get('url', '')})" for src in web_sources)
        sandbox_line = (
            f"used={sandbox_output.get('used')}, ok={sandbox_output.get('ok')}, result={sandbox_output.get('result')}, "
            f"error={sandbox_output.get('error')}"
        )
        context_blob = (
            f"Model metrics: {metric_lines or 'none'}\n"
            f"Sandbox summary: {sandbox_line}\n"
            f"Web references: {web_lines or 'none'}\n"
            f"Evidence summary rows: {evidence_rows[:3] if evidence_rows else 'none'}"
        )
        llm_messages: List[Dict[str, str]] = [{"role": "system", "content": system_prompt}]
        memory_prefix = (
            f"Conversation summary: {conversation_summary or 'none'}\n"
            f"Fact memory: assumptions={fact_memory.get('assumptions', [])}, "
            f"constraints={fact_memory.get('constraints', [])}, "
            f"decisions={fact_memory.get('decisions', [])}, "
            f"kpis={fact_memory.get('kpis', {})}"
        )
        llm_messages.append({"role": "system", "content": memory_prefix})
        llm_messages.extend([{"role": str(m.get("role", "user")), "content": str(m.get("content", ""))} for m in history[-8:]])
        llm_messages.append(
            {
                "role": "user",
                "content": f"User prompt:\n{user_prompt}\n\nContext:\n{context_blob}",
            }
        )
        assistant_text = _call_chat_provider(settings, llm_messages)
        if not assistant_text.strip():
            assistant_text = _build_chatbot_fallback_reply(
                user_prompt=user_prompt,
                metrics=compact_metrics,
                sandbox_output=sandbox_output,
                web_sources=web_sources,
            )
        checklist = _reasoning_checklist(
            response_text=assistant_text,
            user_prompt=user_prompt,
            plan=plan,
            sandbox_output=sandbox_output,
            web_sources=web_sources,
        )
        assistant_text = _enforce_reasoning_checklist(assistant_text, checklist, plan)

        history.append({"role": "user", "content": user_prompt})
        history.append({"role": "assistant", "content": assistant_text})
        st.session_state["chat_history"] = history
        st.session_state["chat_summary"] = _update_conversation_summary(history, conversation_summary)
        st.session_state["chat_fact_memory"] = _update_fact_memory(
            user_prompt=user_prompt,
            assistant_text=assistant_text,
            metrics=compact_metrics,
            existing_facts=fact_memory,
        )

        st.markdown("### Direct answer")
        st.write(assistant_text)
        st.markdown("### Planner step")
        st.write(
            f"Objective: {plan.get('objective')}\n\n"
            f"Clarification: {plan.get('clarification')}\n\n"
            f"Sandbox needed: {plan.get('sandbox_needed')} | Web needed: {plan.get('web_needed')}\n\n"
            f"Execution plan: {plan.get('execution')}"
        )
        st.markdown("### Internal reasoning checklist")
        st.write(
            f"Stated assumptions: {checklist.get('stated_assumptions')}\n\n"
            f"Used requested data/tools: {checklist.get('used_requested_data_tools')}\n\n"
            f"Explained causality: {checklist.get('explained_causality')}\n\n"
            f"Recommendation + caveat: {checklist.get('recommendation_and_caveat')}"
        )
        st.markdown("### Internal reasoning")
        st.info("Reasoning summary is embedded in the assistant response.")
        st.markdown("### Sandbox output")
        if sandbox_output.get("used"):
            if sandbox_output.get("ok"):
                st.write(
                    f"The sandbox ran successfully. Result: {sandbox_output.get('result')}. "
                    f"Variables produced: {', '.join(map(str, sandbox_output.get('locals', [])))}."
                )
            else:
                st.write(f"The sandbox run failed with error: {sandbox_output.get('error')}.")
        else:
            st.caption("Sandbox execution was not used for this turn.")
        st.markdown("### External comparison via web search")
        if evidence_rows:
            evidence_df = pd.DataFrame(evidence_rows)
            st.dataframe(evidence_df, use_container_width=True)
        else:
            st.caption("No web comparison data was collected.")
        st.markdown("### Interpretation")
        st.caption("Use the response to validate assumptions against model metrics and external references.")
        st.markdown("### Recommendation")
        st.caption("Iterate with tighter prompts, explicit assumptions, and optional sandbox scripts for reproducible steps.")
        st.markdown("### Sources")
        for src in web_sources:
            st.markdown(f"- [{src['title']}]({src['url']})")

    with st.expander("Offline evaluation set (quality diagnostics)", expanded=False):
        st.caption("Runs a 20-prompt offline heuristic evaluation across grounding, coherence, tool-use correctness, latency, and token-cost estimate.")
        if st.button("Run offline evaluation", key="chat_run_eval"):
            eval_df = _run_offline_chat_evaluation(model_results)
            st.dataframe(eval_df, use_container_width=True)
            if not eval_df.empty:
                summary = pd.DataFrame(
                    [
                        {
                            "avg_factual_grounding": float(eval_df["factual_grounding"].mean()),
                            "avg_coherence": float(eval_df["coherence"].mean()),
                            "avg_tool_use_correctness": float(eval_df["tool_use_correctness"].mean()),
                            "avg_latency_ms": float(eval_df["latency_ms"].mean()),
                            "avg_cost_tokens_estimate": float(eval_df["cost_tokens_estimate"].mean()),
                        }
                    ]
                )
                st.dataframe(summary, use_container_width=True)


@contextmanager
def _modal_container(title: str):
    """Yield a modal-like container, falling back when `st.modal` is unavailable."""

    if hasattr(st, "modal"):
        with st.modal(title):
            yield True
        return

    # Fallback for older Streamlit releases – use an expander to host the form.
    with st.expander(title, expanded=True):
        st.info(
            "Modal dialogs are not supported in this Streamlit version; "
            "editing is displayed inline instead.",
            icon="ℹ️",
        )
        yield False


def _format_default_for_entry(value: object, dtype: str) -> object:
    """Return a widget-friendly default for modal data-entry fields."""

    if dtype == "bool":
        if value is None or (isinstance(value, float) and math.isnan(value)):
            return False
        return bool(value)
    if value is None:
        return ""
    if isinstance(value, (float, np.floating)) and math.isnan(value):
        return ""
    if isinstance(value, pd.Timestamp):
        return value.strftime("%Y-%m-%d")
    return str(value)


def _parse_modal_entry(value: object, dtype: str) -> object:
    """Convert modal widget values back into schema-compatible data."""

    if dtype == "bool":
        return bool(value)
    text = str(value).strip() if value is not None else ""
    if text == "":
        return np.nan
    try:
        if dtype == "int":
            return int(float(text))
        if dtype == "float":
            return float(text)
    except ValueError:
        raise ValueError(f"'{text}' is not a valid {dtype} value")
    return text


MODEL_IMPORT_ERROR: ModuleNotFoundError | None = None
try:  # noqa: SIM105 - streamlit feedback when dependencies missing
    from model import (
        DEFAULTS,
        FEEDSTOCK_SCENARIOS,
        INPUT_SCHEMAS,
        PRODUCTS,
        InputTables,
        align_with_projection_horizon,
        build_config,
        MONTE_CARLO_DISTRIBUTIONS,
        MONTE_CARLO_VARIABLES,
        MONTE_CARLO_VARIABLE_LABELS,
        compute_risk_profile,
        decision_tree_analysis,
        metaheuristic_optimize,
        monte_carlo,
        neural_forecast_production,
    parse_ramp,
    run_full_model,
    run_scenarios,
    sensitivity_tornado,
    statistical_forecast,
    normalize_key,
    Timeline,
    resolve_excel_engine,
    SIMPLE_XLSX_ENGINE,
    write_simple_xlsx,
)
except ModuleNotFoundError as exc:  # pragma: no cover - executed only when deps missing
    MODEL_IMPORT_ERROR = exc

if MODEL_IMPORT_ERROR is not None:
    _KEY_NORMALIZER = re.compile(r"[^0-9a-zA-Z]+")

    def normalize_key(value: str) -> str:
        key = _KEY_NORMALIZER.sub("_", str(value).strip().lower())
        key = re.sub(r"_+", "_", key).strip("_")
        return key

    PRODUCTS = ("ethanol", "sugar", "electricity", "animal_feed")

    MONTE_CARLO_DISTRIBUTIONS = ("normal", "lognormal", "triangular", "uniform")
    MONTE_CARLO_VARIABLE_ITEMS = (
        ("opex", "Operating expenditure (all)"),
        ("interest_rate", "Interest rate"),
        ("capex", "Total CAPEX"),
        ("initial_investment", "Initial Investment (CAPEX)"),
        ("debt_schedule", "Debt schedule"),
        ("production", "Production volumes (all products)"),
        ("production_ethanol", "Production volumes (annual) – Ethanol"),
        ("production_sugar", "Production volumes (annual) – Sugar"),
        ("production_electricity", "Production volumes (annual) – Electricity"),
        ("production_animal_feed", "Production volumes (annual) – Animal feed"),
        ("pricing", "Pricing (all products)"),
        ("pricing_ethanol", "Pricing – Ethanol"),
        ("pricing_sugar", "Pricing – Sugar"),
        ("pricing_electricity", "Pricing – Electricity"),
        ("pricing_animal_feed", "Pricing – Animal feed"),
        ("revenue", "Revenue"),
        ("sugarcane_yield", "Sugarcane yield"),
        ("operating_cost_direct", "Operating Costs - Direct"),
        ("operating_cost_staff", "Operating Costs - Staff"),
        ("operating_cost_other", "Operating Costs - Other Opex"),
        ("labour", "Labour costs"),
        ("availability", "Plant availability"),
        ("other", "Other"),
    )
    MONTE_CARLO_VARIABLES = tuple(key for key, _ in MONTE_CARLO_VARIABLE_ITEMS)
    MONTE_CARLO_VARIABLE_LABELS = {key: label for key, label in MONTE_CARLO_VARIABLE_ITEMS}

    SIMPLE_XLSX_ENGINE = "__simple_xlsx__"

    def resolve_excel_engine(preferred: Sequence[str] = ("xlsxwriter", "openpyxl")) -> str:
        for engine in preferred:
            module = "openpyxl" if engine == "openpyxl" else engine
            try:
                importlib.import_module(module)
                return engine
            except ImportError:
                continue
        return SIMPLE_XLSX_ENGINE

    def write_simple_xlsx(sheets: Mapping[str, pd.DataFrame], handle) -> None:
        raise RuntimeError(
            "Excel export requires the 'xlsxwriter' or 'openpyxl' package. "
            "Install one of them to enable workbook downloads."
        )

    def compute_risk_profile(risk_params):  # pragma: no cover - fallback stub
        return {}

    def metaheuristic_optimize(*args, **kwargs):  # pragma: no cover - fallback stub
        return pd.DataFrame()

    def neural_forecast_production(*args, **kwargs):  # pragma: no cover - fallback stub
        return pd.DataFrame(columns=["date", "product", "forecast_volume"])

    def statistical_forecast(*args, **kwargs):  # pragma: no cover - fallback stub
        return {"historical": pd.DataFrame(), "forecast": pd.DataFrame(), "residual_std": np.nan, "equipment_failure_risk": np.nan}

    def decision_tree_analysis(*args, **kwargs):  # pragma: no cover - fallback stub
        return {"paths": pd.DataFrame(), "expected_metric": np.nan, "objective": "Project_NPV", "total_probability": 0.0}


if MODEL_IMPORT_ERROR is None:
    RISK_DISTRIBUTION_OPTIONS: Tuple[str, ...] = tuple(dict.fromkeys(MONTE_CARLO_DISTRIBUTIONS))
    _default_targets = set(str(x).lower() for x in DEFAULTS["risk_params"].get("target", []) if pd.notna(x))
    _default_targets.update(["risk", "price", "availability", "opex", "capex"])
    RISK_TARGET_OPTIONS: Tuple[str, ...] = tuple(sorted(_default_targets))
    _default_applies = {"global", "market"}
    for col in ("applies_to",):
        if col in DEFAULTS["risk_params"]:
            _default_applies.update(
                str(x).lower() for x in DEFAULTS["risk_params"][col].dropna().unique()
            )
    _default_applies.update(PRODUCTS)
    RISK_APPLIES_OPTIONS: Tuple[str, ...] = tuple(sorted(_default_applies))
else:  # pragma: no cover - fallback values when engine unavailable
    RISK_DISTRIBUTION_OPTIONS = MONTE_CARLO_DISTRIBUTIONS
    RISK_TARGET_OPTIONS = ("risk", "price", "availability", "opex", "capex")
    RISK_APPLIES_OPTIONS = tuple(sorted({"global", "market", *PRODUCTS}))

MONTE_CARLO_VARIABLE_LABEL_TO_KEY: Dict[str, str] = {
    label: key for key, label in MONTE_CARLO_VARIABLE_LABELS.items()
}

PRODUCT_DISPLAY_NAMES: Dict[str, str] = {
    "ethanol": "Bioethanol",
    "sugar": "Sugar",
    "electricity": "Electricity",
    "animal_feed": "Animal feed",
}


OPTIMIZER_VARIABLE_OPTIONS: Dict[str, str] = {
    "ethanol_price": "Ethanol price multiplier",
    "sugar_price": "Sugar price multiplier",
    "electricity_price": "Electricity tariff multiplier",
    "animal_feed_price": "Animal feed price multiplier",
    "availability": "Plant availability multiplier",
    "capex": "Total CAPEX multiplier",
    "opex": "Operating cost multiplier",
    "debt_rate": "Debt rate shift",
}

STATISTICAL_SERIES_OPTIONS: Dict[str, str] = {
    "revenue": "Revenue",
    "cogs": "Cost of goods sold",
    "opex": "Operating expenditure",
    "ebitda": "EBITDA",
    "staff_costs": "Staff costs",
    "production_ethanol": "Production – Ethanol",
    "production_sugar": "Production – Sugar",
    "production_electricity": "Production – Electricity",
    "production_animal_feed": "Production – Animal feed",
}


def _format_metric(value: object, kind: str = "number") -> str:
    """Format metric values for display in metric cards."""
    if value is None:
        return "N/A"
    if isinstance(value, (int, float)):
        if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
            return "N/A"
        if kind == "currency":
            return f"${value:,.2f}"
        if kind == "percent":
            return f"{value:.2%}"
        if kind == "ratio":
            return f"{value:,.2f}"
        if kind == "year":
            return str(int(round(value)))
        return f"{value:,.2f}"
    return str(value)


def _metric_kind(metric_name: str) -> str:
    key = str(metric_name or "").lower()
    if "irr" in key:
        return "percent"
    if "npv" in key or "value" in key:
        return "currency"
    if "dscr" in key or "ratio" in key:
        return "ratio"
    return "number"


DEFAULT_TABLE_STORE_KEY = "table_defaults_store"
DEFAULT_EDIT_STATE_KEY = "default_edit_state"
_FACTORY_DEFAULT_FRAMES_CACHE: Optional[Dict[str, pd.DataFrame]] = None


def _render_dataframe(df: pd.DataFrame, title: str, key: str) -> None:
    """Render a dataframe without exposing download controls."""
    st.subheader(title)
    st.dataframe(df, use_container_width=True, key=f"df_{key}")


def _ensure_statement_columns(statement_key: str, df: pd.DataFrame) -> pd.DataFrame:
    """Guarantee key financial columns exist so statement views stay complete."""

    required: Dict[str, List[str]] = {
        "pnl": ["DirectCosts", "StaffCosts", "OtherOpexCosts", "Interest"],
        "cashflow": ["CFO", "CFI", "CFF"],
        "balancesheet": ["Inventory", "PPE_Gross", "Debt", "AccountsPayable", "TotalLiabilities"],
    }
    if not isinstance(df, pd.DataFrame):
        return pd.DataFrame()
    columns = required.get(statement_key, [])
    if not columns:
        return df
    out = df.copy()
    for col in columns:
        if col not in out.columns:
            out[col] = 0.0
    return out


def _ensure_matplotlib() -> bool:
    """Return True when matplotlib is available, showing guidance otherwise."""

    global _MATPLOTLIB_WARNING_SHOWN, MATPLOTLIB_INSTALL_ERROR, plt, mdates

    if plt is None:
        if not _MATPLOTLIB_WARNING_SHOWN:
            message = (
                "Matplotlib is required to display charts. Install it with "
                "`pip install matplotlib` and reload the app."
            )
            if MATPLOTLIB_INSTALL_ERROR:
                message += f"\nLast installation attempt: {MATPLOTLIB_INSTALL_ERROR}"
            st.info(message)
            _MATPLOTLIB_WARNING_SHOWN = True
        return False
    return True


def _finalize_plot(fig) -> None:
    """Render and close a matplotlib figure inside Streamlit."""

    if fig is None:
        return
    st.pyplot(fig, use_container_width=True)
    plt.close(fig)


def _friendly_label(label: str) -> str:
    return label.replace("_", " ").title()


def _render_horizon_timeline_chart(
    horizon: Mapping[str, object],
) -> None:
    try:
        proj_start = pd.Timestamp(
            year=int(horizon.get("start_year", 0) or 0),
            month=int(horizon.get("start_month", 1) or 1),
            day=1,
        )
        proj_end = pd.Timestamp(year=int(horizon.get("end_year", 0) or 0), month=12, day=31)
    except Exception:
        st.info("Projection horizon is incomplete; add start and end years to view the timeline.")
        return

    prod_start = proj_start
    prod_end = proj_end

    bars = [
        ("Projection horizon", proj_start, proj_end, "#1f77b4"),
        ("Production horizon", prod_start, prod_end, "#2ca02c"),
    ]

    summary = pd.DataFrame(
        {
            "Horizon": [label for label, *_ in bars],
            "Start": [start.date() for _, start, _, _ in bars],
            "End": [end.date() for _, _, end, _ in bars],
            "Duration (years)": [max((end - start).days / 365.0, 0.0) for _, start, end, _ in bars],
        }
    )

    if not _ensure_matplotlib():
        st.dataframe(summary, hide_index=True, use_container_width=True)
        return

    fig, ax = plt.subplots(figsize=(8, 2.6))
    yticks: List[int] = []
    ylabels: List[str] = []

    for idx, (label, start, end, color) in enumerate(bars):
        if end < start:
            continue
        width = max((end - start).days, 1)
        left = mdates.date2num(start) if mdates else start.toordinal()
        ax.barh(idx, width, left=left, color=color, alpha=0.75)
        ax.text(left, idx, label, va="center", ha="left", fontsize=9, color="white", weight="bold")
        yticks.append(idx)
        ylabels.append(label)

    if not yticks:
        plt.close(fig)
        st.info("Unable to chart the horizons because start/end dates are missing.")
        return

    ax.set_yticks(yticks)
    ax.set_yticklabels([])
    ax.set_xlabel("Year")
    if mdates:
        ax.xaxis_date()
        ax.xaxis.set_major_formatter(mdates.DateFormatter("%Y"))
    ax.grid(axis="x", linestyle="--", alpha=0.3)
    ax.set_title("Projection vs production horizon")
    _finalize_plot(fig)


def _render_stacked_columns(
    df: pd.DataFrame,
    x_col: str,
    value_cols: Sequence[str],
    title: str,
    ylabel: str = "Value",
) -> None:
    if df.empty:
        st.info(f"No data available for {title.lower()}.")
        return

    data = df.copy()
    data = data[[x_col, *value_cols]].fillna(0.0)

    if plt is None or not _ensure_matplotlib():
        chart_data = data.set_index(x_col)[list(value_cols)]
        st.bar_chart(chart_data, use_container_width=True)
        return

    x_labels = data[x_col].astype(str).tolist()
    bottom = np.zeros(len(data))

    fig, ax = plt.subplots(figsize=(8, 4))
    for col in value_cols:
        values = data[col].astype(float).to_numpy()
        ax.bar(x_labels, values, bottom=bottom, label=_friendly_label(col))
        bottom += values

    ax.set_title(title)
    ax.set_ylabel(ylabel)
    ax.set_xlabel(_friendly_label(x_col))
    ax.legend()
    ax.set_xticklabels(x_labels, rotation=45, ha="right")
    ax.grid(axis="y", linestyle="--", alpha=0.3)
    _finalize_plot(fig)


def _render_line_chart(
    df: pd.DataFrame,
    x_col: str,
    value_cols: Sequence[str],
    title: str,
    ylabel: str = "Value",
    secondary: Optional[str] = None,
) -> None:
    data = df.copy()
    data[x_col] = pd.to_datetime(data[x_col], errors="coerce")
    data = data.dropna(subset=[x_col])
    data = data.sort_values(x_col)
    if data.empty:
        st.info(f"No valid dates available for {title.lower()}.")
        return

    if plt is None or not _ensure_matplotlib():
        chart_df = data.set_index(x_col)[list(value_cols)].astype(float)
        st.line_chart(chart_df, use_container_width=True)
        if secondary and secondary in data.columns:
            st.line_chart(data.set_index(x_col)[[secondary]], use_container_width=True)
        return

    fig, ax = plt.subplots(figsize=(8, 4))
    for col in value_cols:
        ax.plot(data[x_col], data[col], label=_friendly_label(col))

    ax.set_title(title)
    ax.set_ylabel(ylabel)
    ax.set_xlabel("Date")
    if mdates:
        ax.xaxis.set_major_formatter(mdates.DateFormatter("%Y-%m"))
    ax.grid(True, linestyle="--", alpha=0.3)

    if secondary and secondary in data.columns:
        ax2 = ax.twinx()
        ax2.plot(data[x_col], data[secondary], color="black", linestyle="--", label=_friendly_label(secondary))
        ax2.set_ylabel(_friendly_label(secondary))
        lines, labels = ax.get_legend_handles_labels()
        lines2, labels2 = ax2.get_legend_handles_labels()
        ax.legend(lines + lines2, labels + labels2, loc="upper left")
    else:
        ax.legend(loc="upper left")

    _finalize_plot(fig)


def _render_area_chart(
    df: pd.DataFrame,
    x_col: str,
    value_cols: Sequence[str],
    title: str,
    ylabel: str = "Value",
    invert_columns: Optional[Sequence[str]] = None,

) -> None:
    if df.empty:
        st.info(f"No data available for {title.lower()}.")
        return

    invert_columns = tuple(invert_columns or [])
    data = df.copy()
    data[x_col] = pd.to_datetime(data[x_col], errors="coerce")
    data = data.dropna(subset=[x_col])
    data = data.sort_values(x_col)
    if data.empty:
        st.info(f"No valid dates available for {title.lower()}.")
        return

    if plt is None or not _ensure_matplotlib():
        chart_df = data.set_index(x_col)[list(value_cols)].astype(float)
        for col in invert_columns:
            if col in chart_df.columns:
                chart_df[col] = -chart_df[col]
        st.area_chart(chart_df, use_container_width=True)
        return

    stack_values = []
    labels: List[str] = []
    for col in value_cols:
        series = data[col].astype(float).fillna(0.0)
        if col in invert_columns:
            series = -series
        stack_values.append(series.to_numpy())
        labels.append(_friendly_label(col))

    fig, ax = plt.subplots(figsize=(8, 4))
    ax.stackplot(data[x_col], *stack_values, labels=labels, alpha=0.85)
    ax.set_title(title)
    ax.set_ylabel(ylabel)
    ax.set_xlabel("Date")
    if mdates:
        ax.xaxis.set_major_formatter(mdates.DateFormatter("%Y-%m"))
    ax.legend(loc="upper left")
    ax.grid(True, linestyle="--", alpha=0.3)
    _finalize_plot(fig)


def _render_grouped_bar_chart(
    df: pd.DataFrame,
    category_col: str,
    value_map: Mapping[str, str],
    title: str,
    ylabel: str = "Value",
) -> None:
    if df.empty:
        st.info(f"No data available for {title.lower()}.")
        return

    categories = df[category_col].astype(str).tolist()
    metrics = list(value_map.keys())
    values = [df[col].astype(float).fillna(0.0).to_numpy() for col in metrics if col in df.columns]
    if not values:
        st.info(f"The required fields are missing for {title.lower()}.")
        return

    if plt is None or not _ensure_matplotlib():
        chart_df = df.set_index(category_col)[metrics].fillna(0.0)
        st.bar_chart(chart_df, use_container_width=True)
        return

    x = np.arange(len(categories))
    width = 0.8 / max(len(values), 1)

    fig, ax = plt.subplots(figsize=(8, 4))
    for idx, col in enumerate(metrics):
        if col not in df.columns:
            continue
        ax.bar(x + idx * width, df[col].astype(float).fillna(0.0), width, label=value_map[col])

    ax.set_title(title)
    ax.set_ylabel(ylabel)
    ax.set_xticks(x + width * (len(values) - 1) / 2)
    ax.set_xticklabels(categories, rotation=45, ha="right")
    ax.legend()
    ax.grid(axis="y", linestyle="--", alpha=0.3)
    _finalize_plot(fig)


def _render_scatter_chart(
    df: pd.DataFrame,
    x_col: str,
    y_col: str,
    category_col: str,
    title: str,
    xlabel: str,
    ylabel: str,
) -> None:
    data = df.dropna(subset=[x_col, y_col]).copy()
    if data.empty:
        st.info(f"Insufficient data to render {title.lower()}.")
        return

    if plt is None or not _ensure_matplotlib():
        chart_df = data[[x_col, y_col]].astype(float)
        chart_df[category_col] = data[category_col].astype(str)
        st.scatter_chart(chart_df, x=x_col, y=y_col, color=category_col, use_container_width=True)
        return

    categories = data[category_col].astype(str).unique()
    colors = plt.cm.tab10(np.linspace(0, 1, len(categories))) if plt else []

    fig, ax = plt.subplots(figsize=(8, 4))
    for color, category in zip(colors, categories):
        mask = data[category_col].astype(str) == category
        ax.scatter(data.loc[mask, x_col], data.loc[mask, y_col], color=color, label=category.title())

    ax.set_title(title)
    ax.set_xlabel(xlabel)
    ax.set_ylabel(ylabel)
    ax.grid(True, linestyle="--", alpha=0.3)
    ax.legend()
    _finalize_plot(fig)


def _render_working_capital_chart(df: pd.DataFrame) -> None:
    columns = [col for col in ("accounts_receivable", "inventory", "accounts_payable") if col in df.columns]
    if not columns:
        st.info("Working capital components are not available.")
        return
    _render_area_chart(
        df,
        "date",
        columns,
        "Working capital components",
        ylabel="Amount",
        invert_columns=["accounts_payable"],
    )


def _render_dscr_trend_chart(df: pd.DataFrame) -> None:
    if "DSCR" not in df.columns or "debt_service" not in df.columns:
        st.info("DSCR information is unavailable for charting.")
        return
    _render_line_chart(
        df,
        "date",
        ["DSCR"],
        "DSCR vs debt service",
        ylabel="DSCR",
        secondary="debt_service",
    )


def _render_debt_waterfall_chart(df: pd.DataFrame) -> None:
    required = [col for col in ("interest", "principal") if col in df.columns]
    if not required:
        st.info("Debt service summary is not available.")
        return
    _render_stacked_columns(df, "year", required, "Annual debt service breakdown", ylabel="Amount")


def _render_cost_structure_chart(df: pd.DataFrame) -> None:
    value_cols = [col for col in ("DirectCosts", "StaffCosts", "OtherOpexCosts") if col in df.columns]
    if not value_cols:
        st.info("Cost structure data is unavailable.")
        return
    _render_stacked_columns(df, "year", value_cols, "Annual operating cost structure", ylabel="Amount")


def _render_labour_area_chart(df: pd.DataFrame) -> None:
    value_cols = [col for col in ("gross_pay", "benefits", "training", "other") if col in df.columns]
    if not value_cols:
        st.info("Labour cost components are unavailable.")
        return
    _render_area_chart(df, "date", value_cols, "Labour cost composition", ylabel="Amount")


def _render_break_even_bar_chart(df: pd.DataFrame) -> None:
    if not {"product", "actual_volume", "break_even_units"}.issubset(df.columns):
        st.info("Break-even comparison data is unavailable.")
        return
    summary = df[["product", "actual_volume", "break_even_units"]].copy()
    summary = summary.replace({"": np.nan}).dropna(subset=["product"])
    if summary.empty:
        st.info("No break-even data to chart.")
        return
    _render_grouped_bar_chart(
        summary,
        "product",
        {"actual_volume": "Actual volume", "break_even_units": "Break-even volume"},
        "Actual vs break-even volumes",
        ylabel="Units",
    )


def _render_cumulative_cash_chart(df: pd.DataFrame) -> None:
    if not {"project_cumulative", "equity_cumulative"}.issubset(df.columns):
        st.info("Cumulative cash flow data is unavailable.")
        return
    _render_line_chart(
        df,
        "date",
        ["project_cumulative", "equity_cumulative"],
        "Cumulative project and equity cash flows",
        ylabel="Amount",
    )


def _render_revenue_vs_production_scatter(df: pd.DataFrame) -> None:
    if not {"volume", "revenue", "product"}.issubset(df.columns):
        st.info("Revenue vs production data is unavailable.")
        return
    _render_scatter_chart(
        df,
        "volume",
        "revenue",
        "product",
        "Revenue vs production",
        xlabel="Production volume",
        ylabel="Revenue",
    )


def _render_pnl_chart(df: pd.DataFrame) -> None:
    columns = [col for col in ("Revenue", "EBITDA", "NetIncome") if col in df.columns]
    if not columns:
        st.info("P&L drivers are unavailable for charting.")
        return
    _render_line_chart(df, "date", columns, "Monthly income statement drivers", ylabel="Amount")


def _render_cashflow_chart(df: pd.DataFrame) -> None:
    columns = [col for col in ("CFO", "CFI", "CFF", "NetCashFlow") if col in df.columns]
    if not columns:
        st.info("Cash flow components are unavailable for charting.")
        return
    _render_line_chart(df, "date", columns, "Monthly cash flow components", ylabel="Amount")


def _render_balance_sheet_chart(df: pd.DataFrame) -> None:
    columns = [col for col in ("TotalAssets", "TotalLiabilities", "Equity") if col in df.columns]
    if not columns:
        st.info("Balance sheet totals are unavailable for charting.")
        return
    _render_line_chart(df, "date", columns, "Balance sheet trend", ylabel="Amount")


def _render_break_even_comparison_chart(df: pd.DataFrame) -> None:
    if df.empty or not {"product", "break_even_revenue", "reference_volume"}.issubset(df.columns):
        st.info("Break-even revenue data is unavailable.")
        return
    subset = df[["product", "break_even_revenue", "reference_volume"]].copy()
    subset.rename(columns={"break_even_revenue": "Break-even revenue", "reference_volume": "Reference volume"}, inplace=True)
    _render_grouped_bar_chart(
        subset,
        "product",
        {"Break-even revenue": "Break-even revenue", "Reference volume": "Reference volume"},
        "Break-even revenue vs reference volume",
        ylabel="Value",
    )


def _render_tornado_chart(df: pd.DataFrame) -> None:
    if df.empty or not {"driver", "scenario", "delta"}.issubset(df.columns):
        st.info("Tornado analysis did not return usable data.")
        return

    pivot = df.pivot_table(index="driver", columns="scenario", values="delta", aggfunc="first").fillna(0.0)
    if pivot.empty:
        st.info("Tornado analysis results are empty.")
        return

    bars = pivot.loc[:, sorted(pivot.columns)]
    plot_df = bars.reset_index().rename(columns={"driver": "Driver"})
    _render_grouped_bar_chart(
        plot_df,
        "Driver",
        {col: f"{col} delta" for col in bars.columns},
        "Tornado sensitivity (Δ Project NPV)",
        ylabel="Delta",
    )


def _render_monte_carlo_histograms(samples: pd.DataFrame) -> None:
    if samples.empty:
        st.info("Monte Carlo samples are empty.")
        return

    metrics = [col for col in ("Project_NPV", "Equity_IRR", "Unit_Margin") if col in samples.columns]
    if not metrics:
        st.info("Monte Carlo outputs do not include the expected metrics.")
        return

    if plt is None or not _ensure_matplotlib():
        for metric in metrics:
            values = samples[metric].dropna()
            if values.empty:
                continue
            counts, bins = np.histogram(values, bins=30)
            hist_df = pd.DataFrame({
                metric: counts,
                "bin_start": bins[:-1],
            }).set_index("bin_start")
            st.bar_chart(hist_df, use_container_width=True)
        return

    fig, axes = plt.subplots(1, len(metrics), figsize=(6 * len(metrics), 4))
    if len(metrics) == 1:
        axes = [axes]
    for ax, metric in zip(axes, metrics):
        ax.hist(samples[metric].dropna(), bins=30, color="#1f77b4", alpha=0.7)
        ax.set_title(_friendly_label(metric))
        ax.set_xlabel(_friendly_label(metric))
        ax.set_ylabel("Frequency")
        ax.grid(True, linestyle="--", alpha=0.3)
    fig.suptitle("Monte Carlo outcome distribution")
    _finalize_plot(fig)


def _render_decision_tree_chart(df: pd.DataFrame, objective: str) -> None:
    if df.empty or not {"path_name", "metric", "probability"}.issubset(df.columns):
        st.info("Decision tree results lack the required columns.")
        return

    data = df[["path_name", "metric", "probability"]].copy()
    if data.empty:
        st.info("Decision tree produced no active paths.")
        return

    if plt is None or not _ensure_matplotlib():
        st.dataframe(data.rename(columns={"path_name": "Path"}), hide_index=True, use_container_width=True)
        chart_df = data.set_index("path_name")["metric"].to_frame()
        st.bar_chart(chart_df, use_container_width=True)
        return

    fig, ax1 = plt.subplots(figsize=(8, 4))
    x = np.arange(len(data))
    width = 0.35
    ax1.bar(x - width / 2, data["metric"], width, color="#1f77b4", label=objective)
    ax1.set_ylabel(_friendly_label(objective))
    ax1.set_xticks(x)
    ax1.set_xticklabels(data["path_name"], rotation=45, ha="right")
    ax1.grid(axis="y", linestyle="--", alpha=0.3)

    ax2 = ax1.twinx()
    ax2.bar(x + width / 2, data["probability"], width, color="#ff7f0e", alpha=0.6, label="Probability")
    ax2.set_ylabel("Probability")

    lines, labels = ax1.get_legend_handles_labels()
    lines2, labels2 = ax2.get_legend_handles_labels()
    ax1.legend(lines + lines2, labels + labels2, loc="upper right")
    ax1.set_title("Decision tree payoffs and probabilities")
    _finalize_plot(fig)


def _render_scenario_metric_chart(df: pd.DataFrame) -> None:
    if df.empty or "scenario" not in df.columns:
        st.info("Scenario metrics are unavailable for charting.")
        return
    metrics = [col for col in ("Project_NPV", "Project_IRR", "Equity_IRR", "Payback_Year", "DSCR_min") if col in df.columns]
    if not metrics:
        st.info("Scenario metrics table does not include the expected KPIs.")
        return
    plot_df = df[["scenario", *metrics]].copy()
    _render_grouped_bar_chart(
        plot_df,
        "scenario",
        {metric: _friendly_label(metric) for metric in metrics},
        "Scenario KPI comparison",
        ylabel="Value",
    )


def _render_scenario_cashflow_stack(results_map: Mapping[str, Mapping[str, object]]) -> None:
    records: List[pd.Series] = []
    for scenario, res in results_map.items():
        cash = res.get("statements_annual", {}).get("cashflow") if isinstance(res.get("statements_annual"), Mapping) else None
        if isinstance(cash, pd.DataFrame) and not cash.empty:
            totals = cash[[col for col in ("CFO", "CFI", "CFF") if col in cash.columns]].sum()
            totals["scenario"] = scenario
            records.append(totals)
    if not records:
        st.info("Cash flow statements are unavailable for the selected scenarios.")
        return
    df = pd.DataFrame(records).fillna(0.0)
    _render_stacked_columns(df, "scenario", [col for col in df.columns if col != "scenario"], "Scenario cash flow composition", ylabel="Amount")


def _render_scenario_dscr_chart(results_map: Mapping[str, Mapping[str, object]]) -> None:
    frames = []
    for scenario, res in results_map.items():
        dashboard = res.get("dashboard") if isinstance(res, Mapping) else None
        dscr = dashboard.get("dscr_trend") if isinstance(dashboard, Mapping) else None
        if isinstance(dscr, pd.DataFrame) and not dscr.empty and "DSCR" in dscr.columns:
            frame = dscr[["date", "DSCR"]].copy()
            frame["scenario"] = scenario
            frames.append(frame)
    if not frames:
        st.info("DSCR trends are unavailable for the selected scenarios.")
        return
    data = pd.concat(frames, ignore_index=True)
    data["date"] = pd.to_datetime(data["date"], errors="coerce")
    data = data.dropna(subset=["date"]).sort_values("date")

    if plt is None or not _ensure_matplotlib():
        chart_df = data.pivot_table(index="date", columns="scenario", values="DSCR")
        st.line_chart(chart_df, use_container_width=True)
        return

    fig, ax = plt.subplots(figsize=(8, 4))
    for scenario, group in data.groupby("scenario"):
        ax.plot(group["date"], group["DSCR"], label=scenario)
    ax.set_title("DSCR trend by scenario")
    ax.set_ylabel("DSCR")
    ax.set_xlabel("Date")
    if mdates:
        ax.xaxis.set_major_formatter(mdates.DateFormatter("%Y-%m"))
    ax.grid(True, linestyle="--", alpha=0.3)
    ax.legend()
    _finalize_plot(fig)


def _render_scenario_scatter_chart(results_map: Mapping[str, Mapping[str, object]]) -> None:
    frames = []
    for scenario, res in results_map.items():
        dashboard = res.get("dashboard") if isinstance(res, Mapping) else None
        rev_prod = dashboard.get("revenue_vs_production") if isinstance(dashboard, Mapping) else None
        if isinstance(rev_prod, pd.DataFrame) and not rev_prod.empty:
            frame = rev_prod[["product", "volume", "revenue"]].copy()
            frame["scenario"] = scenario
            frames.append(frame)
    if not frames:
        st.info("Revenue vs production data is unavailable for the selected scenarios.")
        return
    data = pd.concat(frames, ignore_index=True)
    if plt is None or not _ensure_matplotlib():
        chart_df = data.rename(columns={"scenario": "Scenario"})
        st.scatter_chart(chart_df, x="volume", y="revenue", color="Scenario", use_container_width=True)
        return

    fig, ax = plt.subplots(figsize=(8, 4))
    for scenario, group in data.groupby("scenario"):
        ax.scatter(group["volume"], group["revenue"], label=scenario)
    ax.set_title("Scenario revenue vs production")
    ax.set_xlabel("Volume")
    ax.set_ylabel("Revenue")
    ax.grid(True, linestyle="--", alpha=0.3)
    ax.legend()
    _finalize_plot(fig)
def _scenario_overrides_from_table(scenario_df: pd.DataFrame) -> Dict[str, Dict[str, object]]:
    """Convert the scenario comparison table into override dictionaries."""

    if not isinstance(scenario_df, pd.DataFrame) or scenario_df.empty:
        return {}

    working = scenario_df.copy()
    working = working.replace({"": np.nan})

    enabled = working.get("enabled", True)
    if not isinstance(enabled, pd.Series):
        enabled = pd.Series(True, index=working.index)
    working = working[enabled.fillna(True)]
    working = working.dropna(subset=["scenario_name"], how="any")

    overrides: Dict[str, Dict[str, object]] = {}
    for _, row in working.iterrows():
        name = str(row.get("scenario_name", "")).strip()
        if not name:
            continue
        override: Dict[str, object] = {}
        production_override: Dict[str, object] = {}
        feedstock = str(row.get("feedstock_scenario", "")).strip()
        if feedstock:
            production_override["feedstock_scenario"] = feedstock.upper()
        farm_share_val = row.get("farm_share")
        if pd.notna(farm_share_val):
            try:
                production_override["farm_share"] = float(farm_share_val)
            except (TypeError, ValueError):
                pass
        if production_override:
            override["production"] = production_override
        increment_profile = str(row.get("increment_profile", "")).strip()
        if increment_profile:
            override["increment_profile"] = increment_profile
        if override:
            overrides[name] = override

    return overrides


def _apply_scenario_override(base_cfg: Dict[str, object], override: Dict[str, object]) -> Dict[str, object]:
    """Return a deep-copied configuration with scenario overrides applied."""

    scenario_cfg = copy.deepcopy(base_cfg)
    for section, values in override.items():
        if isinstance(values, Mapping) and isinstance(scenario_cfg.get(section), Mapping):
            target = copy.deepcopy(scenario_cfg.get(section))
            for key, value in values.items():
                target[key] = value
            scenario_cfg[section] = target
        else:
            scenario_cfg[section] = values
    return scenario_cfg


BASE_SCENARIO_LABEL = "Base case"


def _ensure_scenario_payload(
    selected: str,
    base_cfg: Dict[str, object],
    base_results: Mapping[str, object],
    overrides: Mapping[str, Dict[str, object]],
) -> Tuple[Dict[str, object], Mapping[str, object]]:
    """Return (cfg, results) for the selected scenario, caching evaluations."""

    cache: Dict[str, Tuple[Dict[str, object], Mapping[str, object]]] = st.session_state.setdefault(
        "scenario_payload_cache", {}
    )

    if selected == BASE_SCENARIO_LABEL or selected not in overrides:
        cache[BASE_SCENARIO_LABEL] = (copy.deepcopy(base_cfg), base_results)
        st.session_state.scenario_payload_cache = cache
        return base_cfg, base_results

    if selected in cache:
        cfg_cached, results_cached = cache[selected]
        return cfg_cached, results_cached

    override = overrides[selected]
    scenario_cfg = _apply_scenario_override(base_cfg, override)
    scenario_results = run_full_model(scenario_cfg)
    cache[selected] = (scenario_cfg, scenario_results)
    st.session_state.scenario_payload_cache = cache
    return scenario_cfg, scenario_results


def _generate_excel_bytes(
    cfg: Mapping[str, object],
    results: Mapping[str, object],
    scenario_name: str,
) -> bytes:
    """Create an Excel workbook for the provided results and return its bytes."""

    buffer = BytesIO()
    engine = resolve_excel_engine()
    sheets: OrderedDict[str, pd.DataFrame] = OrderedDict()
    chart_builders: OrderedDict[str, Callable[[], Optional["plt.Figure"]]] = OrderedDict()

    def _to_sheet_name(name: str) -> str:
        clean = re.sub(r"[:\\\\/?*\\[\\]]", "_", name).strip()
        return clean[:31] if len(clean) > 31 else clean

    def _make_line_figure(df: pd.DataFrame, x_col: str, value_cols: Sequence[str], title: str, ylabel: str = "Value") -> Optional["plt.Figure"]:
        if plt is None or not isinstance(df, pd.DataFrame) or df.empty or x_col not in df.columns:
            return None
        cols = [c for c in value_cols if c in df.columns]
        if not cols:
            return None
        fig, ax = plt.subplots(figsize=(9, 4))
        for col in cols:
            ax.plot(df[x_col], pd.to_numeric(df[col], errors="coerce"), label=_friendly_label(col))
        ax.set_title(title)
        ax.set_ylabel(ylabel)
        ax.legend(loc="best")
        fig.tight_layout()
        return fig

    def _make_bar_figure(df: pd.DataFrame, category_col: str, value_cols: Sequence[str], title: str, ylabel: str = "Value") -> Optional["plt.Figure"]:
        if plt is None or not isinstance(df, pd.DataFrame) or df.empty or category_col not in df.columns:
            return None
        cols = [c for c in value_cols if c in df.columns]
        if not cols:
            return None
        fig, ax = plt.subplots(figsize=(9, 4))
        plot_df = df[[category_col, *cols]].copy().set_index(category_col)
        plot_df = plot_df.apply(pd.to_numeric, errors="coerce").fillna(0.0)
        plot_df.plot(kind="bar", ax=ax)
        ax.set_title(title)
        ax.set_ylabel(ylabel)
        fig.tight_layout()
        return fig

    dashboard = results.get("dashboard") if isinstance(results, Mapping) else None
    metrics = results.get("metrics", {}) if isinstance(results, Mapping) else {}
    statements_monthly = results.get("statements_monthly", {}) if isinstance(results, Mapping) else {}
    statements_annual = results.get("statements_annual") if isinstance(results, Mapping) else None

    summary_rows = [
        {"Metric": k, "Value": v}
        for k, v in (metrics.items() if isinstance(metrics, Mapping) else [])
    ]
    if summary_rows:
        sheets["Summary"] = pd.DataFrame(summary_rows)
    if isinstance(dashboard, Mapping):
        snapshot = dashboard.get("assumptions_snapshot")
        if isinstance(snapshot, pd.DataFrame) and not snapshot.empty:
            sheets["Summary_Assumptions"] = snapshot
    if "Summary" in sheets:
        chart_builders[_to_sheet_name("Summary_Plots")] = lambda: _make_bar_figure(
            sheets["Summary"].head(10), "Metric", ["Value"], "Summary KPI snapshot"
        )

    financial_sheet = []
    if isinstance(statements_annual, Mapping):
        for key in ("pnl", "cashflow", "balancesheet"):
            df = statements_annual.get(key)
            if isinstance(df, pd.DataFrame) and not df.empty:
                sheets[f"Financial_{key}"] = df
                financial_sheet.append((key, df))
    if financial_sheet:
        def _financial_fig() -> Optional["plt.Figure"]:
            if plt is None:
                return None
            key, df = financial_sheet[0]
            x_col = "year" if "year" in df.columns else "date"
            value_cols = [c for c in df.columns if c not in {x_col}][:4]
            return _make_line_figure(df, x_col, value_cols, f"Financial Statements ({key.upper()})", "Amount")
        chart_builders[_to_sheet_name("Financial_Statements_Plots")] = _financial_fig

    prod_df = results.get("production_monthly") if isinstance(results, Mapping) else None
    price_df = results.get("price_curves") if isinstance(results, Mapping) else None
    revenue_df = results.get("revenue") if isinstance(results, Mapping) else None
    if isinstance(prod_df, pd.DataFrame) and not prod_df.empty:
        sheets["Production_Pricing_Prod"] = prod_df
    if isinstance(price_df, pd.DataFrame) and not price_df.empty:
        sheets["Production_Pricing_Price"] = price_df
    if isinstance(revenue_df, pd.DataFrame) and not revenue_df.empty:
        sheets["Production_Pricing_Revenue"] = revenue_df

    def _prod_price_fig() -> Optional["plt.Figure"]:
        if plt is None:
            return None
        if isinstance(price_df, pd.DataFrame) and not price_df.empty and {"date", "product", "price"}.issubset(price_df.columns):
            pivot = price_df.pivot_table(index="date", columns="product", values="price", aggfunc="mean").reset_index()
            return _make_line_figure(pivot, "date", [c for c in pivot.columns if c != "date"], "Production & Pricing - Price curves", "Price")
        if isinstance(prod_df, pd.DataFrame) and not prod_df.empty and {"date", "product", "volume"}.issubset(prod_df.columns):
            pivot = prod_df.pivot_table(index="date", columns="product", values="volume", aggfunc="sum").reset_index()
            return _make_line_figure(pivot, "date", [c for c in pivot.columns if c != "date"], "Production & Pricing - Volume", "Volume")
        return None
    chart_builders[_to_sheet_name("Production_Pricing_Plots")] = _prod_price_fig

    if isinstance(dashboard, Mapping):
        dashboard_sheet_map: OrderedDict[str, str] = OrderedDict(
            [
                ("Dashboard_Annual_Production", "annual_production"),
                ("Dashboard_Annual_Cashflow", "annual_cashflow"),
                ("Dashboard_DSCR_Trend", "dscr_trend"),
                ("Dashboard_Debt_Service", "debt_service_summary"),
                ("Dashboard_Working_Capital", "working_capital_trend"),
                ("Dashboard_Cost_Structure", "cost_structure_annual"),
                ("Dashboard_Labour_Summary", "labour_summary"),
                ("Dashboard_Break_Even", "break_even_per_product"),
                ("Dashboard_Cumulative_CF", "cumulative_cashflows"),
                ("Dashboard_Revenue_vs_Prod", "revenue_vs_production"),
            ]
        )
        for sheet_name, dashboard_key in dashboard_sheet_map.items():
            df = dashboard.get(dashboard_key)
            if isinstance(df, pd.DataFrame) and not df.empty:
                sheets[sheet_name] = df

        annual_production_df = dashboard.get("annual_production")
        if isinstance(annual_production_df, pd.DataFrame) and not annual_production_df.empty:
            chart_builders[_to_sheet_name("Dashboard_Annual_Production_Plots")] = lambda: _make_bar_figure(
                annual_production_df, "year", [c for c in annual_production_df.columns if c != "year"], "Annual production by product", "Volume"
            )

        annual_cashflow_df = dashboard.get("annual_cashflow")
        if isinstance(annual_cashflow_df, pd.DataFrame) and not annual_cashflow_df.empty:
            chart_builders[_to_sheet_name("Dashboard_Annual_Cashflow_Plots")] = lambda: _make_bar_figure(
                annual_cashflow_df, "year", [c for c in ("NetCashFlow", "CFO") if c in annual_cashflow_df.columns], "Annual cash flow", "Amount"
            )

        dscr_df = dashboard.get("dscr_trend")
        if isinstance(dscr_df, pd.DataFrame) and not dscr_df.empty:
            chart_builders[_to_sheet_name("Dashboard_DSCR_Trend_Plots")] = lambda: _make_line_figure(
                dscr_df, "date", [c for c in ("DSCR", "CFADS", "debt_service") if c in dscr_df.columns], "DSCR and debt service trend", "Value"
            )

        debt_service_df = dashboard.get("debt_service_summary")
        if isinstance(debt_service_df, pd.DataFrame) and not debt_service_df.empty:
            chart_builders[_to_sheet_name("Dashboard_Debt_Service_Plots")] = lambda: _make_bar_figure(
                debt_service_df, "year", [c for c in ("interest", "principal", "draw") if c in debt_service_df.columns], "Debt service summary", "Amount"
            )

        wc_df = dashboard.get("working_capital_trend")
        if isinstance(wc_df, pd.DataFrame) and not wc_df.empty:
            chart_builders[_to_sheet_name("Dashboard_Working_Capital_Plots")] = lambda: _make_line_figure(
                wc_df, "date", [c for c in ("accounts_receivable", "inventory", "accounts_payable") if c in wc_df.columns], "Working capital components", "Amount"
            )

        cost_df = dashboard.get("cost_structure_annual")
        if isinstance(cost_df, pd.DataFrame) and not cost_df.empty:
            chart_builders[_to_sheet_name("Dashboard_Cost_Structure_Plots")] = lambda: _make_bar_figure(
                cost_df, "year", [c for c in cost_df.columns if c != "year"], "Annual cost structure", "Amount"
            )

        labour_df = dashboard.get("labour_summary")
        if isinstance(labour_df, pd.DataFrame) and not labour_df.empty and {"date", "dept", "total_cost"}.issubset(labour_df.columns):
            def _labour_fig() -> Optional["plt.Figure"]:
                labour_plot = labour_df.groupby(["date", "dept"])["total_cost"].sum().unstack(fill_value=0.0).reset_index()
                return _make_line_figure(labour_plot, "date", [c for c in labour_plot.columns if c != "date"], "Labour cost by department", "Amount")
            chart_builders[_to_sheet_name("Dashboard_Labour_Plots")] = _labour_fig

        break_even_df = dashboard.get("break_even_per_product")
        if isinstance(break_even_df, pd.DataFrame) and not break_even_df.empty:
            chart_builders[_to_sheet_name("Dashboard_Break_Even_Plots")] = lambda: _make_bar_figure(
                break_even_df, "product", [c for c in ("actual_volume", "break_even_units") if c in break_even_df.columns], "Actual vs break-even volume", "Units"
            )

        cumulative_cf_df = dashboard.get("cumulative_cashflows")
        if isinstance(cumulative_cf_df, pd.DataFrame) and not cumulative_cf_df.empty:
            chart_builders[_to_sheet_name("Dashboard_Cumulative_CF_Plots")] = lambda: _make_line_figure(
                cumulative_cf_df, "date", [c for c in ("project_cumulative", "equity_cumulative") if c in cumulative_cf_df.columns], "Cumulative cash flows", "Amount"
            )

        rev_prod_df = dashboard.get("revenue_vs_production")
        if isinstance(rev_prod_df, pd.DataFrame) and not rev_prod_df.empty:
            chart_builders[_to_sheet_name("Dashboard_Revenue_vs_Prod_Plots")] = lambda: _make_line_figure(
                rev_prod_df, "volume", [c for c in ("revenue", "average_price") if c in rev_prod_df.columns], "Revenue vs production", "Value"
            )

    # Sensitivity page content
    sens_df = results.get("sensitivities") if isinstance(results, Mapping) else None
    if not isinstance(sens_df, pd.DataFrame) or sens_df.empty:
        try:
            if isinstance(cfg, Mapping):
                sens_df = sensitivity_tornado(cfg, {"metrics": results.get("metrics", {})}, lambda c: run_full_model(c), None)
        except Exception:
            sens_df = pd.DataFrame()
    if isinstance(sens_df, pd.DataFrame) and not sens_df.empty:
        sheets["Sensitivities"] = sens_df
        chart_builders[_to_sheet_name("Sensitivities_Plots")] = lambda: _make_bar_figure(
            sens_df.head(12), "driver" if "driver" in sens_df.columns else sens_df.columns[0], ["delta"] if "delta" in sens_df.columns else [sens_df.columns[-1]], "Sensitivity tornado (top drivers)"
        )

    for label, key in (
        ("IFRS_SoPL_OCI_Monthly", "sopl_oci_monthly"),
        ("IFRS_SoPL_OCI_Annual", "sopl_oci_annual"),
        ("IFRS_SoFP_Monthly", "sofp_monthly"),
        ("IFRS_SoFP_Annual", "sofp_annual"),
        ("IFRS_SoCE_Monthly", "socie_monthly"),
        ("IFRS_SoCE_Annual", "socie_annual"),
        ("IFRS_SCF_Indirect_Monthly", "scf_indirect_monthly"),
        ("IFRS_SCF_Indirect_Annual", "scf_indirect_annual"),
        ("IFRS_Note_PPE", "ifrs_note_ppe_rollforward"),
        ("IFRS_Note_Debt_Maturity", "ifrs_note_debt_maturity"),
        ("IFRS_Note_WC_Bridge", "ifrs_note_wc_bridge"),
        ("IFRS_Note_Deferred_Tax", "ifrs_note_deferred_tax"),
        ("IFRS_Note_Lease", "ifrs_note_lease"),
        ("IFRS_Note_Hedge_Reserve", "ifrs_note_hedge_reserve"),
    ):
        df = results.get(key) if isinstance(results, Mapping) else None
        if isinstance(df, pd.DataFrame) and not df.empty:
            sheets[label] = df

    for label, key in (("CAPEX", "capex"), ("Debt", "debt_schedule"), ("WorkingCapital", "working_capital")):
        df = results.get(key) if isinstance(results, Mapping) else None
        if isinstance(df, pd.DataFrame) and not df.empty:
            sheets[label] = df

    meta_rows: List[Dict[str, object]] = []
    global_inputs = cfg.get("global_inputs") if isinstance(cfg, Mapping) else {}
    if isinstance(global_inputs, Mapping):
        meta_rows.append(
            {
                "Scenario": scenario_name,
                "Discount rate": global_inputs.get("discount_rate"),
                "Corporate tax": global_inputs.get("corp_tax_rate"),
                "Investor share": global_inputs.get("investor_share"),
                "Owner share": global_inputs.get("owner_share"),
            }
        )
    if meta_rows:
        sheets["Scenario"] = pd.DataFrame(meta_rows)

    if engine == SIMPLE_XLSX_ENGINE:
        write_simple_xlsx(sheets, buffer)
    else:
        with pd.ExcelWriter(buffer, engine=engine) as writer:
            for sheet_name, df in sheets.items():
                df.to_excel(writer, sheet_name=_to_sheet_name(sheet_name), index=False)
            # Add chart sheets where supported.
            if plt is not None:
                for raw_sheet_name, fig_builder in chart_builders.items():
                    fig = None
                    try:
                        fig = fig_builder()
                    except Exception:
                        fig = None
                    if fig is None:
                        continue
                    image_stream = BytesIO()
                    fig.savefig(image_stream, format="png", dpi=140, bbox_inches="tight")
                    image_stream.seek(0)
                    plt.close(fig)
                    sheet_name = _to_sheet_name(raw_sheet_name)
                    if engine == "xlsxwriter":
                        worksheet = writer.book.add_worksheet(sheet_name)
                        writer.sheets[sheet_name] = worksheet
                        worksheet.write(0, 0, f"{sheet_name} ({scenario_name})")
                        worksheet.insert_image(2, 0, "chart.png", {"image_data": image_stream})
                    elif engine == "openpyxl":
                        ws = writer.book.create_sheet(title=sheet_name)
                        ws["A1"] = f"{sheet_name} ({scenario_name})"
                        try:
                            from openpyxl.drawing.image import Image as OpenPyxlImage
                            image_stream.seek(0)
                            img = OpenPyxlImage(image_stream)
                            ws.add_image(img, "A3")
                        except Exception:
                            ws["A3"] = "Chart image embedding unavailable (install pillow)."

    buffer.seek(0)
    return buffer.read()

def _get_risk_select_options(tables: "InputTables") -> Tuple[List[str], List[str]]:
    """Return dropdown option sets for risk targets and scope fields."""

    target_set = set(RISK_TARGET_OPTIONS)
    applies_set = set(RISK_APPLIES_OPTIONS)

    default_risk = DEFAULTS.get("risk_params")
    sources: List[pd.DataFrame] = []
    if isinstance(default_risk, pd.DataFrame) and not default_risk.empty:
        sources.append(default_risk)
    try:
        active_risk = tables.ensure_table("risk_params")
    except Exception:
        active_risk = pd.DataFrame()
    if isinstance(active_risk, pd.DataFrame) and not active_risk.empty:
        sources.append(active_risk)

    for source in sources:
        if "target" in source:
            target_set.update(str(val).lower() for val in source["target"].dropna().unique())
        if "applies_to" in source:
            applies_set.update(str(val).lower() for val in source["applies_to"].dropna().unique())

    target_options = sorted(option for option in target_set if option)
    applies_options = sorted(option for option in applies_set if option)
    return target_options, applies_options


def _render_risk_schedule_preview(tables: "InputTables") -> None:
    """Display the normalised risk schedule and consolidated multipliers."""

    risk_df = tables.ensure_table("risk_params").copy()
    risk_df = risk_df.replace({"": np.nan}).dropna(how="all")
    if risk_df.empty:
        st.info("Add at least one risk driver above to populate the active schedule.")
        return

    risk_df = risk_df.reset_index(drop=True)
    _render_dataframe(risk_df, "Active risk schedule", key="risk_schedule_preview")

    profile = compute_risk_profile(risk_df)
    if not isinstance(profile, dict) or not profile:
        return

    summary_rows: List[Dict[str, object]] = []
    profile_labels = [
        ("production", "Production multiplier"),
        ("labour", "Labour multiplier"),
        ("price", "Price multiplier"),
        ("revenue", "Revenue multiplier"),
        ("yield", "Yield multiplier"),
    ]
    for key_name, label in profile_labels:
        value = profile.get(key_name)
        if value is None:
            continue
        try:
            summary_rows.append({"Metric": label, "Multiplier": float(value)})
        except (TypeError, ValueError):  # pragma: no cover - defensive guard
            continue
    if summary_rows:
        summary_df = pd.DataFrame(summary_rows)
        _render_dataframe(summary_df, "Consolidated risk multipliers", key="risk_profile_summary")

    price_map = profile.get("price_by_product")
    if isinstance(price_map, dict) and price_map:
        price_rows: List[Dict[str, object]] = []
        for product, value in price_map.items():
            try:
                multiplier = float(value)
            except (TypeError, ValueError):
                continue
            product_label = str(product).replace("_", " ").title()
            price_rows.append({"Product": product_label, "Multiplier": multiplier})
        if price_rows:
            price_df = pd.DataFrame(price_rows)
            _render_dataframe(price_df, "Price multipliers by product", key="risk_profile_prices")


def _display_default_value(value: object) -> str:
    """Pretty-print values for the default-row manager."""
    if value is None:
        return "—"
    if isinstance(value, (float, np.floating)):
        if math.isnan(float(value)):
            return "—"
        magnitude = abs(float(value))
        if magnitude >= 1_000:
            return f"{float(value):,.2f}"
        if magnitude >= 1:
            return f"{float(value):,.2f}"
        return (f"{float(value):.6f}").rstrip("0").rstrip(".")
    if isinstance(value, (int, np.integer)):
        return f"{int(value):,}"
    if isinstance(value, pd.Timestamp):
        return value.strftime("%Y-%m-%d")
    return str(value)


def _factory_default_frames() -> Dict[str, pd.DataFrame]:
    """Return deep copies of the model's factory default tables."""

    global _FACTORY_DEFAULT_FRAMES_CACHE
    if _FACTORY_DEFAULT_FRAMES_CACHE is None:
        frames: Dict[str, pd.DataFrame] = {}

        horizon_defaults = dict(DEFAULTS["horizon"])
        frames["projection_horizon"] = pd.DataFrame([horizon_defaults])
        frames["production_horizon"] = pd.DataFrame([DEFAULTS["production_horizon"]])
        frames["global_inputs"] = pd.DataFrame([DEFAULTS["global"]])
        frames["working_capital_days"] = pd.DataFrame([DEFAULTS["working_capital"]])

        start_year = int(horizon_defaults["start_year"])
        start_month = int(horizon_defaults.get("start_month", 1))

        capex_rows = [
            {
                "item_name": "Land Acquisition",
                "category": "land",
                "amount": 5_000_000.0,
                "currency": "USD",
                "fx_curve": "",
                "start_date": f"{start_year}-01",
                "end_date": f"{start_year}-12",
                "life_years": 40,
                "depr_method": "straight",
                "depr_rate_override": np.nan,
                "vat_rate": 0.0,
                "vat_recovery_lag_months": 0,
                "capitalized": True,
                "is_farm_capex": True,
            },
            {
                "item_name": "Civil Works",
                "category": "civil",
                "amount": 8_500_000.0,
                "currency": "USD",
                "fx_curve": "",
                "start_date": f"{start_year}-01",
                "end_date": f"{start_year}-12",
                "life_years": 20,
                "depr_method": "straight",
                "depr_rate_override": np.nan,
                "vat_rate": 0.0,
                "vat_recovery_lag_months": 0,
                "capitalized": True,
                "is_farm_capex": False,
            },
            {
                "item_name": "Process Equipment",
                "category": "equipment",
                "amount": 12_500_000.0,
                "currency": "USD",
                "fx_curve": "",
                "start_date": f"{start_year}-01",
                "end_date": f"{start_year}-12",
                "life_years": 12,
                "depr_method": "straight",
                "depr_rate_override": np.nan,
                "vat_rate": 0.0,
                "vat_recovery_lag_months": 0,
                "capitalized": True,
                "is_farm_capex": False,
            },
            {
                "item_name": "Farm Machinery",
                "category": "farm_machinery",
                "amount": 4_000_000.0,
                "currency": "USD",
                "fx_curve": "",
                "start_date": f"{start_year}-01",
                "end_date": f"{start_year}-12",
                "life_years": 10,
                "depr_method": "straight",
                "depr_rate_override": np.nan,
                "vat_rate": 0.0,
                "vat_recovery_lag_months": 0,
                "capitalized": True,
                "is_farm_capex": True,
            },
        ]
        frames["capex_lines"] = pd.DataFrame(capex_rows)

        price_rows = []
        for product, params in DEFAULTS["prices"].items():
            price_rows.append(
                {
                    "product": product,
                    "base_price": params.get("base_price", 0.0),
                    "price_escalation_pa": params.get("price_escalation_pa", 0.0),
                    "price_indexation": params.get("price_indexation", "cpi"),
                    "uom": params.get("uom", ""),
                    "tariff_structure": "",
                    "revenue_share": 1.0,
                }
            )
        frames["revenue_params"] = pd.DataFrame(price_rows)

        prod_defaults = DEFAULTS["production"]
        feedstock = prod_defaults["annual_feedstock_ton"]
        availability = prod_defaults["plant_availability"]
        loss = prod_defaults["loss_factor"]
        ramp = "0.7;0.9;1.0"

        production_rows = [
            {
                "product": "ethanol",
                "annual_volume": feedstock * prod_defaults["ethanol_litre_per_ton"] * availability * (1 - loss),
                "availability": availability,
                "loss_factor": loss,
                "startup_ramp": ramp,
                "boe_conversion": np.nan,
                "sugarcane_yield_ton_per_ha": prod_defaults["sugarcane_yield_ton_per_ha"],
                "farm_area_ha": feedstock / prod_defaults["sugarcane_yield_ton_per_ha"],
            },
            {
                "product": "sugar",
                "annual_volume": feedstock * prod_defaults["sugar_ton_per_ton_cane"] * availability * (1 - loss),
                "availability": availability,
                "loss_factor": loss,
                "startup_ramp": ramp,
                "boe_conversion": np.nan,
                "sugarcane_yield_ton_per_ha": prod_defaults["sugarcane_yield_ton_per_ha"],
                "farm_area_ha": feedstock / prod_defaults["sugarcane_yield_ton_per_ha"],
            },
            {
                "product": "electricity",
                "annual_volume": feedstock * prod_defaults["electricity_mwh_per_ton_cane"] * availability * (1 - loss),
                "availability": availability,
                "loss_factor": loss,
                "startup_ramp": ramp,
                "boe_conversion": np.nan,
                "sugarcane_yield_ton_per_ha": prod_defaults["sugarcane_yield_ton_per_ha"],
                "farm_area_ha": feedstock / prod_defaults["sugarcane_yield_ton_per_ha"],
            },
            {
                "product": "animal_feed",
                "annual_volume": feedstock * prod_defaults["animal_feed_ton_per_ton_cane"] * availability * (1 - loss),
                "availability": availability,
                "loss_factor": loss,
                "startup_ramp": ramp,
                "boe_conversion": np.nan,
                "sugarcane_yield_ton_per_ha": prod_defaults["sugarcane_yield_ton_per_ha"],
                "farm_area_ha": feedstock / prod_defaults["sugarcane_yield_ton_per_ha"],
            },
        ]
        frames["production_annual"] = pd.DataFrame(production_rows)

        first_year_months = pd.date_range(f"{start_year}-{start_month:02d}-01", periods=12, freq="MS")
        monthly_rows = []
        for row in production_rows:
            monthly_volume = float(row["annual_volume"]) / 12.0
            for date in first_year_months:
                monthly_rows.append(
                    {
                        "date": date.strftime("%Y-%m"),
                        "product": row["product"],
                        "volume": monthly_volume,
                        "availability_override": np.nan,
                        "maintenance_downtime": np.nan,
                        "loss_override": np.nan,
                    }
                )
        frames["production_monthly"] = pd.DataFrame(monthly_rows)

        frames["direct_costs_monthly"] = pd.DataFrame(
            [
                {
                    "date": f"{start_year}-01",
                    "cost_type": "feedstock purchase",
                    "product_link": "ethanol",
                    "unit_price": 65.0,
                    "quantity": 8_000.0,
                    "amount": 520_000.0,
                    "currency": "USD",
                }
            ]
        )

        frames["staff_costs_monthly"] = pd.DataFrame(
            [
                {
                    "date": f"{start_year}-01",
                    "dept": "Operations",
                    "headcount": 50,
                    "gross_pay": 150_000.0,
                    "benefits": 25_000.0,
                    "training": 5_000.0,
                    "other": 10_000.0,
                    "gross_pay_per_head": 3_000.0,
                    "benefits_per_head": 500.0,
                    "training_per_head": 100.0,
                    "other_per_head": 200.0,
                    "currency": "USD",
                }
            ]
        )

        frames["other_opex_monthly"] = pd.DataFrame(
            [
                {"date": f"{start_year}-01", "category": "insurance", "amount": 20_000.0, "currency": "USD"},
                {"date": f"{start_year}-01", "category": "service_contract", "amount": 35_000.0, "currency": "USD"},
                {"date": f"{start_year}-01", "category": "general_admin", "amount": 50_000.0, "currency": "USD"},
                {"date": f"{start_year}-01", "category": "energy_cost", "amount": 15_000.0, "currency": "USD"},
            ]
        )

        frames["ar_other_assets"] = pd.DataFrame(
            [
                {
                    "date": f"{start_year}-01",
                    "receivables": 0.0,
                    "prepaid_expenses": 0.0,
                    "other_current_assets": 0.0,
                    "dso_days": DEFAULTS["working_capital"]["dso_days"],
                }
            ]
        )

        frames["inventory_ap"] = pd.DataFrame(
            [
                {
                    "date": f"{start_year}-01",
                    "inventory_raw": 0.0,
                    "inventory_wip": 0.0,
                    "inventory_fg": 0.0,
                    "accounts_payable": 0.0,
                    "dio_days": DEFAULTS["working_capital"]["dio_days"],
                    "dpo_days": DEFAULTS["working_capital"]["dpo_days"],
                }
            ]
        )

        frames["debt_tranches"] = pd.DataFrame(DEFAULTS["debt"]["tranches"])
        frames["tax_schedule"] = pd.DataFrame([DEFAULTS["tax"]])
        frames["inflation_index"] = DEFAULTS["inflation_index"].copy()
        frames["risk_params"] = DEFAULTS["risk_params"].copy()
        frames["tornado_drivers"] = DEFAULTS["tornado_drivers"].copy()
        frames["monte_carlo_settings"] = DEFAULTS["monte_carlo_settings"].copy()
        frames["scenario_comparison"] = DEFAULTS["scenario_comparison"].copy()
        frames["optimizer_settings"] = DEFAULTS["optimizer_settings"].copy()
        frames["neural_forecast_settings"] = DEFAULTS["neural_forecast_settings"].copy()
        frames["statistical_forecast_settings"] = DEFAULTS["statistical_forecast_settings"].copy()
        frames["decision_tree_paths"] = DEFAULTS["decision_tree_paths"].copy()
        frames["break_even_inputs"] = DEFAULTS["break_even_inputs"].copy()

        for table_name, schema in INPUT_SCHEMAS.items():
            frames.setdefault(table_name, pd.DataFrame(columns=list(schema.columns.keys())))

        _FACTORY_DEFAULT_FRAMES_CACHE = {name: df.copy(deep=True) for name, df in frames.items()}

    return {name: df.copy(deep=True) for name, df in _FACTORY_DEFAULT_FRAMES_CACHE.items()}


def _ensure_default_table_store() -> Dict[str, pd.DataFrame]:
    """Ensure session state carries a mutable, schema-aligned default table store."""

    if DEFAULT_TABLE_STORE_KEY not in st.session_state:
        factory_frames = _factory_default_frames()
        seeded_tables = InputTables()
        for table_name, df in factory_frames.items():
            seeded_tables.set_table(table_name, df)
        st.session_state[DEFAULT_TABLE_STORE_KEY] = {
            name: seeded_tables.ensure_table(name).copy()
            for name in INPUT_SCHEMAS.keys()
        }

    store = st.session_state[DEFAULT_TABLE_STORE_KEY]
    for table_name, schema in INPUT_SCHEMAS.items():
        if table_name not in store:
            store[table_name] = pd.DataFrame(columns=list(schema.columns.keys()))
    return store


def _get_default_table(table_name: str) -> pd.DataFrame:
    store = _ensure_default_table_store()
    df = store.get(table_name)
    if df is None:
        schema = INPUT_SCHEMAS[table_name]
        return pd.DataFrame(columns=list(schema.columns.keys()))
    return df.copy(deep=True)


def _save_default_table(table_name: str, df: pd.DataFrame) -> None:
    temp = InputTables()
    temp.set_table(table_name, df)
    store = _ensure_default_table_store()
    store[table_name] = temp.ensure_table(table_name).copy()


def _restore_factory_default_table(table_name: str) -> pd.DataFrame:
    factory_frames = _factory_default_frames()
    df = factory_frames.get(
        table_name,
        pd.DataFrame(columns=list(INPUT_SCHEMAS[table_name].columns.keys())),
    )
    _save_default_table(table_name, df)
    return _get_default_table(table_name)


def _render_default_row_controls(
    table_name: str,
    label: str,
    schema,
    default_df: pd.DataFrame,
) -> None:
    if default_df.empty:
        st.info("No default rows defined yet – save the current table to create a baseline.")
        return

    st.markdown("##### Default baseline rows")
    display_df = default_df.reset_index(drop=True)
    for idx, row in display_df.iterrows():
        col_widths = [1] * len(schema.columns) + [0.6]
        row_cols = st.columns(col_widths)
        for position, column_name in enumerate(schema.columns.keys()):
            value = row.get(column_name, np.nan)
            display_value = _display_default_value(value)
            row_cols[position].markdown(
                f"<span style='font-size:0.7rem;color:#6b6b6b'>{column_name.replace('_', ' ').title()}</span><br>"
                f"<span style='font-weight:600'>{display_value}</span>",
                unsafe_allow_html=True,
            )
        if row_cols[-1].button("Edit", key=f"default_edit_{table_name}_{idx}"):
            st.session_state[DEFAULT_EDIT_STATE_KEY] = {"table": table_name, "row": int(idx)}
            _safe_rerun()


def _render_default_edit_modal(table_name: str, label: str, schema, tables: "InputTables") -> None:
    edit_state = st.session_state.get(DEFAULT_EDIT_STATE_KEY)
    if not edit_state or edit_state.get("table") != table_name:
        return

    row_index = int(edit_state.get("row", -1))
    store = _ensure_default_table_store()
    df = store.get(table_name)
    if df is None or not (0 <= row_index < len(df)):
        st.session_state.pop(DEFAULT_EDIT_STATE_KEY, None)
        return

    row = df.iloc[row_index]
    risk_select_cache: Optional[Tuple[List[str], List[str]]] = None

    def _field_options(column_name: str) -> Optional[List[str]]:
        nonlocal risk_select_cache
        col = column_name.lower()
        if table_name == "risk_params":
            if col == "distribution":
                return list(RISK_DISTRIBUTION_OPTIONS)
            if col in {"target", "applies_to"}:
                if risk_select_cache is None:
                    risk_select_cache = _get_risk_select_options(tables)
                targets, applies = risk_select_cache
                return list(targets if col == "target" else applies)
        elif table_name == "monte_carlo_settings":
            if col == "distribution":
                return list(MONTE_CARLO_DISTRIBUTIONS)
            if col == "variable":
                return list(MONTE_CARLO_VARIABLE_LABELS.values())
            if col == "applies_to":
                return ["global", *PRODUCTS]
        return None

    with _modal_container(f"Edit default row {row_index + 1} – {label}") as modal_supported:
        form = st.form(key=f"default_edit_form_{table_name}_{row_index}")
        updated_values: Dict[str, object] = {}
        for column_name, dtype in schema.columns.items():
            field_label = column_name.replace("_", " ").title()
            current_value = row.get(column_name, np.nan)
            if dtype == "int":
                default_val = schema.defaults.get(column_name, 0)
                base_value = int(default_val) if pd.isna(current_value) else int(current_value)
                input_value = form.number_input(
                    field_label,
                    value=base_value,
                    step=1,
                    format="%d",
                    key=f"default_edit_field_{table_name}_{row_index}_{column_name}",
                )
                updated_values[column_name] = int(input_value)
            elif dtype == "float":
                default_val = schema.defaults.get(column_name, 0.0)
                base_value = float(default_val) if pd.isna(current_value) else float(current_value)
                input_value = form.number_input(
                    field_label,
                    value=base_value,
                    step=0.01,
                    format="%.6f",
                    key=f"default_edit_field_{table_name}_{row_index}_{column_name}",
                )
                updated_values[column_name] = float(input_value)
            elif dtype == "bool":
                default_val = bool(schema.defaults.get(column_name, False))
                base_value = default_val if pd.isna(current_value) else bool(current_value)
                input_value = form.checkbox(
                    field_label,
                    value=base_value,
                    key=f"default_edit_field_{table_name}_{row_index}_{column_name}",
                )
                updated_values[column_name] = bool(input_value)
            else:
                options = _field_options(column_name)
                if options:
                    default_option = schema.defaults.get(column_name)
                    if table_name == "monte_carlo_settings" and column_name.lower() == "variable":
                        default_label = MONTE_CARLO_VARIABLE_LABELS.get(str(default_option), options[0])
                        current_option = MONTE_CARLO_VARIABLE_LABELS.get(
                            str(current_value), default_label
                        )
                    else:
                        if default_option is None or (
                            isinstance(default_option, float) and math.isnan(default_option)
                        ):
                            default_option = options[0] if options else ""
                        if current_value is None or (
                            isinstance(current_value, float) and math.isnan(current_value)
                        ):
                            current_option = str(default_option)
                        else:
                            current_option = str(current_value)
                    index = _option_index(options, current_option, default=0)
                    input_value = form.selectbox(
                        field_label,
                        options,
                        index=index,
                        key=f"default_edit_field_{table_name}_{row_index}_{column_name}",
                    )
                    if table_name == "monte_carlo_settings" and column_name.lower() == "variable":
                        updated_values[column_name] = MONTE_CARLO_VARIABLE_LABEL_TO_KEY.get(
                            str(input_value), str(input_value)
                        )
                    else:
                        updated_values[column_name] = str(input_value)
                else:
                    if current_value is None or (isinstance(current_value, float) and math.isnan(current_value)):
                        text_value = ""
                    else:
                        text_value = str(current_value)
                    input_value = form.text_input(
                        field_label,
                        value=text_value,
                        key=f"default_edit_field_{table_name}_{row_index}_{column_name}",
                    )
                    updated_values[column_name] = input_value if input_value != "" else None

        submit_label = "Save changes" if modal_supported else "Save inline changes"
        if form.form_submit_button(submit_label, type="primary"):
            updated_df = df.copy()
            for column_name, value in updated_values.items():
                updated_df.at[row_index, column_name] = value
            _save_default_table(table_name, updated_df)
            # Apply the refreshed defaults to the live table so downstream
            # schedules pick up the edited baseline immediately.
            try:
                refreshed_defaults = _get_default_table(table_name)
                tables.set_table(table_name, refreshed_defaults)
            except Exception as exc:
                st.error(f"Default saved but unable to update table: {exc}")
            else:
                _sync_table_mutation_state(table_name, tables)
                st.session_state.pop(DEFAULT_EDIT_STATE_KEY, None)
                st.session_state[f"default_feedback_{table_name}"] = "Default row updated and applied."
                _safe_rerun()

        if st.button("Cancel", key=f"default_edit_cancel_{table_name}_{row_index}"):
            st.session_state.pop(DEFAULT_EDIT_STATE_KEY, None)
            _safe_rerun()


def _seed_tables_with_defaults(tables: InputTables) -> None:
    """Populate session tables with the active default frames."""

    default_store = _ensure_default_table_store()
    for table_name, schema in INPUT_SCHEMAS.items():
        default_df = default_store.get(table_name)
        if default_df is None:
            empty_df = pd.DataFrame(columns=list(schema.columns.keys()))
            tables.set_table(table_name, empty_df)
        else:
            tables.set_table(table_name, default_df.copy())


def _get_tables() -> InputTables:
    if "input_tables" not in st.session_state:
        tables = InputTables()
        _seed_tables_with_defaults(tables)
        st.session_state["input_tables"] = tables
    return st.session_state["input_tables"]


def _sync_tables_from_state(tables: InputTables) -> Dict[str, str]:
    errors: Dict[str, str] = {}
    for table_name in INPUT_SCHEMAS.keys():
        state_key = _editor_state_key(table_name)
        value = st.session_state.get(state_key)
        if isinstance(value, pd.DataFrame):
            try:
                tables.set_table(table_name, value)
            except Exception as exc:  # pragma: no cover - validation feedback for UI edits
                errors[table_name] = str(exc)

    # Keep production horizon fused with projection horizon whenever start/end
    # years are edited from any table editor entry point (not only Model Controls).
    try:
        proj_df = tables.ensure_table("projection_horizon")
        if isinstance(proj_df, pd.DataFrame) and not proj_df.empty:
            row = proj_df.iloc[0]
            start_year = int(row.get("start_year", DEFAULTS["horizon"]["start_year"]))
            end_year = int(row.get("end_year", DEFAULTS["horizon"]["end_year"]))
            prod_target = pd.DataFrame([{"start_year": start_year, "end_year": end_year}])
            current_prod = tables.ensure_table("production_horizon")
            current_tuple = None
            if isinstance(current_prod, pd.DataFrame) and not current_prod.empty:
                cur_row = current_prod.iloc[0]
                current_tuple = (int(cur_row.get("start_year", start_year)), int(cur_row.get("end_year", end_year)))
            target_tuple = (start_year, end_year)
            if current_tuple != target_tuple:
                tables.set_table("production_horizon", prod_target)
                _sync_table_mutation_state("production_horizon", tables)
    except Exception as exc:  # pragma: no cover - defensive sync guard
        errors["production_horizon"] = str(exc)
    return errors


LANDING_TABLES: List[Tuple[str, str, Optional[str]]] = [
    ("Projection Horizon", "projection_horizon", "Define the calendar start and end of the modeling period."),
    ("Global Inputs", "global_inputs", "Corporate tax, discount rate, and ownership split."),
    (
        "Working Capital Assumptions",
        "working_capital_days",
        "DSO, DIO, and DPO assumptions controlling receivables, inventory, and payables timing.",
    ),
    (
        "Initial Investment (CAPEX)",
        "capex_lines",
        "Detailed plant and farm investment lines with depreciation lives and VAT timing.",
    ),
    ("Product Pricing Inputs", "revenue_params", "Product pricing, escalation, and indexation."),
    ("Yearly Increment Factors", "yearly_increments", "Apply annual increment percentages and propagate them across price, opex, debt, capex, working capital, and tax."),
    (
        "Production Volumes (Annual)",
        "production_annual",
        "Annual production volumes, availability, and ramp-up assumptions by product.",
    ),
    (
        "Production Volumes (Monthly)",
        "production_monthly",
        "Monthly production schedule by product within the operating window.",
    ),
    ("Operating Costs - Direct", "direct_costs_monthly", "Feedstock and variable operating expenses."),
    ("Operating Costs - Staff", "staff_costs_monthly", "Headcount and payroll assumptions."),
    (
        "Operating Costs - Other Opex",
        "other_opex_monthly",
        "Insurance, services, energy, and overhead costs.",
    ),
    (
        "Working Capital Balances - Receivables",
        "ar_other_assets",
        "Accounts receivable, prepaid expenses, and other current assets.",
    ),
    (
        "Working Capital Balances - Inventory & Payables",
        "inventory_ap",
        "Inventory positions and supplier payables by month.",
    ),
    ("Debt Schedule", "debt_tranches", "Debt facilities with rates, tenors, amortisation, and IDC."),
    ("Tax Schedule", "tax_schedule", "Tax rate, incentives, timing adjustments, and loss carryforwards."),
    ("Inflation Schedule", "inflation_index", "Inflation and FX indexation curves."),
]


HORIZON_SYNC_TABLES: Tuple[str, ...] = (
    "capex_lines",
    "production_monthly",
    "direct_costs_monthly",
    "staff_costs_monthly",
    "other_opex_monthly",
    "ar_other_assets",
    "inventory_ap",
    "inflation_index",
)


YEARLY_INCREMENT_CONFIG = {
    "production_annual": {
        "kind": "production",
        "columns": {
            "annual_volume": "Annual volume annual change (%)",
        },
    },
    "direct_costs_monthly": {
        "kind": "monthly",
        "columns": {
            "unit_price": "Unit price annual change (%)",
            "quantity": "Quantity annual change (%)",
        },
        "reset_amount": True,
    },
    "staff_costs_monthly": {
        "kind": "monthly",
        "columns": {
            "headcount": "Headcount annual change (%)",
            "gross_pay": "Gross pay annual change (%)",
            "benefits": "Benefits annual change (%)",
            "training": "Training annual change (%)",
            "other": "Other staff cost annual change (%)",
        },
        "per_head_map": {
            "gross_pay": "gross_pay_per_head",
            "benefits": "benefits_per_head",
            "training": "training_per_head",
            "other": "other_per_head",
        },
    },
    "other_opex_monthly": {
        "kind": "monthly",
        "columns": {
            "amount": "Other opex annual change (%)",
        },
    },
    "inflation_index": {
        "kind": "monthly",
        "columns": {
            "cpi": "CPI annual change (%)",
            "fx_index": "FX index annual change (%)",
        },
    },
}


def _sync_tables_to_horizon(tables: InputTables, cfg: Dict[str, object]) -> None:
    """Update time-indexed tables in the UI after horizon edits."""

    errors: List[str] = []
    prod_horizon_cfg = cfg.get("production_horizon", {})
    if isinstance(prod_horizon_cfg, Mapping):
        try:
            prod_df = pd.DataFrame(
                [
                    {
                        "start_year": int(prod_horizon_cfg.get("start_year", cfg.get("projection_horizon", {}).get("start_year", DEFAULTS["horizon"]["start_year"]))),
                        "end_year": int(prod_horizon_cfg.get("end_year", cfg.get("projection_horizon", {}).get("end_year", DEFAULTS["horizon"]["end_year"]))),
                    }
                ]
            )
            tables.set_table("production_horizon", prod_df)
            _update_editor_state("production_horizon", tables)
        except Exception as exc:  # pragma: no cover
            errors.append(f"production_horizon: {exc}")
    for table in HORIZON_SYNC_TABLES:
        df = cfg.get(table)
        if not isinstance(df, pd.DataFrame):
            continue
        try:
            tables.set_table(table, df.copy())
        except Exception as exc:  # pragma: no cover - UI validation feedback
            errors.append(f"{table}: {exc}")
        else:
            _update_editor_state(table, tables)
            cfg[table] = tables.ensure_table(table).copy()
    for message in errors:
        st.warning(f"Unable to align {message}")


def _auto_step(value: float) -> float:
    magnitude = abs(float(value))
    if magnitude == 0:
        return 1.0
    step = 10 ** math.floor(math.log10(magnitude))
    return max(step * 0.1, 0.01)


def _infer_date_format(template: object) -> str:
    if isinstance(template, str):
        if re.fullmatch(r"\d{4}-\d{2}-\d{2}", template):
            return "%Y-%m-%d"
        if re.fullmatch(r"\d{4}-\d{2}", template):
            return "%Y-%m"
    return "%Y-%m-%d"


def _format_date_like(date_value: pd.Timestamp, template: object) -> str:
    fmt = _infer_date_format(template)
    return date_value.strftime(fmt)


def _growth_factor(percent: float, offset: int, mode: str) -> float:
    if offset <= 0:
        return 1.0
    pct = float(percent or 0.0)
    if mode == "increase":
        pct = abs(pct)
    elif mode == "decrease":
        pct = -abs(pct)
    elif mode == "copy":
        pct = 0.0
    base = 1.0 + pct / 100.0
    if base <= 0:
        return 0.0
    return base ** offset


def _apply_yearly_increment_monthly(
    df: pd.DataFrame,
    value_columns: Iterable[str],
    percent_map: Dict[str, float],
    timeline: Timeline,
    base_year: int,
    mode: str,
) -> pd.DataFrame:
    if df is None or df.empty:
        raise ValueError("Add at least one row for the selected table before applying yearly increments.")
    if "date" not in df.columns:
        raise ValueError("The selected table does not include a 'date' column.")

    working = df.copy()
    working["__date_ts"] = pd.to_datetime(working["date"], errors="coerce")
    working = working.dropna(subset=["__date_ts"])  # remove rows without valid dates
    if working.empty:
        raise ValueError("No valid dates detected in the current table.")

    target_year = int(base_year)
    available_years = (
        working["__date_ts"].dt.year.dropna().astype(int).sort_values().unique().tolist()
    )
    if target_year not in available_years:
        fallback_year = next((y for y in available_years if y >= target_year), None)
        if fallback_year is None and available_years:
            fallback_year = available_years[-1]
        if fallback_year is None:
            raise ValueError(
                "No dated rows available to seed the yearly increment helper."
            )
        target_year = int(fallback_year)

    base_rows = working[working["__date_ts"].dt.year == target_year].copy()
    if base_rows.empty:
        raise ValueError("No rows found for the selected base year. Update the table and try again.")

    for column in value_columns:
        if column in base_rows.columns:
            base_rows[column] = pd.to_numeric(base_rows[column], errors="coerce").fillna(0.0)

    template_sample = base_rows.iloc[0]["date"]
    years = [year for year in timeline.annual_index() if year >= target_year]
    if not years:
        raise ValueError("Projection horizon does not extend beyond the selected base year.")

    generated_rows: List[Dict[str, object]] = []
    for year in years:
        offset = year - target_year
        for _, row in base_rows.iterrows():
            new_row = row.copy()
            new_date = row["__date_ts"].replace(year=int(year))
            new_row["date"] = _format_date_like(new_date, template_sample)
            for column in value_columns:
                if column not in new_row:
                    continue
                base_value = float(row.get(column, 0.0) or 0.0)
                factor = _growth_factor(percent_map.get(column, 0.0), offset, mode)
                new_row[column] = base_value * factor
            generated_rows.append(new_row)

    result = pd.DataFrame(generated_rows)
    if result.empty:
        raise ValueError("No rows generated from the yearly increment helper.")

    result = result.drop(columns=["__date_ts"], errors="ignore")
    result["__sort"] = pd.to_datetime(result["date"], errors="coerce")
    result = result.sort_values(["__sort"] + [col for col in base_rows.columns if col not in {"__date_ts", "date"}])
    result = result.drop(columns=["__sort"], errors="ignore").reset_index(drop=True)

    # Reorder columns to match the original DataFrame structure
    result = result.reindex(columns=list(df.columns), fill_value=np.nan)
    return result


def _apply_yearly_increment_production(
    df: pd.DataFrame,
    base_values: Dict[str, float],
    percent_map: Dict[str, float],
    timeline: Timeline,
    production_horizon: Mapping[str, object],
    mode: str,
) -> pd.DataFrame:
    if df is None or df.empty:
        raise ValueError("Add at least one product row before applying yearly increments.")

    years = timeline.annual_index()
    if not years:
        raise ValueError("Projection horizon is not defined.")

    prod_start = max(int(production_horizon.get("start_year", years[0])), years[0])
    prod_end = min(int(production_horizon.get("end_year", years[-1])), years[-1])
    if prod_end < prod_start:
        raise ValueError("Production horizon is not aligned with the projection horizon.")

    updated = df.copy()
    for idx, row in updated.iterrows():
        product = str(row.get("product", "")).strip()
        if not product:
            continue
        base_volume = float(base_values.get(product, row.get("annual_volume", 0.0)) or 0.0)
        base_volume = max(base_volume, 0.0)
        percent = float(percent_map.get(product, 0.0) or 0.0)

        volumes_by_year: List[float] = []
        for year in years:
            if year < prod_start or year > prod_end:
                volumes_by_year.append(0.0)
                continue
            offset = year - prod_start
            factor = _growth_factor(percent, offset, mode)
            volumes_by_year.append(base_volume * factor)

        if base_volume <= 0:
            ramp_values = [0.0 for _ in years]
        else:
            ramp_values = []
            for year, volume in zip(years, volumes_by_year):
                if year < prod_start or year > prod_end:
                    ramp_values.append(0.0)
                else:
                    ramp_values.append(volume / base_volume if base_volume else 0.0)

        ramp_str = ";".join(f"{value:.6f}" for value in ramp_values)
        updated.at[idx, "annual_volume"] = base_volume
        updated.at[idx, "startup_ramp"] = ramp_str

    return updated


def _render_yearly_increment_helper(
    tables: InputTables,
    table_name: str,
    df: pd.DataFrame,
    timeline: Timeline,
    production_horizon: Mapping[str, object],
) -> None:
    config = YEARLY_INCREMENT_CONFIG.get(table_name)
    if not config:
        return

    helper_key = f"{table_name}_yearly_helper"
    with st.expander("Yearly increment helper", expanded=False):
        st.caption(
            "Fill the first year's values and specify the annual percentage change to populate future years automatically."
        )

        if config["kind"] == "production":
            if df.empty:
                st.info("Add production rows before applying yearly changes.")
                return

            years = timeline.annual_index()
            if not years:
                st.warning("Projection horizon years are not available.")
                return

            prod_start = max(int(production_horizon.get("start_year", years[0])), years[0])
            prod_end = min(int(production_horizon.get("end_year", years[-1])), years[-1])
            st.write(
                f"First production year: **{prod_start}**. Values outside {prod_start}–{prod_end} are set to zero."
            )

            base_values: Dict[str, float] = {}
            percent_map: Dict[str, float] = {}
            controls = []
            for _, row in df.iterrows():
                product = str(row.get("product", "")).strip()
                if not product:
                    continue
                ramp = parse_ramp(row.get("startup_ramp"), len(years))
                base_index = years.index(prod_start) if prod_start in years else 0
                base_actual = float(row.get("annual_volume", 0.0) or 0.0)
                if ramp:
                    base_actual *= float(ramp[base_index])
                base_actual = max(base_actual, 0.0)

                base_key = f"{helper_key}_{product}_base"
                change_key = f"{helper_key}_{product}_pct"

                base_input = st.number_input(
                    f"{product.replace('_', ' ').title()} first production year volume",
                    min_value=0.0,
                    value=base_actual,
                    step=_auto_step(base_actual),
                    format="%.4f",
                    key=base_key,
                )
                pct_input = st.number_input(
                    f"{product.replace('_', ' ').title()} annual change (%)",
                    value=0.0,
                    step=0.1,
                    format="%.4f",
                    key=change_key,
                )
                base_values[product] = float(base_input)
                percent_map[product] = float(pct_input)
                controls.append(product)

            if not controls:
                st.info("No product rows detected. Add products to the Production Volumes table first.")
                return

            action_cols = st.columns(3)
            copy_clicked = action_cols[0].button("Copy forward", key=f"{helper_key}_copy")
            inc_clicked = action_cols[1].button("Apply increases", key=f"{helper_key}_inc")
            dec_clicked = action_cols[2].button("Apply decreases", key=f"{helper_key}_dec")

            mode = None
            if copy_clicked:
                mode = "copy"
            elif inc_clicked:
                mode = "increase"
            elif dec_clicked:
                mode = "decrease"

            if mode:
                try:
                    updated_df = _apply_yearly_increment_production(
                        df,
                        base_values,
                        percent_map,
                        timeline,
                        production_horizon,
                        mode,
                    )
                except ValueError as exc:
                    st.error(f"Unable to apply yearly increments: {exc}")
                else:
                    try:
                        tables.set_table(table_name, updated_df)
                    except Exception as exc:
                        st.error(f"Unable to update production table: {exc}")
                    else:
                        st.success("Production ramp updated from yearly changes.")
                        _update_editor_state(table_name, tables)
                        _safe_rerun()

        else:  # monthly-style tables
            if df.empty:
                st.info("Add at least one row before applying yearly changes.")
                return

            years = timeline.annual_index()
            if not years:
                st.warning("Projection horizon years are not available.")
                return

            base_year = st.number_input(
                "Base year",
                min_value=int(years[0]),
                max_value=int(years[-1]),
                value=int(years[0]),
                step=1,
                key=f"{helper_key}_base_year",
            )

            percent_map = {
                column: float(
                    st.number_input(
                        label,
                        value=0.0,
                        step=0.1,
                        format="%.4f",
                        key=f"{helper_key}_{column}_pct",
                    )
                )
                for column, label in config["columns"].items()
            }

            action_cols = st.columns(3)
            copy_clicked = action_cols[0].button("Copy forward", key=f"{helper_key}_copy")
            inc_clicked = action_cols[1].button("Apply increases", key=f"{helper_key}_inc")
            dec_clicked = action_cols[2].button("Apply decreases", key=f"{helper_key}_dec")

            mode = None
            if copy_clicked:
                mode = "copy"
            elif inc_clicked:
                mode = "increase"
            elif dec_clicked:
                mode = "decrease"

            if mode:
                try:
                    updated_df = _apply_yearly_increment_monthly(
                        df,
                        config["columns"].keys(),
                        percent_map,
                        timeline,
                        int(base_year),
                        mode,
                    )
                except ValueError as exc:
                    st.error(f"Unable to apply yearly increments: {exc}")
                else:
                    if config.get("reset_amount") and "amount" in updated_df.columns:
                        updated_df["amount"] = np.nan

                    if table_name == "staff_costs_monthly":
                        headcount = pd.to_numeric(updated_df.get("headcount"), errors="coerce").fillna(0.0)
                        updated_df["headcount"] = headcount
                        per_head_map = config.get("per_head_map", {})
                        for total_col, per_col in per_head_map.items():
                            if total_col in updated_df.columns:
                                updated_df[total_col] = pd.to_numeric(
                                    updated_df[total_col], errors="coerce"
                                ).fillna(0.0)
                                if per_col in updated_df.columns:
                                    updated_df[per_col] = np.where(
                                        headcount > 0,
                                        updated_df[total_col] / headcount,
                                        0.0,
                                    )

                    try:
                        tables.set_table(table_name, updated_df)
                    except Exception as exc:
                        st.error(f"Unable to update table: {exc}")
                    else:
                        st.success("Yearly changes applied to the table.")
                        _update_editor_state(table_name, tables)
                        _safe_rerun()
def _render_table_editor(
    tables: InputTables,
    table_name: str,
    label: str,
    error_message: Optional[str] = None,
    description: Optional[str] = None,
    helper: Optional[Callable[[InputTables, str, pd.DataFrame], None]] = None,
    column_config: Optional[Dict[str, object]] = None,
) -> None:
    st.markdown(f"#### {label}")
    if description:
        st.caption(description)
    schema = INPUT_SCHEMAS[table_name]
    df = tables.ensure_table(table_name).copy()

    if helper is not None:
        helper(tables, table_name, df)
        df = tables.ensure_table(table_name).copy()

    feedback_key = f"default_feedback_{table_name}"
    feedback_message = st.session_state.pop(feedback_key, None)
    if feedback_message:
        st.success(feedback_message)
    controls = st.columns(2)
    add_modal_key = f"show_add_modal_{table_name}"
    if controls[0].button(f"Add row", key=f"add_{table_name}"):
        st.session_state[add_modal_key] = True

    if st.session_state.get(add_modal_key):
        modal_title = f"Add {label} row"
        with _modal_container(modal_title):
            schema_columns = list(schema.columns.items())
            with st.form(f"add_row_form_{table_name}"):
                form_values: Dict[str, object] = {}
                for col_name, dtype in schema_columns:
                    pretty_label = col_name.replace("_", " ").title()
                    default_value = schema.defaults.get(col_name, np.nan)
                    if dtype == "bool":
                        form_values[col_name] = st.checkbox(
                            pretty_label,
                            value=_format_default_for_entry(default_value, dtype),
                            key=f"add_row_{table_name}_{col_name}",
                        )
                    else:
                        form_values[col_name] = st.text_input(
                            pretty_label,
                            value=_format_default_for_entry(default_value, dtype),
                            key=f"add_row_{table_name}_{col_name}",
                        )

                share_amount_key: Optional[str] = None
                share_amount_value: Optional[str] = None
                if table_name == "debt_tranches":
                    share_amount_key = f"add_row_{table_name}_share_amount"
                    total_capex = _total_capex_from_inputs(tables)
                    share_amount_help = (
                        "Optional: enter the tranche debt amount to derive the share automatically. "
                        f"Current total CAPEX = {total_capex:,.2f}."
                        if total_capex > 0
                        else "Enter the tranche share directly or populate CAPEX to derive it from an amount."
                    )
                    share_amount_value = st.text_input(
                        "Debt amount (optional)",
                        value="",
                        help=share_amount_help,
                        key=share_amount_key,
                    )

                submit_col, cancel_col = st.columns(2)
                submitted = submit_col.form_submit_button("Save row", use_container_width=True)
                cancelled = cancel_col.form_submit_button("Cancel", use_container_width=True, type="secondary")

            if cancelled:
                st.session_state.pop(add_modal_key, None)
                for col_name, _ in schema_columns:
                    st.session_state.pop(f"add_row_{table_name}_{col_name}", None)
                if share_amount_key is not None:
                    st.session_state.pop(share_amount_key, None)
                _safe_rerun()
            elif submitted:
                try:
                    row_payload: Dict[str, object] = {}
                    for col_name, dtype in schema_columns:
                        widget_value = form_values[col_name]
                        row_payload[col_name] = _parse_modal_entry(widget_value, dtype)
                    if table_name == "debt_tranches" and share_amount_key is not None:
                        share_amount_raw = (share_amount_value or "").strip()
                        if share_amount_raw:
                            amount_value = _parse_modal_entry(share_amount_raw, "float")
                            total_capex = _total_capex_from_inputs(tables)
                            if total_capex <= 0:
                                raise ValueError(
                                    "Populate the CAPEX schedule before deriving a debt amount."
                                )
                            row_payload["share"] = float(amount_value) / float(total_capex)
                    tables.add_row(table_name, row_payload)
                except Exception as exc:
                    st.error(f"Unable to add row: {exc}")
                else:
                    st.session_state[feedback_key] = (
                        "Row added successfully. You can fine-tune the values directly "
                        "in the table below."
                    )
                    st.session_state.pop(add_modal_key, None)
                    for col_name, _ in schema_columns:
                        st.session_state.pop(f"add_row_{table_name}_{col_name}", None)
                    if share_amount_key is not None:
                        st.session_state.pop(share_amount_key, None)
                    _sync_table_mutation_state(table_name, tables)
                    _safe_rerun()

    if not df.empty:
        remove_idx = controls[1].selectbox(
            "Row to remove",
            options=list(df.index),
            format_func=lambda idx: f"Row {idx + 1}",
            key=f"remove_select_{table_name}",
        )
        if controls[1].button("Remove selected", key=f"remove_{table_name}"):
            try:
                tables.remove_row(table_name, int(remove_idx))
            except Exception as exc:  # pragma: no cover - defensive feedback
                st.error(f"Unable to remove row: {exc}")
            df = tables.ensure_table(table_name).copy()
            _sync_table_mutation_state(table_name, tables)
            _safe_rerun()
    state_key = _editor_state_key(table_name)
    # Clear any legacy widget state that may have been set by previous builds
    # using the old key to avoid Streamlit policy violations on render.
    st.session_state.pop(f"editor_{table_name}", None)

    display_df = df.copy()
    effective_column_config: Dict[str, object] = dict(column_config) if column_config else {}

    if table_name == "debt_tranches":
        total_capex = _total_capex_from_inputs(tables)
        if "share" in display_df.columns:
            share_series = pd.to_numeric(display_df["share"], errors="coerce").fillna(0.0)
        else:
            share_series = pd.Series(0.0, index=display_df.index)
        display_df["share_amount"] = share_series * total_capex
        if "share" in display_df.columns and "share_amount" in display_df.columns:
            columns_order = list(display_df.columns)
            share_idx = columns_order.index("share")
            share_amount_idx = columns_order.index("share_amount")
            if share_amount_idx != share_idx + 1:
                columns_order.insert(share_idx + 1, columns_order.pop(share_amount_idx))
                display_df = display_df[columns_order]
        share_help = (
            "Calculated debt amount based on the tranche share multiplied by the "
            f"current total CAPEX ({total_capex:,.2f})."
        )
        effective_column_config.setdefault(
            "share",
            st.column_config.NumberColumn(
                "Share",
                min_value=0.0,
                max_value=1.0,
                format="%.4f",
            ),
        )
        effective_column_config.setdefault(
            "start_year",
            st.column_config.NumberColumn(
                "Start year",
                step=1,
                format="%d",
            ),
        )
        effective_column_config["share_amount"] = st.column_config.NumberColumn(
            "Debt amount",
            help=share_help,
            format="%.2f",
            disabled=True,
        )

    if table_name == "monte_carlo_settings" and "variable" in display_df.columns:
        display_df["variable"] = display_df["variable"].map(MONTE_CARLO_VARIABLE_LABELS).fillna(
            display_df["variable"].astype(str)
        )

    editor = st.data_editor(
        display_df,
        num_rows="dynamic",
        use_container_width=True,
        key=state_key,
        column_config=effective_column_config if effective_column_config else column_config,
    )
    if isinstance(editor, pd.DataFrame):
        editor_clean = editor.copy()
        if "_index" in editor_clean.columns:
            editor_clean = editor_clean.drop(columns=["_index"])

        drop_columns: List[str] = []
        rename_map: Dict[str, str] = {}
        for col in list(editor_clean.columns):
            if col in schema.columns:
                continue
            norm = normalize_key(col)
            if norm in schema.columns:
                rename_map[col] = norm
            else:
                drop_columns.append(col)
        if drop_columns:
            editor_clean = editor_clean.drop(columns=drop_columns)
        if rename_map:
            editor_clean = editor_clean.rename(columns=rename_map)

        if table_name == "monte_carlo_settings" and "variable" in editor_clean.columns:
            editor_clean["variable"] = editor_clean["variable"].map(MONTE_CARLO_VARIABLE_LABEL_TO_KEY).fillna(
                editor_clean["variable"]
            )

        canonical_cols = list(schema.columns.keys())
        editor_clean = editor_clean.reindex(columns=canonical_cols, fill_value=np.nan)
        editor_clean = editor_clean.replace({None: np.nan})
        editor_clean = editor_clean.dropna(how="all").reset_index(drop=True)

        try:
            assert_frame_equal(
                df.reset_index(drop=True),
                editor_clean.reset_index(drop=True),
                check_dtype=False,
                check_like=True,
            )
            frames_equal = True
        except AssertionError:
            frames_equal = False

        if not frames_equal:
            try:
                tables.set_table(table_name, editor_clean)
            except Exception as exc:
                st.error(f"Unable to update table: {exc}")
            else:
                _sync_table_mutation_state(table_name, tables)
                _safe_rerun()

    current_df = tables.ensure_table(table_name).copy()

    with st.expander("Manage defaults & clean start", expanded=False):
        st.caption(
            "Reset this table to the stored defaults or update the baseline values used when creating a clean workbook."
        )
        action_cols = st.columns(3)
        if action_cols[0].button("Reset table to defaults", key=f"reset_defaults_{table_name}"):
            try:
                tables.set_table(table_name, _get_default_table(table_name))
            except Exception as exc:
                st.error(f"Unable to reset table: {exc}")
            else:
                st.session_state[feedback_key] = "Table reset to stored defaults."
                _sync_table_mutation_state(table_name, tables)
                _safe_rerun()
        if action_cols[1].button("Save current as defaults", key=f"save_defaults_{table_name}"):
            try:
                _save_default_table(table_name, current_df)
            except Exception as exc:
                st.error(f"Unable to save defaults: {exc}")
            else:
                st.session_state[feedback_key] = "Stored defaults updated from current table."
                _sync_table_mutation_state(table_name, tables)
                _safe_rerun()
        if action_cols[2].button("Restore factory defaults", key=f"factory_defaults_{table_name}"):
            try:
                restored_df = _restore_factory_default_table(table_name)
                tables.set_table(table_name, restored_df)
            except Exception as exc:
                st.error(f"Unable to restore factory defaults: {exc}")
            else:
                st.session_state[feedback_key] = "Factory defaults restored and applied."
                _sync_table_mutation_state(table_name, tables)
                _safe_rerun()

        defaults_df = _get_default_table(table_name)
        _render_default_row_controls(table_name, label, schema, defaults_df)

    _render_default_edit_modal(table_name, label, schema, tables)

    if error_message:
        st.error(f"Validation error: {error_message}")
    st.divider()


def main() -> None:
    try:
        if _streamlit_runtime_exists():
            st.set_page_config(title="Sugarcane Bioethanol Finance Model", layout="wide")
    except (StreamlitAPIException, RuntimeError, Exception):  # pragma: no cover - defensive guard
        # Some Streamlit versions raise a generic Exception when page config is invoked
        # outside a live runtime; swallow and continue so local execution still works.
        pass

    # Ensure the hero title has sufficient breathing room while keeping the layout
    # nearly full-width on large monitors.
    st.markdown(
        """
        <style>
        html, body, .stApp {
            margin: 0 !important;
            padding: 0 !important;
            width: 100% !important;
            max-width: 100% !important;
            overflow-x: hidden !important;
        }
        [data-testid="stAppViewContainer"] {
            margin: 0 !important;
            padding: 3.5rem 0 2.5rem !important;
            width: 100% !important;
            max-width: 100% !important;
        }
        [data-testid="stAppViewContainer"] > .main {
            margin: 0 auto !important;
            width: 100% !important;
            max-width: 100% !important;
        }
        [data-testid="stAppViewContainer"] .main .block-container,
        .block-container {
            margin: 0 auto !important;
            width: 100% !important;
            max-width: none !important;
            padding-left: clamp(20px, 3vw, 48px) !important;
            padding-right: clamp(20px, 3vw, 48px) !important;
        }
        [data-testid="stVerticalBlock"],
        [data-testid="stHorizontalBlock"] {
            margin-left: 0 !important;
            margin-right: 0 !important;
            padding-left: 0 !important;
            padding-right: 0 !important;
            width: 100% !important;
            max-width: 100% !important;
        }
        h1, h1 span, .stMarkdown h1 {
            font-size: clamp(2.2rem, 3vw, 2.8rem) !important;
            line-height: 1.2 !important;
            margin-top: 0 !important;
            margin-bottom: 1.2rem !important;
            white-space: normal !important;
            overflow-wrap: anywhere !important;
        }
        </style>
        """,
        unsafe_allow_html=True,
    )

    if MATPLOTLIB_INSTALL_ERROR and plt is None:
        st.info(
            "Matplotlib charts will use Streamlit fallbacks until the library is "
            "installed. Last attempt: "
            f"{MATPLOTLIB_INSTALL_ERROR}",
            icon="ℹ️",
        )

    if MODEL_IMPORT_ERROR is not None:
        st.error(
            "Required dependency missing when importing the finance engine: "
            f"{MODEL_IMPORT_ERROR}. Install the project requirements (numpy, pandas, etc.) "
            "and restart the app."
        )
        st.stop()

    st.title("Sugarcane Bioethanol Project Finance Model")
    st.markdown(
        "Use this Streamlit interface to explore the integrated bioethanol, sugar, "
        "electricity, and animal feed project finance model. Adjust critical drivers "
        "in the control tabs above and review the resulting statements, dashboards, "
        "sensitivities, and scenarios."
    )

    page_tabs_container = st.container()
    (
        model_controls_tab,
        landing_tab,
        summary_tab,
        financial_tab,
        production_tab,
        sensitivity_tab,
        chatbot_tab,
    ) = page_tabs_container.tabs(
        [
            "Model Controls",
            "Input & Assumptions",
            "Summary",
            "Financial Statements",
            "Production & Pricing",
            "Sensitivities",
            "AI Chatbot",
        ]
    )

    tables = _get_tables()
    sync_errors = _sync_tables_from_state(tables)
    assumptions: Dict[str, object] = {}
    cfg = build_config(assumptions, tables)
    st.session_state.pop("scenario_payload_cache", None)

    horizon = cfg["projection_horizon"]
    production_horizon = {"start_year": horizon["start_year"], "end_year": horizon["end_year"]}

    with model_controls_tab:
        st.subheader("Model Controls")
        control_tabs = st.tabs(
            [
                "Projection",
                "Financial",
                "Production",
                "Pricing",
                "Risk & Scenarios",
            ]
        )

        with control_tabs[0]:
            st.markdown("### Projection horizon")
            start_year = st.number_input("Start year", value=int(horizon["start_year"]), step=1)
            end_year = st.number_input(
                "End year",
                value=int(horizon["end_year"]),
                min_value=int(start_year),
                step=1,
            )
            start_month = st.number_input(
                "Start month",
                min_value=1,
                max_value=12,
                value=int(horizon.get("start_month", 1)),
            )
            horizon.update(
                {
                    "start_year": int(start_year),
                    "end_year": int(end_year),
                    "start_month": int(start_month),
                }
            )
            production_horizon.update({"start_year": int(start_year), "end_year": int(end_year)})
            cfg["production_horizon"] = production_horizon
            st.caption("Production horizon is automatically aligned to the projection horizon.")
            try:
                tables.set_table("projection_horizon", pd.DataFrame([horizon]))
                _update_editor_state("projection_horizon", tables)
                tables.set_table("production_horizon", pd.DataFrame([production_horizon]))
                _update_editor_state("production_horizon", tables)
            except Exception as exc:
                st.warning(f"Projection inputs not saved due to validation error: {exc}")
            else:
                cfg = align_with_projection_horizon(cfg)
                _sync_tables_to_horizon(tables, cfg)

        global_inputs = cfg["global_inputs"]
        with control_tabs[1]:
            st.markdown("### Financial assumptions")
            discount_rate = st.number_input(
                "Discount rate (WACC)",
                min_value=0.0,
                max_value=1.0,
                value=float(global_inputs.get("discount_rate", 0.12)),
                step=0.005,
                format="%.4f",
            )
            corp_tax = st.number_input(
                "Corporate tax rate",
                min_value=0.0,
                max_value=1.0,
                value=float(global_inputs.get("corp_tax_rate", 0.28)),
                step=0.01,
                format="%.4f",
            )
            investor_options = _float_option_values(0.0, 1.0, 0.05, digits=2)
            investor_share_default = float(global_inputs.get("investor_share", 0.6))
            investor_share = st.selectbox(
                "Investor equity share",
                investor_options,
                index=min(range(len(investor_options)), key=lambda idx: abs(investor_options[idx] - investor_share_default)),
                format_func=lambda value: f"{value:.0%}",
            )
            global_inputs.update(
                {
                    "discount_rate": float(discount_rate),
                    "corp_tax_rate": float(corp_tax),
                    "investor_share": float(investor_share),
                    "owner_share": float(1.0 - investor_share),
                }
            )
            st.caption(f"Owner equity share automatically set to {1.0 - investor_share:.2f}")
            try:
                tables.set_table("global_inputs", pd.DataFrame([global_inputs]))
                _update_editor_state("global_inputs", tables)
                tables.set_table("working_capital_days", pd.DataFrame([cfg["working_capital"]]))
                _update_editor_state("working_capital_days", tables)
            except Exception as exc:
                st.warning(f"Global inputs not saved due to validation error: {exc}")

        production_cfg = cfg.setdefault("production", {})
        with control_tabs[2]:
            st.markdown("### Production assumptions")
            feedstock = st.number_input(
                "Annual feedstock (t)",
                min_value=0.0,
                value=float(production_cfg.get("annual_feedstock_ton", 100_000.0)),
                step=10_000.0,
                format="%.0f",
            )
            availability_options = _float_option_values(0.5, 1.0, 0.01, digits=2)
            availability_default = float(production_cfg.get("plant_availability", 0.9))
            availability = st.selectbox(
                "Plant availability",
                availability_options,
                index=min(range(len(availability_options)), key=lambda idx: abs(availability_options[idx] - availability_default)),
                format_func=lambda value: f"{value:.0%}",
            )
            loss_factor_options = _float_option_values(0.0, 0.2, 0.005, digits=3)
            loss_factor_default = float(production_cfg.get("loss_factor", 0.02))
            loss_factor = st.selectbox(
                "Process loss factor",
                loss_factor_options,
                index=min(range(len(loss_factor_options)), key=lambda idx: abs(loss_factor_options[idx] - loss_factor_default)),
                format_func=lambda value: f"{value:.1%}",
            )
            production_cfg.update(
                {
                    "annual_feedstock_ton": float(feedstock),
                    "plant_availability": float(availability),
                    "loss_factor": float(loss_factor),
                }
            )
            scenario = st.selectbox(
                "Feedstock sourcing scenario",
                FEEDSTOCK_SCENARIOS,
                index=FEEDSTOCK_SCENARIOS.index(production_cfg.get("feedstock_scenario", "HYBRID")),
            )
            production_cfg["feedstock_scenario"] = scenario
            if scenario == "HYBRID":
                farm_share_options = _float_option_values(0.0, 1.0, 0.05, digits=2)
                farm_share_default = float(production_cfg.get("farm_share", 0.5))
                farm_share = st.selectbox(
                    "Hybrid farm share",
                    farm_share_options,
                    index=min(range(len(farm_share_options)), key=lambda idx: abs(farm_share_options[idx] - farm_share_default)),
                    format_func=lambda value: f"{value:.0%}",
                )
                production_cfg["farm_share"] = float(farm_share)

        pricing_cfg = cfg.setdefault("prices", {})
        with control_tabs[3]:
            st.markdown("### Product pricing")
            for product in PRODUCTS:
                params = pricing_cfg.setdefault(product, {})
                base_price = st.number_input(
                    f"{product.replace('_', ' ').title()} base price",
                    min_value=0.0,
                    value=float(params.get("base_price", 0.0 if product != "ethanol" else 0.7)),
                    step=1.0 if product != "ethanol" else 0.05,
                    format="%.4f",
                )
                escalation = st.number_input(
                    f"{product.replace('_', ' ').title()} price escalation (pa)",
                    min_value=0.0,
                    max_value=0.3,
                    value=float(params.get("price_escalation_pa", 0.02)),
                    step=0.005,
                    format="%.4f",
                )
                params.update({"base_price": float(base_price), "price_escalation_pa": float(escalation)})
            revenue_table = tables.ensure_table("revenue_params").copy()
            if revenue_table.empty:
                revenue_table = pd.DataFrame(
                    columns=list(INPUT_SCHEMAS["revenue_params"].columns.keys())
                )
            for product in PRODUCTS:
                params = pricing_cfg.get(product, {})
                mask = (
                    revenue_table["product"].astype(str).str.lower() == product
                    if "product" in revenue_table
                    else pd.Series(dtype=bool)
                )
                updated = {
                    "product": product,
                    "base_price": params.get("base_price", np.nan),
                    "price_escalation_pa": params.get("price_escalation_pa", np.nan),
                    "price_indexation": params.get("price_indexation", "cpi"),
                    "uom": params.get("uom", ""),
                    "tariff_structure": params.get("tariff_structure", ""),
                    "revenue_share": params.get("revenue_share", 1.0),
                }
                if mask.any():
                    for key, value in updated.items():
                        revenue_table.loc[mask, key] = value
                else:
                    revenue_table = pd.concat(
                        [revenue_table, pd.DataFrame([updated])],
                        ignore_index=True,
                    )
            try:
                tables.set_table("revenue_params", revenue_table)
            except Exception as exc:
                st.warning(f"Pricing table not saved due to validation error: {exc}")
            else:
                _update_editor_state("revenue_params", tables)

        with control_tabs[4]:
            st.markdown("### Risk and scenario options")
            risk_targets, risk_applies = _get_risk_select_options(tables)
            risk_column_config = {
                "distribution": st.column_config.SelectboxColumn(
                    "Distribution",
                    options=list(RISK_DISTRIBUTION_OPTIONS),
                    help="Probability distribution used when sampling this risk driver.",
                ),
                "target": st.column_config.SelectboxColumn(
                    "Target",
                    options=risk_targets,
                    help="Choose whether the draw scales general risk multipliers, prices, availability, opex, or capex.",
                ),
                "applies_to": st.column_config.SelectboxColumn(
                    "Applies to",
                    options=risk_applies,
                    help="Restrict the driver to a specific product/segment or leave as global.",
                ),
                "p1": st.column_config.NumberColumn("P1", help="Distribution parameter 1 (e.g., mean or minimum)."),
                "p2": st.column_config.NumberColumn("P2", help="Distribution parameter 2 (e.g., stdev or mode)."),
                "p3": st.column_config.NumberColumn("P3", help="Distribution parameter 3 (e.g., triangular maximum)."),
                "production_multiplier": st.column_config.NumberColumn(
                    "Production multiplier",
                    min_value=0.0,
                    help="Baseline production multiplier applied before any sampled draw.",
                ),
                "labour_multiplier": st.column_config.NumberColumn(
                    "Labour multiplier",
                    min_value=0.0,
                    help="Baseline labour cost multiplier applied before any sampled draw.",
                ),
                "price_multiplier": st.column_config.NumberColumn(
                    "Price multiplier",
                    min_value=0.0,
                    help="Baseline price multiplier applied before any sampled draw.",
                ),
                "revenue_multiplier": st.column_config.NumberColumn(
                    "Revenue multiplier",
                    min_value=0.0,
                    help="Baseline revenue multiplier applied before any sampled draw.",
                ),
                "yield_multiplier": st.column_config.NumberColumn(
                    "Yield multiplier",
                    min_value=0.0,
                    help="Baseline yield multiplier applied before any sampled draw.",
                ),
            }
            _render_table_editor(
                tables,
                "risk_params",
                "Risk Schedule",
                sync_errors.get("risk_params"),
                "Political, environmental, and market risk multipliers for production, pricing, and labour assumptions.",
                column_config=risk_column_config,
            )
            _render_risk_schedule_preview(tables)

    timeline = Timeline(
        int(horizon["start_year"]),
        int(horizon["end_year"]),
        int(horizon.get("start_month", 1)),
    )

    with st.spinner("Running base model..."):
        try:
            results = run_full_model(cfg)
        except Exception as exc:  # pragma: no cover - runtime feedback for the UI
            st.error(f"Model execution failed: {exc}")
            st.stop()

    metrics = results["metrics"]
    dashboard = results["dashboard"]

    metric_items = [
        ("Project_NPV", "Project NPV", "currency"),
        ("Project_IRR", "Project IRR", "percent"),
        ("Equity_IRR", "Equity IRR", "percent"),
        ("Payback_Year", "Payback Year", "year"),
        ("DSCR_min", "Min DSCR", "ratio"),
        ("DSCR_avg", "Avg DSCR", "ratio"),
    ]

    scenario_table_for_download = tables.ensure_table("scenario_comparison").copy()
    scenario_overrides = _scenario_overrides_from_table(scenario_table_for_download)
    scenario_options: List[str] = [BASE_SCENARIO_LABEL, *list(scenario_overrides.keys())]


    with landing_tab:
        top_left, top_right = st.columns([3, 2])
        with top_left:
            st.subheader("Input & assumptions tables")
            st.markdown(
                "Review, add, or remove records from each canonical input table. Updates apply across the model "
                "on the next run."
            )
        with top_right:
            st.markdown("#### Excel model download")
            if not scenario_options:
                scenario_options = [BASE_SCENARIO_LABEL]
            default_option = st.session_state.get("excel_download_selected", scenario_options[0])
            if default_option not in scenario_options:
                default_option = scenario_options[0]
            selected_scenario = st.selectbox(
                "Scenario",
                scenario_options,
                index=scenario_options.index(default_option),
                key="excel_download_scenario",
            )
            st.session_state["excel_download_selected"] = selected_scenario

            download_container = st.container()
            excel_map: Dict[str, bytes] = st.session_state.setdefault("excel_bytes_map", {})
            stale_keys = [key for key in excel_map if key not in scenario_options]
            for key in stale_keys:
                excel_map.pop(key, None)
            st.session_state.excel_bytes_map = excel_map

            scenario_cfg_payload, scenario_results_payload = _ensure_scenario_payload(
                selected_scenario,
                cfg,
                results,
                scenario_overrides,
            )
            cfg_for_excel = copy.deepcopy(scenario_cfg_payload)
            metadata = cfg_for_excel.setdefault("metadata", {}) if isinstance(cfg_for_excel, dict) else {}
            if isinstance(metadata, dict):
                metadata["scenario"] = selected_scenario
            st.session_state.model_results = (cfg_for_excel, scenario_results_payload)

            excel_bytes = excel_map.get(selected_scenario)

            with download_container:
                if not excel_bytes:
                    if st.button(
                        "Prepare Excel Model",
                        key=f"prepare_excel_{normalize_key(selected_scenario) or 'base'}",
                    ):
                        with st.spinner("Preparing Excel workbook..."):
                            try:
                                excel_bytes = _generate_excel_bytes(
                                    cfg_for_excel,
                                    scenario_results_payload,
                                    selected_scenario,
                                )
                            except RuntimeError as exc:
                                st.error(str(exc))
                                excel_bytes = None
                            else:
                                excel_map[selected_scenario] = excel_bytes
                                st.session_state.excel_bytes_map = excel_map
                if excel_bytes:
                    file_scenario = normalize_key(selected_scenario) or "base"
                    st.download_button(
                        "Download Excel Model",
                        data=excel_bytes,
                        file_name=f"Sugarcane_Financial_Model_{file_scenario}.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        key=f"download_excel_{file_scenario}",
                    )
                    if st.button(
                        "Clear Prepared Excel",
                        key=f"clear_excel_{normalize_key(selected_scenario) or 'base'}",
                    ):
                        excel_map.pop(selected_scenario, None)
                        st.session_state.excel_bytes_map = excel_map
                        excel_bytes = None
                if not excel_bytes:
                    st.info("Click 'Prepare Excel Model' to generate the workbook for download.")

        st.markdown("### Horizon overview")
        _render_horizon_timeline_chart(horizon)

        st.markdown("### Assumption-driven visuals")
        st.caption("Charts below reflect the current input schedules and help validate annualised assumptions.")

        prod_annual_df = results.get("production_annual") if isinstance(results, Mapping) else None
        if isinstance(prod_annual_df, pd.DataFrame) and not prod_annual_df.empty:
            prod_chart = (
                prod_annual_df.pivot(index="year", columns="product", values="volume").fillna(0.0).reset_index()
            )
            prod_chart.rename(columns={"year": "Year"}, inplace=True)
            value_cols = [col for col in prod_chart.columns if col != "Year"]
            if value_cols:
                _render_stacked_columns(prod_chart, "Year", value_cols, "Annual production by product", ylabel="Volume")

        direct_input = tables.ensure_table("direct_costs_monthly").copy()
        if not direct_input.empty and "amount" in direct_input.columns:
            direct_input["date"] = pd.to_datetime(direct_input["date"], errors="coerce")
            direct_input = direct_input.dropna(subset=["date"])
            if not direct_input.empty:
                direct_input["year"] = direct_input["date"].dt.year
                direct_input["cost_type"] = direct_input.get("cost_type", "").astype(str).replace("", "Unspecified")
                direct_year = (
                    direct_input.groupby(["year", "cost_type"], dropna=False)["amount"].sum().unstack(fill_value=0.0).reset_index()
                )
                direct_year.rename(columns={"year": "Year"}, inplace=True)
                value_cols = [col for col in direct_year.columns if col != "Year"]
                if value_cols:
                    _render_stacked_columns(
                        direct_year,
                        "Year",
                        value_cols,
                        "Annual direct costs by type",
                        ylabel="Amount",
                    )

        staff_input = tables.ensure_table("staff_costs_monthly").copy()
        if not staff_input.empty:
            staff_input["date"] = pd.to_datetime(staff_input["date"], errors="coerce")
            staff_input = staff_input.dropna(subset=["date"])
            if not staff_input.empty:
                staff_input["year"] = staff_input["date"].dt.year
                staff_year = (
                    staff_input.groupby("year")[[col for col in ("gross_pay", "benefits", "training", "other") if col in staff_input.columns]]
                    .sum()
                    .reset_index()
                )
                staff_year.rename(columns={"year": "Year"}, inplace=True)
                value_cols = [col for col in staff_year.columns if col != "Year"]
                if value_cols:
                    _render_stacked_columns(
                        staff_year,
                        "Year",
                        value_cols,
                        "Annual staff cost components",
                        ylabel="Amount",
                    )

        other_input = tables.ensure_table("other_opex_monthly").copy()
        if not other_input.empty and "amount" in other_input.columns:
            other_input["date"] = pd.to_datetime(other_input["date"], errors="coerce")
            other_input = other_input.dropna(subset=["date"])
            if not other_input.empty:
                other_input["year"] = other_input["date"].dt.year
                other_input["category"] = other_input.get("category", "").astype(str).replace("", "Unspecified")
                other_year = (
                    other_input.groupby(["year", "category"], dropna=False)["amount"].sum().unstack(fill_value=0.0).reset_index()
                )
                other_year.rename(columns={"year": "Year"}, inplace=True)
                value_cols = [col for col in other_year.columns if col != "Year"]
                if value_cols:
                    _render_stacked_columns(
                        other_year,
                        "Year",
                        value_cols,
                        "Annual other opex by category",
                        ylabel="Amount",
                    )

        inflation_input = tables.ensure_table("inflation_index").copy()
        if not inflation_input.empty:
            inflation_input["date"] = pd.to_datetime(inflation_input["date"], errors="coerce")
            inflation_input = inflation_input.dropna(subset=["date"])
            if not inflation_input.empty:
                inflation_input["year"] = inflation_input["date"].dt.year
                value_cols = [col for col in ("cpi", "fx_index") if col in inflation_input.columns]
                if value_cols:
                    inflation_year = inflation_input.groupby("year")[value_cols].mean().reset_index()
                    inflation_year.rename(columns={"year": "Year"}, inplace=True)
                    _render_stacked_columns(
                        inflation_year,
                        "Year",
                        value_cols,
                        "Average inflation and FX indices",
                        ylabel="Index",
                    )

        if sync_errors:
            for table_name, message in sync_errors.items():
                st.error(f"{table_name}: {message}")
        for label, table_name, description in LANDING_TABLES:
            helper = None
            if table_name in YEARLY_INCREMENT_CONFIG:
                helper = lambda tbls, _name, df, tn=table_name: _render_yearly_increment_helper(
                    tbls,
                    tn,
                    df,
                    timeline,
                    production_horizon,
                )
            _render_table_editor(
                tables,
                table_name,
                label,
                sync_errors.get(table_name),
                description,
                helper=helper,
            )

    with summary_tab:
        st.subheader("Headline metrics")
        for idx in range(0, len(metric_items), 3):
            cols = st.columns(3)
            for col, (key, label, kind) in zip(cols, metric_items[idx: idx + 3]):
                col.metric(label, _format_metric(metrics.get(key), kind))
        st.subheader("Assumptions snapshot")
        st.dataframe(dashboard["assumptions_snapshot"], use_container_width=True)
        st.subheader("Global block")
        st.dataframe(dashboard["global_block"], use_container_width=True)
        latest = dashboard.get("latest_drivers", {})
        if latest:
            st.subheader("Latest drivers")
            latest_df = pd.DataFrame([latest])
            st.dataframe(latest_df, use_container_width=True)
        annual_prod = dashboard.get("annual_production")
        if isinstance(annual_prod, pd.DataFrame) and not annual_prod.empty:
            st.subheader("Annual production by product")
            prod_chart = annual_prod.set_index("year")
            st.bar_chart(prod_chart)
        annual_cashflow = dashboard.get("annual_cashflow")
        if isinstance(annual_cashflow, pd.DataFrame) and not annual_cashflow.empty:
            st.subheader("Annual cash flows")
            cash_chart = annual_cashflow.set_index("year")[["CFO", "CFI", "CFF", "NetCashFlow"]]
            st.bar_chart(cash_chart)

        wc_trend = dashboard.get("working_capital_trend")
        if isinstance(wc_trend, pd.DataFrame) and not wc_trend.empty:
            _render_working_capital_chart(wc_trend)

        dscr_trend = dashboard.get("dscr_trend")
        if isinstance(dscr_trend, pd.DataFrame) and not dscr_trend.empty:
            _render_dscr_trend_chart(dscr_trend)

        debt_summary = dashboard.get("debt_service_summary")
        if isinstance(debt_summary, pd.DataFrame) and not debt_summary.empty:
            _render_debt_waterfall_chart(debt_summary)

        cost_structure_annual = dashboard.get("cost_structure_annual")
        if isinstance(cost_structure_annual, pd.DataFrame) and not cost_structure_annual.empty:
            _render_cost_structure_chart(cost_structure_annual)

        labour_summary = dashboard.get("labour_summary")
        if isinstance(labour_summary, pd.DataFrame) and not labour_summary.empty:
            _render_labour_area_chart(labour_summary)

        break_even_products = dashboard.get("break_even_per_product")
        if isinstance(break_even_products, pd.DataFrame) and not break_even_products.empty:
            _render_break_even_bar_chart(break_even_products)

        cumulative_cf = dashboard.get("cumulative_cashflows")
        if isinstance(cumulative_cf, pd.DataFrame) and not cumulative_cf.empty:
            _render_cumulative_cash_chart(cumulative_cf)

        revenue_vs_production = dashboard.get("revenue_vs_production")
        if isinstance(revenue_vs_production, pd.DataFrame) and not revenue_vs_production.empty:
            _render_revenue_vs_production_scatter(revenue_vs_production)

        staff_monthly = results.get("staff_costs_detail")
        if isinstance(staff_monthly, pd.DataFrame) and not staff_monthly.empty:
            staff_monthly_view = staff_monthly.copy()
            staff_monthly_view["date"] = pd.to_datetime(staff_monthly_view["date"])
            staff_monthly_view = staff_monthly_view.sort_values(["date", "dept"]).reset_index(drop=True)
            _render_dataframe(
                staff_monthly_view,
                "Staff operating cost breakdown (monthly)",
                key="staff_costs_detail_monthly",
            )

        staff_annual = results.get("staff_costs_detail_annual")
        if isinstance(staff_annual, pd.DataFrame) and not staff_annual.empty:
            sort_cols = ["year"]
            if "dept" in staff_annual.columns:
                sort_cols.append("dept")
            if "currency" in staff_annual.columns and "currency" not in sort_cols:
                sort_cols.append("currency")
            staff_annual_view = staff_annual.sort_values(sort_cols).reset_index(drop=True)
            _render_dataframe(
                staff_annual_view,
                "Staff operating cost breakdown (annual)",
                key="staff_costs_detail_annual",
            )

    statement_configs = [
        ("Income Statement (P&L)", "pnl"),
        ("Statement of Cash Flows", "cashflow"),
        ("Statement of Financial Position", "balancesheet"),
    ]
    with financial_tab:
        with st.expander("Where key operating/financing fields appear", expanded=False):
            mapping_df = pd.DataFrame(
                [
                    {"Field": "DirectCosts", "Statement": "Income Statement (P&L)", "Column": "DirectCosts"},
                    {"Field": "StaffCosts", "Statement": "Income Statement (P&L)", "Column": "StaffCosts"},
                    {"Field": "Other Opex", "Statement": "Income Statement (P&L)", "Column": "OtherOpexCosts"},
                    {"Field": "Interest", "Statement": "Income Statement (P&L)", "Column": "Interest"},
                    {"Field": "CFI", "Statement": "Cash Flow", "Column": "CFI"},
                    {"Field": "CFF", "Statement": "Cash Flow", "Column": "CFF"},
                    {"Field": "Inventory", "Statement": "Balance Sheet", "Column": "Inventory"},
                    {"Field": "PPE Gross", "Statement": "Balance Sheet", "Column": "PPE_Gross"},
                    {"Field": "Debt", "Statement": "Balance Sheet", "Column": "Debt"},
                    {"Field": "Accounts Payable", "Statement": "Balance Sheet", "Column": "AccountsPayable"},
                    {"Field": "Total Liabilities", "Statement": "Balance Sheet", "Column": "TotalLiabilities"},
                ]
            )
            st.dataframe(mapping_df, use_container_width=True)

        fs_tabs = st.tabs([label for label, _ in statement_configs])
        for tab, (label, key) in zip(fs_tabs, statement_configs):
            with tab:
                monthly_df = results["statements_monthly"].get(key, pd.DataFrame())
                annual_df = results["statements_annual"].get(key, pd.DataFrame())
                monthly_df = _ensure_statement_columns(key, monthly_df)
                annual_df = _ensure_statement_columns(key, annual_df)
                safe_key = re.sub(r"[^a-z0-9]+", "_", label.lower())
                _render_dataframe(monthly_df, f"Monthly {label}", key=f"monthly_{safe_key}")
                _render_dataframe(annual_df, f"Annual {label}", key=f"annual_{safe_key}")
                if key == "pnl" and isinstance(monthly_df, pd.DataFrame) and not monthly_df.empty:
                    _render_pnl_chart(monthly_df)
                if key == "cashflow" and isinstance(monthly_df, pd.DataFrame) and not monthly_df.empty:
                    _render_cashflow_chart(monthly_df)
                    cumulative_cf = dashboard.get("cumulative_cashflows")
                    if isinstance(cumulative_cf, pd.DataFrame) and not cumulative_cf.empty:
                        _render_cumulative_cash_chart(cumulative_cf)
                if key == "balancesheet" and isinstance(monthly_df, pd.DataFrame) and not monthly_df.empty:
                    _render_balance_sheet_chart(monthly_df)

        with st.expander("IFRS primary statements (SPV + consolidation adjustments)", expanded=False):
            ifrs_configs = [
                ("Statement of Profit or Loss and OCI", "sopl_oci_monthly", "sopl_oci_annual"),
                ("Statement of Financial Position", "sofp_monthly", "sofp_annual"),
                ("Statement of Changes in Equity", "socie_monthly", "socie_annual"),
                ("Statement of Cash Flows (Indirect method)", "scf_indirect_monthly", "scf_indirect_annual"),
            ]
            ifrs_tabs = st.tabs([label for label, _, _ in ifrs_configs])
            for tab, (label, monthly_key, annual_key) in zip(ifrs_tabs, ifrs_configs):
                with tab:
                    monthly_df = results.get(monthly_key, pd.DataFrame())
                    annual_df = results.get(annual_key, pd.DataFrame())
                    safe_key = re.sub(r"[^a-z0-9]+", "_", label.lower())
                    _render_dataframe(monthly_df, f"Monthly {label}", key=f"monthly_ifrs_{safe_key}")
                    _render_dataframe(annual_df, f"Annual {label}", key=f"annual_ifrs_{safe_key}")
            note_configs = [
                ("PPE roll-forward note", "ifrs_note_ppe_rollforward"),
                ("Debt maturity note", "ifrs_note_debt_maturity"),
                ("Working capital bridge note", "ifrs_note_wc_bridge"),
                ("Deferred tax note", "ifrs_note_deferred_tax"),
                ("Lease note (IFRS 16)", "ifrs_note_lease"),
                ("Hedge reserve note (IFRS 9)", "ifrs_note_hedge_reserve"),
            ]
            for title, key in note_configs:
                note_df = results.get(key, pd.DataFrame())
                _render_dataframe(note_df, title, key=f"ifrs_note_{key}")

    with production_tab:
        prod_monthly = results["production_monthly"].copy()
        prod_monthly["date"] = pd.to_datetime(prod_monthly["date"])
        _render_dataframe(prod_monthly, "Monthly production", key="production_monthly")
        prod_annual_src = results.get("production_annual")
        if (
            isinstance(prod_annual_src, pd.DataFrame)
            and {"year", "product", "volume"}.issubset(prod_annual_src.columns)
            and not prod_annual_src.empty
        ):
            prod_annual = (
                prod_annual_src.pivot_table(index="year", columns="product", values="volume", aggfunc="sum")
                .reset_index()
                .fillna(0.0)
            )
        else:
            prod_annual = pd.DataFrame(columns=["year", *PRODUCTS])
        _render_dataframe(prod_annual, "Annual production", key="production_annual")
        if not prod_monthly.empty:
            monthly_chart = (
                prod_monthly.groupby(["date", "product"], as_index=False)["volume"].sum().pivot(
                    index="date", columns="product", values="volume"
                )
            )
            monthly_chart = monthly_chart.fillna(0.0).reset_index()
            value_cols = [col for col in monthly_chart.columns if col != "date"]
            if value_cols:
                _render_line_chart(monthly_chart, "date", value_cols, "Monthly production trends", ylabel="Volume")
        if not prod_annual.empty:
            value_cols = [col for col in prod_annual.columns if col != "year"]
            if value_cols:
                annual_chart = prod_annual.rename(columns={"year": "Year"})
                _render_stacked_columns(annual_chart, "Year", value_cols, "Annual production totals", ylabel="Volume")
        price_curves = results["price_curves"].copy()
        price_curves["date"] = pd.to_datetime(price_curves["date"])
        _render_dataframe(price_curves, "Price curves", key="price_curves")
        if not price_curves.empty:
            price_chart = (
                price_curves.pivot_table(index="date", columns="product", values="price", aggfunc="mean")
                .fillna(0.0)
                .reset_index()
            )
            value_cols = [col for col in price_chart.columns if col != "date"]
            if value_cols:
                _render_line_chart(price_chart, "date", value_cols, "Price curves by product", ylabel="Price")
        revenue_df = results["revenue"].copy()
        revenue_df["date"] = pd.to_datetime(revenue_df["date"])
        _render_dataframe(revenue_df, "Revenue stack", key="revenue")
        if not revenue_df.empty:
            revenue_chart = (
                revenue_df.groupby(["date", "product"], as_index=False)["revenue"].sum().pivot(
                    index="date", columns="product", values="revenue"
                )
            )
            revenue_chart = revenue_chart.fillna(0.0).reset_index()
            value_cols = [col for col in revenue_chart.columns if col != "date"]
            if value_cols:
                _render_area_chart(revenue_chart, "date", value_cols, "Revenue mix by product", ylabel="Revenue")

        st.markdown("### Break-even analysis")
        be_results = results.get("break_even", {})
        per_product_be = (
            be_results.get("per_product")
            if isinstance(be_results, dict)
            and isinstance(be_results.get("per_product"), pd.DataFrame)
            else pd.DataFrame()
        )
        overall_be = be_results.get("overall") if isinstance(be_results, dict) else {}

        break_even_inputs_table = tables.ensure_table("break_even_inputs").copy()
        default_break_even = DEFAULTS["break_even_inputs"].copy()
        default_break_even["product"] = (
            default_break_even["product"].astype(str).str.strip().str.lower()
        )

        if not break_even_inputs_table.empty:
            break_even_inputs_table["product"] = (
                break_even_inputs_table["product"].astype(str).str.strip().str.lower()
            )
            if "product" in break_even_inputs_table.columns:
                break_even_inputs_table = (
                    break_even_inputs_table.dropna(subset=["product"])
                    .drop_duplicates(subset=["product"], keep="last")
                    .reset_index(drop=True)
                )
        else:
            break_even_inputs_table = default_break_even.copy()
            try:
                tables.set_table("break_even_inputs", break_even_inputs_table)
            except Exception as exc:
                st.warning(f"Break-even defaults not applied due to validation error: {exc}")
            else:
                _update_editor_state("break_even_inputs", tables)

        def _be_value_changed(old: object, new: object) -> bool:
            if pd.isna(old) and pd.isna(new):
                return False
            if pd.isna(old) or pd.isna(new):
                return True
            try:
                return not math.isclose(float(old), float(new), rel_tol=1e-6, abs_tol=1e-4)
            except (TypeError, ValueError):
                return True

        break_even_tabs = st.tabs([PRODUCT_DISPLAY_NAMES.get(p, p.replace("_", " ").title()) for p in PRODUCTS])
        break_even_updated = False

        for tab, product in zip(break_even_tabs, PRODUCTS):
            with tab:
                label = PRODUCT_DISPLAY_NAMES.get(product, product.replace("_", " ").title())
                st.markdown(f"#### {label} break-even inputs")

                mask = break_even_inputs_table.get("product", pd.Series(dtype=str)) == product
                if mask.any():
                    row_index = break_even_inputs_table.index[mask][0]
                    row_series = break_even_inputs_table.loc[row_index]
                else:
                    default_row = default_break_even[default_break_even["product"] == product]
                    if default_row.empty:
                        row_series = pd.Series(
                            {
                                "product": product,
                                "unit_price": np.nan,
                                "variable_cost_per_unit": np.nan,
                                "fixed_cost": 0.0,
                                "reference_volume": np.nan,
                            }
                        )
                    else:
                        row_series = default_row.iloc[0]
                    row_index = None

                unit_price_default = float(row_series.get("unit_price", np.nan))
                if not np.isfinite(unit_price_default):
                    unit_price_default = 0.0
                variable_cost_default = float(row_series.get("variable_cost_per_unit", np.nan))
                if not np.isfinite(variable_cost_default):
                    variable_cost_default = 0.0
                fixed_cost_default = float(row_series.get("fixed_cost", 0.0) or 0.0)
                reference_volume_default = row_series.get("reference_volume", np.nan)

                actual_volume = 0.0
                if isinstance(per_product_be, pd.DataFrame) and not per_product_be.empty:
                    product_match = per_product_be[per_product_be["product"] == product]
                    if not product_match.empty:
                        actual_volume = float(product_match.iloc[0].get("actual_volume", 0.0))

                col_inputs_left, col_inputs_right = st.columns(2)
                with col_inputs_left:
                    price_step = max(abs(unit_price_default) * 0.05, 0.01)
                    unit_price_val = st.number_input(
                        "Unit price",
                        value=float(unit_price_default),
                        step=price_step,
                        format="%.4f",
                        key=f"break_even_unit_price_{product}",
                    )
                    variable_step = max(abs(variable_cost_default) * 0.05, 0.01)
                    variable_cost_val = st.number_input(
                        "Variable cost per unit",
                        value=float(variable_cost_default),
                        step=variable_step,
                        format="%.4f",
                        key=f"break_even_variable_cost_{product}",
                    )
                with col_inputs_right:
                    fixed_step = max(abs(fixed_cost_default) * 0.1, 1000.0)
                    fixed_cost_val = st.number_input(
                        "Fixed cost allocation",
                        value=float(fixed_cost_default),
                        min_value=0.0,
                        step=fixed_step,
                        format="%.2f",
                        key=f"break_even_fixed_cost_{product}",
                    )
                    reference_default = reference_volume_default
                    if not np.isfinite(reference_default) or reference_default <= 0:
                        reference_default = actual_volume if actual_volume > 0 else 0.0
                    volume_step = max(abs(reference_default) * 0.1, 1.0)
                    reference_volume_val = st.number_input(
                        "Reference production volume",
                        value=float(reference_default),
                        min_value=0.0,
                        step=volume_step,
                        format="%.2f",
                        key=f"break_even_reference_volume_{product}",
                    )

                new_record = {
                    "product": product,
                    "unit_price": float(unit_price_val),
                    "variable_cost_per_unit": float(variable_cost_val),
                    "fixed_cost": float(fixed_cost_val),
                    "reference_volume": float(reference_volume_val),
                }

                changed = row_index is None
                if not changed and row_index is not None:
                    for field, value in new_record.items():
                        existing = break_even_inputs_table.at[row_index, field]
                        if _be_value_changed(existing, value):
                            changed = True
                            break

                if changed:
                    if row_index is not None:
                        for field, value in new_record.items():
                            break_even_inputs_table.at[row_index, field] = value
                    else:
                        break_even_inputs_table = pd.concat(
                            [break_even_inputs_table, pd.DataFrame([new_record])],
                            ignore_index=True,
                        )
                    break_even_updated = True

                if isinstance(per_product_be, pd.DataFrame) and not per_product_be.empty:
                    product_match = per_product_be[per_product_be["product"] == product]
                else:
                    product_match = pd.DataFrame()

                if not product_match.empty:
                    result_row = product_match.iloc[0]
                    unit_table = pd.DataFrame(
                        [
                            {"Metric": "Unit of measure", "Value": result_row.get("unit_of_measure", "")},
                            {"Metric": "Unit price", "Value": result_row.get("unit_price")},
                            {
                                "Metric": "Variable cost per unit",
                                "Value": result_row.get("variable_cost_per_unit"),
                            },
                            {
                                "Metric": "Contribution margin per unit",
                                "Value": result_row.get("contribution_margin_per_unit"),
                            },
                            {
                                "Metric": "Actual average price",
                                "Value": result_row.get("actual_average_price"),
                            },
                        ]
                    )
                    st.dataframe(
                        unit_table,
                        use_container_width=True,
                        key=f"break_even_unit_table_{product}",
                    )

                    margin_percent = result_row.get("margin_of_safety_percent")
                    if np.isfinite(margin_percent):
                        margin_percent_display = margin_percent * 100.0
                    else:
                        margin_percent_display = np.nan
                    reference_ratio = result_row.get("break_even_vs_reference")
                    if np.isfinite(reference_ratio):
                        reference_ratio_display = reference_ratio * 100.0
                    else:
                        reference_ratio_display = np.nan

                    production_table = pd.DataFrame(
                        [
                            {"Metric": "Fixed cost allocation", "Value": result_row.get("fixed_cost")},
                            {"Metric": "Break-even volume (units)", "Value": result_row.get("break_even_units")},
                            {"Metric": "Reference volume", "Value": result_row.get("reference_volume")},
                            {"Metric": "Actual volume", "Value": result_row.get("actual_volume")},
                            {
                                "Metric": "Break-even vs reference (%)",
                                "Value": reference_ratio_display,
                            },
                            {
                                "Metric": "Margin of safety (units)",
                                "Value": result_row.get("margin_of_safety_units"),
                            },
                            {
                                "Metric": "Margin of safety (%)",
                                "Value": margin_percent_display,
                            },
                            {
                                "Metric": "Break-even revenue",
                                "Value": result_row.get("break_even_revenue"),
                            },
                            {
                                "Metric": "Actual revenue",
                                "Value": result_row.get("actual_revenue"),
                            },
                        ]
                    )
                    st.dataframe(
                        production_table,
                        use_container_width=True,
                        key=f"break_even_production_table_{product}",
                    )
                else:
                    st.info("No break-even results available for this product. Adjust inputs and rerun the model.")

        if break_even_updated:
            try:
                tables.set_table("break_even_inputs", break_even_inputs_table)
            except Exception as exc:
                st.warning(f"Break-even inputs not saved due to validation error: {exc}")
            else:
                _update_editor_state("break_even_inputs", tables)
                _safe_rerun()

        if isinstance(per_product_be, pd.DataFrame) and not per_product_be.empty:
            summary_cols = [
                "product",
                "unit_price",
                "variable_cost_per_unit",
                "contribution_margin_per_unit",
                "fixed_cost",
                "break_even_units",
                "reference_volume",
                "actual_volume",
                "break_even_revenue",
                "actual_revenue",
                "margin_of_safety_percent",
                "break_even_vs_reference",
            ]
            available_cols = [col for col in summary_cols if col in per_product_be.columns]
            summary_df = per_product_be[available_cols].copy()
            summary_df["product"] = summary_df["product"].map(PRODUCT_DISPLAY_NAMES).fillna(
                summary_df["product"].str.replace("_", " ").str.title()
            )
            if "margin_of_safety_percent" in summary_df:
                summary_df["margin_of_safety_percent"] = summary_df["margin_of_safety_percent"] * 100.0
            if "break_even_vs_reference" in summary_df:
                summary_df["break_even_vs_reference"] = summary_df["break_even_vs_reference"] * 100.0
            summary_df = summary_df.rename(
                columns={
                    "product": "Product",
                    "unit_price": "Unit price",
                    "variable_cost_per_unit": "Variable cost per unit",
                    "contribution_margin_per_unit": "Contribution margin per unit",
                    "fixed_cost": "Fixed cost",
                    "break_even_units": "Break-even units",
                    "reference_volume": "Reference volume",
                    "actual_volume": "Actual volume",
                    "break_even_revenue": "Break-even revenue",
                    "actual_revenue": "Actual revenue",
                    "margin_of_safety_percent": "Margin of safety (%)",
                    "break_even_vs_reference": "Break-even vs reference (%)",
                }
            )
            _render_dataframe(summary_df, "Break-even summary by product", key="break_even_summary")
            _render_break_even_bar_chart(per_product_be)
            _render_break_even_comparison_chart(per_product_be)

        if isinstance(overall_be, dict) and overall_be:
            overall_table = pd.DataFrame(
                [
                    {
                        "Metric": "Total fixed costs",
                        "Value": overall_be.get("fixed_costs_total"),
                    },
                    {
                        "Metric": "Variable cost per unit",
                        "Value": overall_be.get("variable_cost_per_unit"),
                    },
                    {
                        "Metric": "Average price",
                        "Value": overall_be.get("average_price"),
                    },
                    {
                        "Metric": "Break-even volume (total)",
                        "Value": overall_be.get("break_even_volume"),
                    },
                    {
                        "Metric": "Margin of safety (%)",
                        "Value": overall_be.get("margin_of_safety", np.nan) * 100.0
                        if overall_be.get("margin_of_safety") is not None
                        else np.nan,
                    },
                ]
            )
            st.dataframe(
                overall_table,
                use_container_width=True,
                key="break_even_overall_table",
            )

    with sensitivity_tab:
        st.subheader("Advanced sensitivity analytics")
        sensitivity_sections = st.tabs(
            [
                "Metaheuristic optimiser",
                "Neural forecasts",
                "Statistical forecasts",
                "Decision tree",
                "Sensitivity tornado",
                "Monte Carlo simulation",
                "Scenario comparison",
            ]
        )

        with sensitivity_sections[0]:
            optimizer_column_config = {
                "enabled": st.column_config.CheckboxColumn("Enabled"),
                "variable": st.column_config.SelectboxColumn(
                    "Variable",
                    options=list(OPTIMIZER_VARIABLE_OPTIONS.keys()),
                    format_func=lambda key: OPTIMIZER_VARIABLE_OPTIONS.get(key, key.replace("_", " ").title()),
                ),
                "lower_bound": st.column_config.NumberColumn("Lower bound"),
                "upper_bound": st.column_config.NumberColumn("Upper bound"),
                "notes": st.column_config.TextColumn("Notes"),
            }
            _render_table_editor(
                tables,
                "optimizer_settings",
                "Optimizer variables",
                sync_errors.get("optimizer_settings"),
                "Define variable bounds for the metaheuristic optimiser.",
                column_config=optimizer_column_config,
            )

            optimizer_table = tables.ensure_table("optimizer_settings").copy()
            enabled_mask = optimizer_table.get("enabled", True)
            if not isinstance(enabled_mask, pd.Series):
                enabled_mask = pd.Series(True, index=optimizer_table.index)
            optimizer_active = optimizer_table[enabled_mask.fillna(True)]
            if optimizer_active.empty:
                st.info("Add at least one enabled variable to run the optimiser.")
            else:
                objective_options = list(metrics.keys())
                default_objective = "Project_NPV" if "Project_NPV" in objective_options else objective_options[0]
                selected_objective = st.selectbox(
                    "Objective metric",
                    objective_options,
                    index=objective_options.index(default_objective),
                )
                iterations = int(
                    st.number_input("Iterations", min_value=1, value=10, step=1, key="optimizer_iterations")
                )
                population = int(
                    st.number_input("Population size", min_value=2, value=6, step=1, key="optimizer_population")
                )
                seed_value = st.number_input("Random seed (optional)", value=42, step=1, key="optimizer_seed")
                if st.button("Run metaheuristic optimisation", key="run_optimizer"):
                    with st.spinner("Running optimisation across variable bounds..."):
                        optimiser_results = metaheuristic_optimize(
                            cfg,
                            metrics,
                            lambda c: run_full_model(c),
                            optimizer_active,
                            selected_objective,
                            iterations=iterations,
                            population=population,
                            seed=int(seed_value),
                        )
                    if optimiser_results.empty:
                        st.info("Optimiser did not produce any results with the current configuration.")
                    else:
                        _render_dataframe(
                            optimiser_results,
                            f"Optimiser results – {selected_objective}",
                            key="optimizer_results",
                        )
                        best_row = optimiser_results.sort_values("objective", ascending=False).iloc[0]
                        metric_kind = _metric_kind(selected_objective)
                        st.metric(
                            f"Best {selected_objective}",
                            _format_metric(best_row.get("objective"), metric_kind),
                            delta=_format_metric(best_row.get("delta_vs_base"), metric_kind),
                        )

        with sensitivity_sections[1]:
            neural_column_config = {
                "enabled": st.column_config.CheckboxColumn("Enabled"),
                "product": st.column_config.SelectboxColumn(
                    "Product",
                    options=list(PRODUCTS),
                    format_func=lambda key: key.replace("_", " ").title(),
                ),
                "lookback_months": st.column_config.NumberColumn("Lookback (months)", min_value=3, step=1),
                "forecast_months": st.column_config.NumberColumn("Forecast horizon (months)", min_value=1, step=1),
                "hidden_units": st.column_config.NumberColumn("Hidden units", min_value=1, step=1),
                "learning_rate": st.column_config.NumberColumn("Learning rate", min_value=0.0001, step=0.0001, format="%.4f"),
                "epochs": st.column_config.NumberColumn("Epochs", min_value=50, step=50),
            }
            _render_table_editor(
                tables,
                "neural_forecast_settings",
                "Neural forecasting settings",
                sync_errors.get("neural_forecast_settings"),
                "Configure neural network parameters to forecast production volumes.",
                column_config=neural_column_config,
            )

            neural_cfg = tables.ensure_table("neural_forecast_settings").copy()
            enabled_mask = neural_cfg.get("enabled", True)
            if not isinstance(enabled_mask, pd.Series):
                enabled_mask = pd.Series(True, index=neural_cfg.index)
            neural_active = neural_cfg[enabled_mask.fillna(True)]
            if neural_active.empty:
                st.info("Enable at least one neural forecast configuration to generate projections.")
            elif st.button("Run neural forecasts", key="run_neural"):
                for idx, row in neural_active.iterrows():
                    product_key = str(row.get("product", "ethanol") or "ethanol")
                    lookback_val = int(row.get("lookback_months", 12) or 12)
                    horizon_val = int(row.get("forecast_months", 12) or 12)
                    hidden_units = int(row.get("hidden_units", 8) or 8)
                    learning_rate = float(row.get("learning_rate", 0.01) or 0.01)
                    epochs_val = int(row.get("epochs", 300) or 300)
                    with st.spinner(f"Forecasting {product_key.title()} volumes..."):
                        forecast_df = neural_forecast_production(
                            results,
                            product_key,
                            lookback=lookback_val,
                            horizon=horizon_val,
                            hidden_units=hidden_units,
                            learning_rate=learning_rate,
                            epochs=epochs_val,
                            seed=int(idx + 1),
                        )
                    if forecast_df.empty:
                        st.warning(f"No forecast generated for {product_key.title()} (insufficient data).")
                    else:
                        _render_dataframe(
                            forecast_df,
                            f"Neural forecast – {product_key.title()}",
                            key=f"neural_{idx}",
                        )

        with sensitivity_sections[2]:
            stat_column_config = {
                "enabled": st.column_config.CheckboxColumn("Enabled"),
                "series": st.column_config.SelectboxColumn(
                    "Series",
                    options=list(STATISTICAL_SERIES_OPTIONS.keys()),
                    format_func=lambda key: STATISTICAL_SERIES_OPTIONS.get(key, key.replace("_", " ").title()),
                ),
                "alpha": st.column_config.NumberColumn("Alpha", min_value=0.01, max_value=1.0, format="%.2f"),
                "forecast_months": st.column_config.NumberColumn("Forecast horizon (months)", min_value=1, step=1),
            }
            _render_table_editor(
                tables,
                "statistical_forecast_settings",
                "Statistical forecast settings",
                sync_errors.get("statistical_forecast_settings"),
                "Set exponential-smoothing parameters for demand and cost forecasting.",
                column_config=stat_column_config,
            )

            stat_cfg = tables.ensure_table("statistical_forecast_settings").copy()
            enabled_mask = stat_cfg.get("enabled", True)
            if not isinstance(enabled_mask, pd.Series):
                enabled_mask = pd.Series(True, index=stat_cfg.index)
            stat_active = stat_cfg[enabled_mask.fillna(True)]
            if stat_active.empty:
                st.info("Enable at least one statistical series to forecast.")
            elif st.button("Run statistical forecasts", key="run_statistical"):
                for idx, row in stat_active.iterrows():
                    series_key = str(row.get("series", "revenue") or "revenue")
                    alpha_val = float(row.get("alpha", 0.3) or 0.3)
                    horizon_val = int(row.get("forecast_months", 12) or 12)
                    with st.spinner(f"Forecasting {STATISTICAL_SERIES_OPTIONS.get(series_key, series_key)}..."):
                        forecast_result = statistical_forecast(
                            results,
                            series_key,
                            alpha=alpha_val,
                            horizon=horizon_val,
                        )
                    historical_df = forecast_result.get("historical", pd.DataFrame())
                    forecast_df = forecast_result.get("forecast", pd.DataFrame())
                    if historical_df.empty and forecast_df.empty:
                        st.warning(f"No data available for {STATISTICAL_SERIES_OPTIONS.get(series_key, series_key)}.")
                        continue
                    if not historical_df.empty:
                        _render_dataframe(
                            historical_df,
                            f"Historical series – {STATISTICAL_SERIES_OPTIONS.get(series_key, series_key)}",
                            key=f"stat_hist_{idx}",
                        )
                    if not forecast_df.empty:
                        _render_dataframe(
                            forecast_df,
                            f"Forecast – {STATISTICAL_SERIES_OPTIONS.get(series_key, series_key)}",
                            key=f"stat_forecast_{idx}",
                        )
                    st.caption(
                        f"Residual standard deviation: {forecast_result.get('residual_std', 0.0):.2f} · "
                        f"Equipment failure risk proxy: {forecast_result.get('equipment_failure_risk', 0.0):.2%}"
                    )

        with sensitivity_sections[3]:
            decision_column_config = {
                "enabled": st.column_config.CheckboxColumn("Enabled"),
                "path_name": st.column_config.TextColumn("Path name"),
                "probability": st.column_config.NumberColumn("Probability", min_value=0.0, max_value=1.0, format="%.2f"),
                "ethanol_price_multiplier": st.column_config.NumberColumn("Ethanol price multiplier", format="%.3f"),
                "sugar_price_multiplier": st.column_config.NumberColumn("Sugar price multiplier", format="%.3f"),
                "electricity_price_multiplier": st.column_config.NumberColumn("Electricity price multiplier", format="%.3f"),
                "animal_feed_price_multiplier": st.column_config.NumberColumn("Animal feed price multiplier", format="%.3f"),
                "capex_multiplier": st.column_config.NumberColumn("CAPEX multiplier", format="%.3f"),
                "opex_multiplier": st.column_config.NumberColumn("Opex multiplier", format="%.3f"),
                "debt_rate_shift": st.column_config.NumberColumn("Debt rate shift", format="%.4f"),
                "notes": st.column_config.TextColumn("Notes"),
            }
            _render_table_editor(
                tables,
                "decision_tree_paths",
                "Decision tree paths",
                sync_errors.get("decision_tree_paths"),
                "Define scenario branches with probabilities and multipliers for pricing, CAPEX, and OPEX.",
                column_config=decision_column_config,
            )

            decision_table = tables.ensure_table("decision_tree_paths").copy()
            if decision_table.empty:
                st.info("Add decision paths to evaluate expected outcomes.")
            else:
                objective_options = list(metrics.keys())
                default_objective = "Project_NPV" if "Project_NPV" in objective_options else objective_options[0]
                selected_objective = st.selectbox(
                    "Objective metric",
                    objective_options,
                    index=objective_options.index(default_objective),
                    key="decision_objective",
                )
                if st.button("Evaluate decision tree", key="run_decision_tree"):
                    with st.spinner("Evaluating decision tree paths..."):
                        decision_result = decision_tree_analysis(
                            cfg,
                            lambda c: run_full_model(c),
                            decision_table,
                            objective=selected_objective,
                        )
                    paths_df = decision_result.get("paths", pd.DataFrame())
                    if paths_df.empty:
                        st.info("No active decision paths with valid probabilities were found.")
                    else:
                        _render_dataframe(paths_df, "Decision tree evaluation", key="decision_tree_results")
                        metric_kind = _metric_kind(selected_objective)
                        st.metric(
                            f"Expected {selected_objective}",
                            _format_metric(decision_result.get("expected_metric"), metric_kind),
                        )
                        st.caption(
                            f"Total probability weight: {decision_result.get('total_probability', 0.0):.2f}"
                        )
                        _render_decision_tree_chart(paths_df, selected_objective)

        with sensitivity_sections[4]:
            _render_table_editor(
                tables,
                "tornado_drivers",
                "Tornado drivers",
                sync_errors.get("tornado_drivers"),
                "Enable drivers and adjust percentage shocks to analyse NPV sensitivity.",
            )
            tornado_df_cfg = tables.ensure_table("tornado_drivers").copy()
            if tornado_df_cfg.empty:
                st.info("Add at least one driver row to evaluate the tornado chart.")
            else:
                enabled_mask = tornado_df_cfg.get("enabled", True)
                if not isinstance(enabled_mask, pd.Series):
                    enabled_mask = pd.Series(True, index=tornado_df_cfg.index)
                enabled_rows = tornado_df_cfg[enabled_mask.fillna(True)]
                enabled_rows = enabled_rows.replace({"": np.nan})
                enabled_rows = enabled_rows.dropna(subset=["driver", "pct_change"], how="any")
                drivers: List[Tuple[str, float]] = []
                for _, driver_row in enabled_rows.iterrows():
                    driver_key = str(driver_row.get("driver", "")).strip()
                    if not driver_key:
                        continue
                    driver_key = normalize_key(driver_key)
                    try:
                        pct_change = float(driver_row.get("pct_change", 0.0))
                    except (TypeError, ValueError):
                        pct_change = 0.0
                    drivers.append((driver_key, pct_change))
                if not drivers:
                    st.info("Enable at least one driver with a valid percentage change to run the tornado analysis.")
                else:
                    with st.spinner("Calculating sensitivity tornado..."):
                        tornado_results = sensitivity_tornado(
                            cfg,
                            {"metrics": metrics},
                            lambda c: run_full_model(c),
                            drivers,
                        )
                    _render_dataframe(tornado_results, "Tornado sensitivity", key="tornado")
                    _render_tornado_chart(tornado_results)

        with sensitivity_sections[5]:
            distribution_options = list(MONTE_CARLO_DISTRIBUTIONS)
            variable_label_map = dict(MONTE_CARLO_VARIABLE_LABELS)
            variable_options = list(variable_label_map.values())
            applies_options = ["global", *PRODUCTS]
            monte_column_config = {
                "enabled": st.column_config.CheckboxColumn("Enabled"),
                "distribution": st.column_config.SelectboxColumn(
                    "Probability distribution",
                    options=distribution_options,
                    help="Select the probability distribution used to draw this Monte Carlo driver.",
                ),
                "variable": st.column_config.SelectboxColumn(
                    "Variable",
                    options=variable_options,
                    help="Choose which driver the sampled draw should adjust during the simulation.",
                ),
                "applies_to": st.column_config.SelectboxColumn(
                    "Applies to",
                    options=applies_options,
                    help="Restrict the driver to a specific product or leave as global for portfolio-wide adjustments.",
                ),
                "p1": st.column_config.NumberColumn("P1", help="Distribution parameter 1 (mean/min)."),
                "p2": st.column_config.NumberColumn("P2", help="Distribution parameter 2 (std/mode/max)."),
                "p3": st.column_config.NumberColumn("P3", help="Distribution parameter 3 (triangular max)."),
            }
            _render_table_editor(
                tables,
                "monte_carlo_settings",
                "Monte Carlo settings",
                sync_errors.get("monte_carlo_settings"),
                "Set iterations, random seed, and enable the simulation to sample risk drivers.",
                column_config=monte_column_config,
            )
            monte_cfg = tables.ensure_table("monte_carlo_settings").copy()
            if not monte_cfg.empty:
                row_indices = list(monte_cfg.index)

                def _format_driver(idx: int) -> str:
                    row = monte_cfg.loc[idx]
                    variable_key = str(row.get("variable", "driver") or "driver")
                    variable = variable_label_map.get(variable_key, variable_key)
                    distribution = str(row.get("distribution", "normal") or "normal")
                    scope = str(row.get("applies_to", "global") or "global")
                    return f"Row {idx + 1}: {variable} · {distribution} ({scope})"

                selected_idx = st.selectbox(
                    "Select Monte Carlo driver to edit",
                    row_indices,
                    format_func=_format_driver,
                    key="monte_driver_select",
                )

                selected_row = monte_cfg.loc[selected_idx]
                with st.form(f"monte_driver_form_{selected_idx}"):
                    enabled_value = st.checkbox(
                        "Enable driver",
                        value=bool(selected_row.get("enabled", False)),
                        key=f"monte_enabled_{selected_idx}",
                    )
                    iterations_value = st.number_input(
                        "Iterations",
                        min_value=1,
                        value=int(selected_row.get("iterations", 1000)
                                  if pd.notna(selected_row.get("iterations"))
                                  else 1000),
                        step=1,
                        key=f"monte_iterations_{selected_idx}",
                    )
                    random_seed_value = st.number_input(
                        "Random seed",
                        value=int(selected_row.get("random_seed", 42)
                                  if pd.notna(selected_row.get("random_seed"))
                                  else 42),
                        step=1,
                        key=f"monte_seed_{selected_idx}",
                    )
                    distribution_value = st.selectbox(
                        "Probability distribution type",
                        distribution_options,
                        index=_option_index(
                            distribution_options,
                            str(selected_row.get("distribution", "normal") or "normal"),
                        ),
                        key=f"monte_distribution_{selected_idx}",
                    )
                    selected_key = str(selected_row.get("variable", "") or "")
                    selected_label = variable_label_map.get(selected_key, variable_options[0])
                    variable_value = st.selectbox(
                        "Variable",
                        variable_options,
                        index=_option_index(variable_options, selected_label),
                        key=f"monte_variable_{selected_idx}",
                    )
                    applies_value = st.selectbox(
                        "Applies to",
                        applies_options,
                        index=_option_index(
                            applies_options,
                            str(selected_row.get("applies_to", "global") or "global"),
                        ),
                        key=f"monte_applies_{selected_idx}",
                    )
                    p1_value = st.number_input(
                        "P1",
                        value=float(selected_row.get("p1", 0.0)
                                    if pd.notna(selected_row.get("p1"))
                                    else 0.0),
                        key=f"monte_p1_{selected_idx}",
                    )
                    p2_value = st.number_input(
                        "P2",
                        value=float(selected_row.get("p2", 0.05)
                                    if pd.notna(selected_row.get("p2"))
                                    else 0.05),
                        key=f"monte_p2_{selected_idx}",
                    )
                    p3_value = st.number_input(
                        "P3",
                        value=float(selected_row.get("p3", 0.0)
                                    if pd.notna(selected_row.get("p3"))
                                    else 0.0),
                        key=f"monte_p3_{selected_idx}",
                    )
                    submitted = st.form_submit_button("Save Monte Carlo driver")

                if submitted:
                    updated_cfg = monte_cfg.copy()
                    updated_cfg.at[selected_idx, "enabled"] = bool(enabled_value)
                    updated_cfg.at[selected_idx, "iterations"] = int(iterations_value)
                    updated_cfg.at[selected_idx, "random_seed"] = int(random_seed_value)
                    updated_cfg.at[selected_idx, "distribution"] = str(distribution_value)
                    updated_cfg.at[selected_idx, "variable"] = MONTE_CARLO_VARIABLE_LABEL_TO_KEY.get(
                        str(variable_value), str(variable_value)
                    )
                    updated_cfg.at[selected_idx, "applies_to"] = str(applies_value)
                    updated_cfg.at[selected_idx, "p1"] = float(p1_value)
                    updated_cfg.at[selected_idx, "p2"] = float(p2_value)
                    updated_cfg.at[selected_idx, "p3"] = float(p3_value)
                    tables.set_table("monte_carlo_settings", updated_cfg)
                    _update_editor_state("monte_carlo_settings", tables)
                    st.success("Monte Carlo driver updated.")
                    _safe_rerun()
            enabled_mc = pd.Series(dtype=bool)
            if not monte_cfg.empty:
                enabled_mc = monte_cfg.get("enabled", False)
                if not isinstance(enabled_mc, pd.Series):
                    enabled_mc = pd.Series(False, index=monte_cfg.index)
                enabled_mc = enabled_mc.fillna(False)
            if monte_cfg.empty or not enabled_mc.any():
                st.info("Enable a Monte Carlo row to execute the simulation.")
            else:
                active_row = monte_cfg.loc[enabled_mc].iloc[0]
                try:
                    iterations = int(float(active_row.get("iterations", 1000)))
                except (TypeError, ValueError):
                    iterations = 1000
                iterations = max(iterations, 1)
                try:
                    random_seed = int(float(active_row.get("random_seed", 42)))
                except (TypeError, ValueError):
                    random_seed = 42
                with st.spinner("Running Monte Carlo simulation..."):
                    monte_results = monte_carlo(
                        cfg,
                        lambda c: run_full_model(c),
                        iterations=iterations,
                        random_seed=random_seed,
                    )
                percentiles = (
                    monte_results["percentiles"].reset_index().rename(columns={"index": "Percentile"})
                )
                _render_dataframe(percentiles, "Monte Carlo percentiles", key="monte_percentiles")
                _render_dataframe(monte_results["samples"], "Monte Carlo samples", key="monte_samples")
                _render_monte_carlo_histograms(monte_results["samples"])

        with sensitivity_sections[6]:
            _render_table_editor(
                tables,
                "scenario_comparison",
                "Scenario definitions",
                sync_errors.get("scenario_comparison"),
                "Toggle and edit scenario overrides (feedstock sourcing and farm share) for comparison against the base case.",
            )
            scenario_cfg = tables.ensure_table("scenario_comparison").copy()
            scenario_overrides_active = _scenario_overrides_from_table(scenario_cfg)
            if not scenario_overrides_active:
                st.info("Add scenario rows with overrides to compare against the base configuration.")
            else:
                with st.spinner("Evaluating scenarios..."):
                    scenario_df = run_scenarios(
                        cfg,
                        lambda c: run_full_model(c),
                        scenario_overrides_active,
                    )
                base_metrics = pd.DataFrame([metrics]).assign(scenario=BASE_SCENARIO_LABEL)
                scenario_df = pd.concat([base_metrics, scenario_df], ignore_index=True)
                _render_dataframe(scenario_df, "Scenario comparison", key="scenarios")
                scenario_results_map: Dict[str, Mapping[str, object]] = {}
                for scenario_name in scenario_df["scenario"].dropna().astype(str).unique():
                    cfg_payload, res_payload = _ensure_scenario_payload(
                        scenario_name,
                        cfg,
                        results,
                        scenario_overrides_active,
                    )
                    scenario_results_map[scenario_name] = res_payload
                _render_scenario_metric_chart(scenario_df)
                if scenario_results_map:
                    _render_scenario_cashflow_stack(scenario_results_map)
                    _render_scenario_dscr_chart(scenario_results_map)
                    _render_scenario_scatter_chart(scenario_results_map)

    with chatbot_tab:
        _render_chatbot_tab(results)

    st.success("Model run complete.")


if __name__ == "__main__":  # pragma: no cover - manual invocation helper
    if _streamlit_runtime_exists():
        main()
    else:
        try:
            from streamlit.web import bootstrap  # type: ignore[attr-defined]

            bootstrap.run(__file__, "", [])
        except ModuleNotFoundError as exc:
            print(
                "Streamlit is not installed in this environment. Install it with "
                "'pip install streamlit' and re-run the app.",
                file=sys.stderr,
            )
            raise SystemExit(1) from exc
        except Exception as exc:  # pragma: no cover - defensive feedback
            print(
                "Unable to start the Streamlit runtime automatically. "
                "If the Streamlit CLI is unavailable, try installing Streamlit or "
                "launching via 'python -m streamlit run streamlit_app.py'.",
                file=sys.stderr,
            )
            raise SystemExit(1) from exc
