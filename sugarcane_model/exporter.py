"""Banker-friendly Excel and PDF exports for the refactored model."""
from __future__ import annotations

import importlib.util
from io import BytesIO
from pathlib import Path
import sys
from typing import Any

import pandas as pd

from .driver_schedules import SCHEDULE_DEFINITIONS
from .inputs import SugarcaneBioethanolInputs, input_values


def _safe_sheet_name(name: str, used: set[str]) -> str:
    base = name.replace("&", "and")[:31]
    candidate = base
    counter = 2
    while candidate in used:
        suffix = f"_{counter}"
        candidate = f"{base[:31-len(suffix)]}{suffix}"
        counter += 1
    used.add(candidate)
    return candidate


def build_excel_report(
    inputs: SugarcaneBioethanolInputs,
    metrics: dict[str, Any],
    tables: dict[str, pd.DataFrame],
) -> bytes:
    from openpyxl.styles import Alignment, Font, PatternFill

    output = BytesIO()
    used: set[str] = set()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        summary = pd.DataFrame(
            [{"Metric": key, "Value": value} for key, value in metrics.items()]
        )
        summary.to_excel(writer, sheet_name="Summary", index=False)
        used.add("Summary")

        payload = input_values(inputs)
        for section, values in payload.items():
            if section == "yearly_schedules":
                for schedule_key, rows in values.items():
                    label = SCHEDULE_DEFINITIONS.get(schedule_key, {}).get(
                        "label", schedule_key.replace("_", " ").title()
                    )
                    pd.DataFrame(rows).to_excel(
                        writer,
                        sheet_name=_safe_sheet_name(f"{label} Drivers", used),
                        index=False,
                    )
                continue

            if section == "capex":
                frame = pd.DataFrame(values.get("items", []))
            elif section == "financing":
                facilities = values.get("additional_debt_facilities", [])
                frame = pd.DataFrame(
                    [
                        {"Assumption": key, "Value": value}
                        for key, value in values.items()
                        if key != "additional_debt_facilities"
                    ]
                )
            else:
                frame = pd.DataFrame(
                    [{"Assumption": key, "Value": value} for key, value in values.items()]
                )
            frame.to_excel(
                writer,
                sheet_name=_safe_sheet_name(section.replace("_", " ").title(), used),
                index=False,
            )
            if section == "financing":
                pd.DataFrame(facilities).to_excel(
                    writer,
                    sheet_name=_safe_sheet_name("Additional Debt Inputs", used),
                    index=False,
                )

        for name, frame in tables.items():
            export_frame = frame.copy()
            if isinstance(export_frame.index, pd.DatetimeIndex):
                export_frame = export_frame.reset_index(names="Date")
            elif export_frame.index.name is not None:
                export_frame = export_frame.reset_index()
            export_frame.to_excel(
                writer,
                sheet_name=_safe_sheet_name(name, used),
                index=False,
            )

        header_fill = PatternFill("solid", fgColor="143D2A")
        section_fill = PatternFill("solid", fgColor="DDEBE3")
        for sheet in writer.book.worksheets:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.fill = header_fill
                cell.font = Font(color="FFFFFF", bold=True)
                cell.alignment = Alignment(horizontal="center")
            if sheet.title == "Summary":
                for cell in sheet[2]:
                    cell.fill = section_fill
            for column in sheet.columns:
                cells = list(column)
                width = max(
                    len(str(cell.value)) if cell.value is not None else 0
                    for cell in cells[:250]
                )
                sheet.column_dimensions[cells[0].column_letter].width = min(
                    max(11, width + 2), 34
                )
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    if isinstance(cell.value, float):
                        cell.number_format = "#,##0.00;[Red](#,##0.00);-"
    return output.getvalue()


def _load_pdf_style_module():
    module_name = "_sugarcane_refactored_pdf_style"
    cached = sys.modules.get(module_name)
    if cached is not None and hasattr(cached, "body"):
        return cached
    sys.modules.pop(module_name, None)
    search_roots = [
        *Path(__file__).resolve().parents,
        *(Path(entry) for entry in sys.path if entry),
    ]
    path = next(
        (
            root / "app" / "reports" / "pdf_style.py"
            for root in search_roots
            if (root / "app" / "reports" / "pdf_style.py").is_file()
        ),
        None,
    )
    if path is None:
        raise ModuleNotFoundError("NumQuants PDF style module is not available")
    spec = importlib.util.spec_from_file_location(module_name, path)
    if spec is None or spec.loader is None:
        raise ImportError(f"Cannot load PDF style module from {path}")
    module = importlib.util.module_from_spec(spec)
    sys.modules[module_name] = module
    try:
        spec.loader.exec_module(module)
    except Exception:
        sys.modules.pop(module_name, None)
        raise
    return module


def _fallback_pdf(metrics: dict[str, Any], user_email: str) -> bytes:
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.backends.backend_pdf import PdfPages

    output = BytesIO()
    rows = [
        ("Scenario", metrics.get("scenario", "n/a")),
        ("Project NPV", f"USD {float(metrics.get('project_npv', 0.0)):,.0f}"),
        ("Project IRR", "n/a" if metrics.get("project_irr") is None else f"{metrics['project_irr']:.1%}"),
        ("Equity IRR", "n/a" if metrics.get("equity_irr") is None else f"{metrics['equity_irr']:.1%}"),
        ("Minimum DSCR", "n/a" if metrics.get("min_dscr") is None else f"{metrics['min_dscr']:.2f}x"),
        ("Total CAPEX", f"USD {float(metrics.get('total_capex', 0.0)):,.0f}"),
        ("Model status", metrics.get("model_status", "n/a")),
    ]
    with PdfPages(output) as pdf:
        figure = plt.figure(figsize=(8.27, 11.69), facecolor="white")
        figure.text(0.08, 0.93, "Sugar Cane Bioethanol", fontsize=22, weight="bold", color="#143D2A")
        figure.text(0.08, 0.895, "Cassava-style modular financial model", fontsize=12, color="#4B6357")
        figure.text(0.08, 0.865, f"Prepared for {user_email}", fontsize=9, color="#6B7280")
        y = 0.78
        for label, value in rows:
            figure.text(0.10, y, str(label), fontsize=11, weight="bold", color="#143D2A")
            figure.text(0.48, y, str(value), fontsize=11, color="#111827")
            y -= 0.055
        figure.text(
            0.08,
            0.08,
            "The Excel pack contains all assumptions, operating schedules, component financials, consolidated statements, and model checks.",
            fontsize=8,
            color="#6B7280",
            wrap=True,
        )
        plt.axis("off")
        pdf.savefig(figure, bbox_inches="tight")
        plt.close(figure)
    return output.getvalue()


def build_pdf_report(
    metrics: dict[str, Any],
    tables: dict[str, pd.DataFrame],
    user_email: str,
    watermark: str | None = None,
) -> bytes:
    try:
        pdf_style = _load_pdf_style_module()
        from reportlab.platypus import PageBreak
    except ModuleNotFoundError:
        return _fallback_pdf(metrics, user_email)

    headline = {
        "Project NPV": f"USD {float(metrics.get('project_npv', 0.0)):,.0f}",
        "Project IRR": "n/a" if metrics.get("project_irr") is None else f"{metrics['project_irr']:.1%}",
        "Equity IRR": "n/a" if metrics.get("equity_irr") is None else f"{metrics['equity_irr']:.1%}",
        "Minimum DSCR": "n/a" if metrics.get("min_dscr") is None else f"{metrics['min_dscr']:.2f}x",
    }
    sections = [
        pdf_style.heading("Executive summary"),
        pdf_style.body(
            "Integrated farm, sourcing, processing, commercialization, component-financial, and consolidated-financial model."
        ),
        pdf_style.metric_grid(headline, columns=4),
        PageBreak(),
        pdf_style.heading("Scenario comparison"),
        *pdf_style.dataframe_table(tables["Scenario Comparison"], max_rows=10),
        PageBreak(),
        pdf_style.heading("Component financial summary"),
        *pdf_style.dataframe_table(tables["Component Financials"], max_rows=42),
        PageBreak(),
        pdf_style.heading("Consolidated profit and loss"),
        *pdf_style.dataframe_table(tables["Consolidated P&L"], max_rows=20),
        PageBreak(),
        pdf_style.heading("Model checks"),
        *pdf_style.dataframe_table(tables["Checks"], max_rows=20),
    ]
    return pdf_style.build_pdf(
        title="Sugar Cane Bioethanol Financial Model",
        subtitle=f"Prepared for {user_email}",
        sections=sections,
        watermark=watermark,
    )
